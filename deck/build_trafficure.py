# CFI x TraffiCure / Google — discussion deck for the working call. 15 slides.
# Ideas to explore together, led by the School Safety Index; every idea
# pressure-tested. Built on the NRSB white-paper design system.
import os
import build_nrsb3 as B
import nrsb2_style as ST
from nrsb2_style import *
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

I=Inches
V6=os.path.join(SP,"assets/v6/ready")
RMI=os.path.join(SP,"assets/rmi")
def a6(n): return os.path.join(V6,n)
def ar(n): return os.path.join(RMI,n)

TOTAL=15
AMB="#B26A00"

def ftc(s,idx):
    hline(s,ML,I(7.06),CW)
    text(s,ML,I(7.14),I(9),I(0.25),"Crashfree India  ·  Discussion deck · Crashfree India × TraffiCure, Google Maps Platform",size=8.5,color=MUT,ls=1.0)
    text(s,SW-MR-I(1.2),I(7.14),I(1.2),I(0.25),f"{idx:02d} / {TOTAL:02d}",size=8.5,color=MUT,align=PP_ALIGN.RIGHT,ls=1.0)

B.footer=lambda s,wg,idx,total: ftc(s,idx)
def _sect_tab(s,cur,name):
    w=I(3.6)
    text(s,SW-MR-w,I(0.30),w,I(0.24),name,font=HEAD,size=8.5,bold=True,color=MUT,tracking=1.6,align=PP_ALIGN.RIGHT)
ST.sect_tab=_sect_tab

def patch_rail(s,old,new):
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text and old in r.text:
                        r.text=r.text.replace(old,new)

# --- pressure-test panel: consistent right column on every idea slide ---
def ptest(s, x, y, w, h, rows):
    rect(s,x,y,w,h,fill=AM_T,rounded=True,radius=0.06)
    chip(s,x+I(0.2),y+I(0.18),I(1.7),I(0.34),"PRESSURE TEST",fill=RGBColor(0xB2,0x6A,0x00),color=RGBColor(0xFF,0xFF,0xFF),size=8.5)
    labels=[("shield-check","Must be true"),("file-x","Where it breaks"),("timer","How we find out fast")]
    yy=y+I(0.68)
    per=(h-I(0.78))/3
    for (ic,lab),body in zip(labels,rows):
        icon(s,ic,x+I(0.2),yy+I(0.02),I(0.26),AMB)
        text(s,x+I(0.56),yy,w-I(0.76),I(0.26),lab,font=HEAD,size=9.5,bold=True,color=INK)
        text(s,x+I(0.56),yy+I(0.27),w-I(0.76),per-I(0.3),body,size=8.5,color=MUT,ls=1.18)
        yy=yy+per
    return s

# ---------------------------------------------------------------- 01 COVER
def cover():
    s=slide()
    s.shapes.add_picture(LOGO_BLUE,ML,I(0.5),width=I(1.85))
    text(s,SW-MR-I(3.5),I(0.55),I(3.5),I(0.3),"AUGUST 2026  ·  NEW DELHI",font=HEAD,size=9,bold=True,color=MUT,tracking=1.8,align=PP_ALIGN.RIGHT)
    text(s,ML,I(1.72),I(8.4),I(0.3),"FOR DISCUSSION  ·  CRASHFREE INDIA × TRAFFICURE · GOOGLE MAPS PLATFORM",font=HEAD,size=10,bold=True,color=BRAND,tracking=2.0)
    text(s,ML,I(2.14),I(8.3),I(0.85),"The safety layer",font=HEAD,size=46,bold=True,color=INK,ls=1.0)
    text(s,ML,I(3.1),I(8.3),I(0.85),"Traffic data already manages congestion.\nTogether it can start preventing crashes.",font=HEAD,size=16.5,bold=True,color=BRAND,ls=1.25)
    rows=[("gauge","TraffiCure has proved the traffic layer: Pune from 20 to 26.8 km/h, 964 roads live"),
          ("zap","RMI's May 2026 feeds carry safety signals: harsh braking, vehicle counts, incidents"),
          ("school","We bring the safety products and the government seats to put them to work")]
    y=I(4.65)
    for i,(ic,t) in enumerate(rows):
        icon_disc(s,ic,ML,y+i*I(0.62),I(0.4),BR_T,"#4A35FF")
        text(s,ML+I(0.56),y+i*I(0.62)+I(0.05),I(8.0),I(0.4),t,font=HEAD,size=11.5,bold=True,color=INK,ls=1.1)
    text(s,ML,I(6.9),I(8.6),I(0.3),"Ideas to shape together, each pressure-tested. Nothing here is a commitment.",size=9.5,color=MUT)
    text(s,SW-MR-I(2.6),I(6.9),I(2.6),I(0.3),"crashfreeindia.org",font=HEAD,size=9.5,bold=True,color=BRAND,align=PP_ALIGN.RIGHT)
    px=I(9.15); pw=I(3.55)
    photo_fit(s,ar("slide_ssi_index.png"),px,I(1.15),pw,I(2.9),caption="School Safety Index, Delhi NCR prototype, live.")
    photo(s,ar("defect_live.png"),px,I(4.55),pw,I(1.6),fill=True,anchor=0.0,
          caption="Road-Infrastructure Defect Repository, live.")
    return s

# ---------------------------------------------------------------- 02 WHY NOW
def whynow():
    s=slide()
    header(s,"01 · Why this conversation, now",
           [("Three facts that were not true ",False),("six months ago",True)],
           lead="Our read of the moment, stated plainly so it can be corrected.",
           sect=(1,"WHY NOW"))
    cols=[("gauge","TRAFFICURE PROVED THE TRAFFIC LAYER","#4A35FF",BR_T,
           [("Pune, Jan 2026:",True),("964 roads watched every 2 minutes, no hardware.",False),
            ("20 → 26.8 km/h",True),("average speed, measured by the police.",False),
            ("Data became a ranked junction list and an engineering brief to the municipal corporation. Fixes got built.",False)]),
          ("zap","RMI NOW CARRIES SAFETY SIGNALS","#B26A00",AM_T,
           [("May 2026, experimental:",True),("hourly vehicle counts, harsh-braking events at intersection scale, and user-reported incidents.",False),
            ("“The safety signal nobody else has… the first proxy for road-design risk at network scale.”",False),
            ("— TraffiCure's own RMI guide",False)]),
          ("school","NOBODY HAS BUILT THE SAFETY PRODUCT","#0A9955",GR_T,
           [("Congestion has a product.",True),("Safety does not, yet.",False),
            ("Turning these signals into safer school zones, junctions and corridors needs safety method, field verification and standing inside government.",False),
            ("That is what we do all day.",True)])]
    cwid=I(3.92); gx=I(0.16)
    for j,(ic,t,col,tint,lines) in enumerate(cols):
        x=ML+j*(cwid+gx); y=I(2.1)
        rect(s,x,y,cwid,I(3.5),fill=PAPER,line=LINE,rounded=True,radius=0.06)
        icon_disc(s,ic,x+I(0.2),y+I(0.2),I(0.44),tint,col)
        text(s,x+I(0.78),y+I(0.26),cwid-I(0.95),I(0.5),t,font=HEAD,size=9.5,bold=True,color=INK,tracking=0.6,ls=1.15)
        tb=text(s,x+I(0.2),y+I(0.92),cwid-I(0.4),I(2.4),
                [(seg+" ",{'bold':b,'color':INK if b else MUT,'size':9.5}) for seg,b in lines],ls=1.3)
    rect(s,ML,I(5.85),CW,I(0.78),fill=BR_T,rounded=True,radius=0.16)
    text(s,ML+I(0.25),I(5.95),CW-I(0.5),I(0.6),
         [("The shape we see: ",{'bold':True,'color':BRAND}),
          ("TraffiCure proves time saved. Together, the same feed can prove lives saved — two products, one pipeline, and a wider government door for both.",{'color':BRAND})],size=11.5,ls=1.3)
    ftc(s,2)
    return s

# ---------------------------------------------------------------- 03 AGENDA
def agenda():
    s=slide()
    text(s,ML,I(0.30),I(8),I(0.24),"HOW THIS DECK IS ORGANISED",font=HEAD,size=9,bold=True,color=BRAND,tracking=2.2)
    hline(s,ML,I(0.62),I(0.55),color=BRAND,wt=2.0)
    text(s,ML,I(0.76),I(12),I(0.5),[("Mostly ideas, ",{'color':INK}),("each one pressure-tested",{'color':BRAND})],font=HEAD,size=22.5,bold=True,ls=1.1)
    rows=[("school","The idea we lead with","School Safety Index × speed data: prototype built, launch with government planned.","04 – 05"),
          ("layers","Three more to shape together","The near-miss map, proving every fix, and SATARK deployment intelligence.","06 – 08"),
          ("lightbulb","Further out on the shelf","Three smaller directions, each with its own caveat attached.","09"),
          ("handshake","What each side brings","The partnership shape, and the government doors we can open together.","10"),
          ("building-2","About us · SATARK · the plan","Who we are, SATARK in the field, a 90-day plan with a decision gate, and our questions.","11 – 14")]
    y=I(1.72)
    for i,(ic,t,d,pg) in enumerate(rows):
        yy=y+i*I(1.02)
        rect(s,ML,yy,CW,I(0.88),fill=MIST if i%2==0 else PAPER,line=LINE if i%2 else None,rounded=True,radius=0.10)
        num_disc(s,ML+I(0.22),yy+I(0.2),I(0.48),i+1)
        icon_disc(s,ic,ML+I(0.95),yy+I(0.2),I(0.48),BR_T,"#4A35FF")
        text(s,ML+I(1.68),yy+I(0.13),I(6.2),I(0.4),t,font=HEAD,size=13,bold=True,color=INK,ls=1.0)
        text(s,ML+I(1.68),yy+I(0.46),I(8.2),I(0.35),d,size=10,color=MUT,ls=1.15)
        text(s,SW-MR-I(1.75),yy+I(0.24),I(1.5),I(0.4),pg,font=HEAD,size=12,bold=True,color=BRAND,align=PP_ALIGN.RIGHT)
    ftc(s,3)
    return s

# ---------------------------------------------------------------- 04 IDEA 1a: SSI BUILT
def ssi_built():
    s=slide()
    header(s,"02 · Idea 01, the one we lead with",
           [("School Safety Index: ",False),("the prototype is already built",True)],
           lead="A public 0–100 grade for the road outside every school gate, anchored to IRC:SP:32, India's school-zone standard.",
           sect=(2,"IDEA 01 · SCHOOLS"))
    rows=[("map-pin","Delhi NCR prototype, live","3,814 schools mapped via the Places API; 90% with usable Street View imagery; pilot report cards published with pinned photo evidence.",BR_T,"#4A35FF"),
          ("scan-line","Method before scale","AI desk audit against a 44-check IRC:SP:32 rubric, independently verified before anything is published. It grades the road, never the school.",GR_T,"#0A9955"),
          ("landmark","Launch with government, planned","We are preparing the public launch with a government partner: our four District Road Safety Committee seats carry the worst-graded zones into the statutory forum, and student teams field-verify them.",AM_T,"#B26A00"),
          ("rocket","The next city is ready","Jaipur: 1,572 schools mapped, 361 GPS-located blackspot crashes from police records, corridors already defined on our Google Cloud project.",TE_T,"#0E8A8A")]
    y=I(2.25)
    for i,(ic,t,d,tint,col) in enumerate(rows):
        irow(s,ML,y+i*I(1.02),I(7.55),ic,t,d,tint,col,tsize=11,dsize=9)
    text(s,ML,I(6.55),I(7.6),I(0.3),
         [("What it cannot see today: ",{'bold':True,'color':INK}),
          ("the speed of traffic past the gate. A photograph never shows it.",{'color':MUT})],size=9.5,ls=1.2)
    px=I(8.55); pw=I(4.15)
    photo_fit(s,ar("slide_ssi_index.png"),px,I(2.2),pw,I(2.3),caption="The index, worst grades first.")
    photo_fit(s,ar("slide_ssi_school.png"),px,I(4.75),pw,I(1.95),caption="A school's report card, evidence pinned.")
    ftc(s,4)
    return s

# ---------------------------------------------------------------- 05 IDEA 1b: SSI x SPEED
def ssi_speed():
    s=slide()
    header(s,"02 · Idea 01, working together",
           [("Add the one input a photograph cannot give: ",False),("speed past the gate",True)],
           lead="Imagery shows the missing crossing. TraffiCure's feed shows the danger moving through it, hour by hour, school-run by school-run.",
           sect=(2,"IDEA 01 · SCHOOLS"))
    rows=[("gauge","Speed-exposure per school zone","Share of school-run hours with traffic above safe speed at the gate, from RMI segment speeds; harsh-braking points near gates as the risk flag.",BR_T,"#4A35FF"),
          ("git-compare","One ranking, by measured danger","Imagery grade × measured speed × vehicle counts as exposure: every school ranked by danger, not by anecdote — and re-measured after a district builds the fix.",GR_T,"#0A9955"),
          ("handshake","How we would split it","We hold method, verification, the public index and the DRSC seats. TraffiCure holds the feed and corridor tooling. The launch names both.",TE_T,"#0E8A8A")]
    y=I(2.3)
    for i,(ic,t,d,tint,col) in enumerate(rows):
        irow(s,ML,y+i*I(1.06),I(7.5),ic,t,d,tint,col,tsize=11,dsize=9)
    rect(s,ML,I(5.75),I(7.45),I(0.95),fill=BR_T,rounded=True,radius=0.12)
    text(s,ML+I(0.22),I(5.86),I(7.0),I(0.75),
         [("A first for India: ",{'bold':True,'color':BRAND}),
          ("no one has published measured speed outside school gates at city scale. The index gives that number a public home and a government audience.",{'color':BRAND})],size=10.5,ls=1.25)
    ptest(s,I(8.5),I(2.2),I(4.2),I(4.5),
          ["RMI segments resolve the short streets outside gates, and school-run-hour slices are supported; terms allow a nonprofit public index to cite the data with attribution.",
           "If segments are too coarse for 30-metre gate zones, or the experimental feeds are thin in Delhi and Jaipur, the speed layer stays anecdote. And a public index must survive a school district challenging its grade.",
           "One working session on 10 gate corridors we have already defined. Pull real segments, check granularity and density against our field audits of the same gates. A week, not a quarter."])
    ftc(s,5)
    return s

# ---------------------------------------------------------------- 06 IDEA 2: NEAR-MISS MAP
def nearmiss():
    s=slide()
    header(s,"02 · Idea 02",
           [("Harsh braking: ",False),("the near-miss map, validated against real crashes",True)],
           lead="Crash data arrives after the funeral. Braking data arrives every day. The question worth answering: does one predict the other on Indian roads?",
           sect=(2,"IDEA 02 · NEAR-MISSES"))
    rows=[("map-pin","What we hold","361 GPS-located crashes at Jaipur's notified blackspots (police e-DAR, via our DRSC seat), 96 mapped clusters, and a live national defect repository with 500+ locations.",BR_T,"#4A35FF"),
          ("zap","What the feed adds","Anonymised harsh-braking points at intersection scale — TraffiCure calls it the safety signal nobody else has. Overlaid on known crash sites, it can be tested, not assumed.",AM_T,"#B26A00"),
          ("file-search","What it becomes","An India-first validation study: where braking predicts crashes, authorities get a risk-ranked corridor list before the next fatality — and Google gets published, citable evidence from India.",GR_T,"#0A9955")]
    y=I(2.35)
    for i,(ic,t,d,tint,col) in enumerate(rows):
        irow(s,ML,y+i*I(1.1),I(7.5),ic,t,d,tint,col,tsize=11,dsize=9)
    text(s,ML,I(5.85),I(7.5),I(0.8),
         [("Worst case, we learn the feed's limits. Best case, ",{'color':MUT}),
          ("India's road-risk maps stop waiting for people to die first.",{'bold':True,'color':INK})],size=10.5,ls=1.3)
    ptest(s,I(8.5),I(2.2),I(4.2),I(4.5),
          ["Braking events are dense enough per junction to be statistically usable, and the anonymisation thresholds still leave signal on two-wheeler-heavy Indian streets.",
           "Braking correlates with congestion, not danger — chaos that never crashes. If the correlation with e-DAR crash sites is weak, this is a research note, not a product.",
           "A 4-week overlay on Jaipur's 96 clusters: braking density at crash sites vs matched control junctions. Clear threshold agreed up front — publish either way."])
    ftc(s,6)
    return s

# ---------------------------------------------------------------- 07 IDEA 3: PROVE THE FIX
def provefix():
    s=slide()
    header(s,"02 · Idea 03",
           [("Prove every fix: ",False),("before-and-after for road safety spending",True)],
           lead="Pune used the feed to prove a decongestion plan worked. The same instrument can prove safety works — nobody has pointed it there.",
           sect=(2,"IDEA 03 · PROOF"))
    rows=[("hard-hat","The pipeline exists","15 Rakshak sites are in implementation now and 75+ implementations are targeted this year — speed humps, crossings, junction geometry, each with a known build date.",BR_T,"#4A35FF"),
          ("git-compare","The measurement is one subscription away","Corridor speeds and counts before the fix, after the fix, against control corridors. Today India evaluates road-safety works by inspection photos; this gives them a measured outcome.",GR_T,"#0A9955"),
          ("landmark","Where it lands","Every before-after goes to the District Road Safety Committee in the statutory record, onto our public dashboard — and into TraffiCure's case-study shelf as the first safety deployments.",TE_T,"#0E8A8A")]
    y=I(2.35)
    for i,(ic,t,d,tint,col) in enumerate(rows):
        irow(s,ML,y+i*I(1.1),I(7.5),ic,t,d,tint,col,tsize=11,dsize=9)
    text(s,ML,I(5.9),I(7.5),I(0.7),
         [("This is the cheapest idea in the deck, ",{'bold':True,'color':INK}),
          ("and the fastest to show working: the fixes are being built either way.",{'color':MUT})],size=10.5,ls=1.3)
    ptest(s,I(8.5),I(2.2),I(4.2),I(4.5),
          ["Corridors can be subscribed before construction starts (baseline needs weeks of history), and speed changes at fixed sites clear the noise of seasons, weather and traffic growth.",
           "Attribution: speeds move for many reasons. Without control corridors and honest stats, a sceptical engineer dismisses it — and one over-claimed result poisons the method.",
           "Pick 2 sites entering construction this quarter, subscribe now, pre-register the analysis with control corridors. First honest result inside 90 days of the build."])
    ftc(s,7)
    return s

# ---------------------------------------------------------------- 08 IDEA 4: SATARK x FEED
def satark_feed():
    s=slide()
    header(s,"02 · Idea 04, exploratory",
           [("SATARK × the feed: ",False),("enforcement standing where the risk is",True)],
           lead="An early idea, listed honestly. It builds only on assets we already run with police partners, and carries no commitments.",
           sect=(2,"IDEA 04 · ENFORCEMENT"))
    rows=[("cpu","What already runs","SATARK screens vehicles for police in Jaipur, Bengaluru and Gurugram: 3,00,000+ screened, 350+ intercepted, challan billboards live. MoUs govern all of it.",BR_T,"#4A35FF"),
          ("gauge","What the feed could add","A weekly view of where and when speeding concentrates near people — so patrols, camera vans, billboards and SATARK sites stand where the risk is, corridor-hour by corridor-hour.",AM_T,"#B26A00"),
          ("shield-check","How we keep it honest","Success measured as speeds before and after a deployment, never as challan counts. RMI stays aggregate — no plates, no individual vehicles. Any pilot starts small, inside an existing MoU.",GR_T,"#0A9955")]
    y=I(2.35)
    for i,(ic,t,d,tint,col) in enumerate(rows):
        irow(s,ML,y+i*I(1.1),I(7.5),ic,t,d,tint,col,tsize=11,dsize=9)
    text(s,ML,I(5.9),I(7.5),I(0.6),
         [("Everything named here is ours or the police's. ",{'bold':True,'color':INK}),("No third party is required for a first version.",{'color':MUT})],size=10.5,ls=1.3)
    ptest(s,I(8.5),I(2.2),I(4.2),I(4.5),
          ["Police continue to pull this — the MoUs and three years of working trust are the asset — and the speeding-concentration view is granular enough to move a patrol roster.",
           "Optics: 'Google data guiding enforcement' can be misread as surveillance. The aggregate-only line must hold absolutely, and the police must own every decision.",
           "One city, four weeks: hand Jaipur Traffic Police a weekly speeding-concentration map for existing patrol corridors. If rosters change, it works. If not, we shelve it."])
    ftc(s,8)
    return s

# ---------------------------------------------------------------- 09 SHELF
def shelf():
    s=slide()
    header(s,"02 · Further out",
           [("Three more on the shelf, ",False),("each with its caveat attached",True)],
           lead="Not proposals yet. Listed so the conversation can pick anything up, or put it down.",
           sect=(2,"THE SHELF"))
    cards=[("bot","Incident feed × the under-reporting gap",
            "Maps users report crashes police never record. Compare the incident feed against FIR and e-DAR records in our districts to size India's crash under-reporting, a number policy people cite for a decade.",
            "Caveat: incident reports are noisy and unverified; needs careful matching before any claim."),
           ("route","Safe-route research for two-wheelers",
            "Join our audit and crash data with speed layers to ask whether a marginally slower route is materially safer, starting with the delivery riders we already study.",
            "Caveat: research first; a consumer nudge is Google's call, never ours."),
           ("school","School-run advisories",
            "Where the index and speed data agree a gate is dangerous at 8 am, that knowledge could reach parents and schools through channels Google already owns.",
            "Caveat: only after the index and speed layer are validated; advisories without method invite panic.")]
    cwid=I(3.92); gx=I(0.16)
    for j,(ic,t,d,cv) in enumerate(cards):
        x=ML+j*(cwid+gx); y=I(2.2)
        rect(s,x,y,cwid,I(4.0),fill=PAPER,line=LINE,rounded=True,radius=0.06)
        icon_disc(s,ic,x+I(0.2),y+I(0.22),I(0.44),BR_T,"#4A35FF")
        text(s,x+I(0.78),y+I(0.26),cwid-I(0.95),I(0.65),t,font=HEAD,size=11,bold=True,color=INK,ls=1.12)
        text(s,x+I(0.2),y+I(1.05),cwid-I(0.4),I(1.85),d,size=9.5,color=MUT,ls=1.28)
        rect(s,x+I(0.14),y+I(3.05),cwid-I(0.28),I(0.8),fill=AM_T,rounded=True,radius=0.12)
        text(s,x+I(0.26),y+I(3.13),cwid-I(0.52),I(0.66),cv,size=8,color=INK,ls=1.18)
    ftc(s,9)
    return s

# ---------------------------------------------------------------- 10 PARTNERSHIP SHAPE
def shape():
    s=slide()
    header(s,"03 · What each side brings",
           [("A partnerships channel into government, ",False),("and a clean split",True)],
           lead="Cities buy operations software. Safety products stay public and free. The pilots that prove both run through doors we already hold open.",
           sect=(3,"THE SHAPE"))
    x1=ML; w=I(5.9)
    rect(s,x1,I(2.15),w,I(3.3),fill=BR_T,rounded=True,radius=0.06)
    text(s,x1+I(0.22),I(2.32),w-I(0.44),I(0.28),"WHAT WE BRING",font=HEAD,size=9.5,bold=True,color=BRAND,tracking=1.6)
    rows1=[("landmark","Standing inside government: seats on 4 District Road Safety Committees, MoUs with Jaipur and Gurugram police, 25+ written authority approvals across six states, EoIs before NRSB working groups."),
           ("search-check","Safety method that survives review: IRC and MoRTH-anchored audits, adversarial verification, student field teams in 18 cities."),
           ("globe","Public products and credibility: live index and repository, national press coverage, MS Dhoni as Goodwill Ambassador.")]
    yy=I(2.7)
    for ic,t in rows1:
        icon(s,ic,x1+I(0.22),yy+I(0.02),I(0.26),"#4A35FF")
        text(s,x1+I(0.6),yy,w-I(0.82),I(0.85),t,size=9.5,color=INK,ls=1.25)
        yy=yy+I(0.92)
    x2=ML+w+I(0.28)
    rect(s,x2,I(2.15),w,I(3.3),fill=GR_T,rounded=True,radius=0.06)
    text(s,x2+I(0.22),I(2.32),w-I(0.44),I(0.28),"WHAT TRAFFICURE + GOOGLE BRING",font=HEAD,size=9.5,bold=True,color=GREEN,tracking=1.6)
    rows2=[("gauge","The data layer nothing else matches: every road, every 2 minutes, plus the May 2026 safety feeds — no hardware, weeks to deploy."),
           ("cpu","A working product and playbook: TraffiCure live with Pune City Traffic Police, corridors measured in four cities, activation in 4–8 weeks."),
           ("handshake","Reach and legitimacy: a Premier Partner network and the Google name, which opens rooms and closes doubts.")]
    yy=I(2.7)
    for ic,t in rows2:
        icon(s,ic,x2+I(0.22),yy+I(0.02),I(0.26),"#0A9955")
        text(s,x2+I(0.6),yy,w-I(0.82),I(0.85),t,size=9.5,color=INK,ls=1.25)
        yy=yy+I(0.92)
    rect(s,ML,I(5.7),CW,I(0.95),fill=MIST,rounded=True,radius=0.12)
    text(s,ML+I(0.25),I(5.82),CW-I(0.5),I(0.75),
         [("The clean split: ",{'bold':True,'color':INK}),
          ("no money needs to move between us. Cities subscribe to TraffiCure for operations; our safety products stay public, free and attributed; joint district pilots become the reference deployments that open the next city for everyone.",{'color':MUT})],size=10.5,ls=1.3)
    ftc(s,10)
    return s

# ---------------------------------------------------------------- 11 ABOUT (portfolio)
def about_tiles():
    s=slide()
    header(s,"04 · About us",[("Six lines of work, ",False),("one connected system",True)],
           lead="A nonprofit registered as the Vision Zero Trust · Co-founded by Cars24 and the Indian Road Safety Council · 18 cities · everything below is live.",
           sect=(4,"ABOUT US"))
    tiles=[("Project Rakshak","Student-led road audits, fixed with authorities","31 sites · 15 in implementation",a5("implementation_work.jpg"),0.30),
           ("SATARK enforcement","AI support that helps police find high-risk vehicles","3,00,000+ vehicles screened",a5("satark_billboard.jpg"),0.42),
           ("Victim support","Aasha chatbot, calculator, hospital helpdesks","100+ victims helped in first 4 days",a5("ph_helpdesk_desk.jpg"),0.25),
           ("Open public data","Four self-updating data platforms, free to use","36 states & 50 cities covered",a5("sctracker.png"),0.05),
           ("School Safety Index","Public grading of school-zone road safety","3,814 Delhi NCR schools mapped",ar("slide_ssi_school.png"),0.0),
           ("Youth & policy","NextMile ideathon and national dialogue series","2,000+ participants engaged",a5("ph_ideathon_winners.jpg"),0.25)]
    tw=I(3.95); th=I(2.32); gx=I(0.12); gy=I(0.16)
    for i,(t,d,st,img,anc) in enumerate(tiles):
        r,c=divmod(i,3)
        x=ML+c*(tw+gx); y=I(2.12)+r*(th+gy)
        rect(s,x,y,tw,th,fill=PAPER,line=LINE,rounded=True,radius=0.06)
        photo(s,img,x,y,tw,I(1.28),fill=True,border=False,anchor=anc)
        rect(s,x,y,tw,I(1.28),fill=None,line=LINE,lw=1.0)
        text(s,x+I(0.14),y+I(1.36),tw-I(0.28),I(0.3),t,font=HEAD,size=11.5,bold=True,color=INK,ls=1.0)
        text(s,x+I(0.14),y+I(1.62),tw-I(0.28),I(0.35),d,size=9,color=MUT,ls=1.12)
        text(s,x+I(0.14),y+I(2.02),tw-I(0.28),I(0.26),st,font=HEAD,size=9,bold=True,color=BRAND,ls=1.0)
    ftc(s,11)
    return s

# ---------------------------------------------------------------- 12 SATARK (reuse + patch rail)
def satark_slide():
    s=B.satark(12,TOTAL,None)
    patch_rail(s,"02 · What we have built, enforcement technology","04 · About us, SATARK in the field")
    patch_rail(s,"02 · WHAT WE HAVE BUILT, ENFORCEMENT TECHNOLOGY","04 · ABOUT US, SATARK IN THE FIELD")
    patch_rail(s,"WHAT WE HAVE BUILT","ABOUT US")
    return s

# ---------------------------------------------------------------- 13 90-DAY PLAN
def plan90():
    s=slide()
    header(s,"05 · The plan",
           [("Ninety days, ",False),("with a decision gate at the end",True)],
           lead="Small, funded on our side, and honest: each idea carries its own kill-test. What survives, scales.",
           sect=(5,"THE PLAN"))
    cards=[("1","WEEKS 1 – 2","The technical session","Lepton + Maps + CFI at one table, on real data: 10 school-gate corridors and Jaipur's 96 crash clusters, already defined on our configured cloud project. Granularity and density answered in the room."),
          ("2","WEEKS 3 – 8","Two pressure tests run","The SSI speed layer scored on ~100 school zones in one city, and the harsh-braking overlay against Jaipur's crash sites — thresholds agreed up front, results shared either way."),
          ("3","WEEKS 9 – 13","Into the statutory record","Surviving results go to the District Road Safety Committees and into the public index; before-after subscriptions start on 2 Rakshak sites entering construction; the government launch of the index is scoped with a named partner.")]
    cwid=I(3.92); gx=I(0.16)
    for j,(n,wk,t,d) in enumerate(cards):
        x=ML+j*(cwid+gx); y=I(2.2)
        rect(s,x,y,cwid,I(3.35),fill=PAPER,line=LINE,rounded=True,radius=0.06)
        num_disc(s,x+I(0.2),y+I(0.2),I(0.44),int(n))
        text(s,x+I(0.78),y+I(0.24),cwid-I(0.95),I(0.26),wk,font=HEAD,size=9,bold=True,color=BRAND,tracking=1.4)
        text(s,x+I(0.78),y+I(0.5),cwid-I(0.95),I(0.35),t,font=HEAD,size=12,bold=True,color=INK,ls=1.05)
        text(s,x+I(0.2),y+I(1.05),cwid-I(0.4),I(2.2),d,size=9.5,color=MUT,ls=1.3)
    rect(s,ML,I(5.8),CW,I(0.85),fill=AM_T,rounded=True,radius=0.12)
    text(s,ML+I(0.25),I(5.9),CW-I(0.5),I(0.65),
         [("Day 90, the gate: ",{'bold':True,'color':INK}),
          ("each idea is scaled, reshaped, or shelved — in writing. Our side runs at no cost to Google. The one ask: feed access for public-good use, with attribution on everything we publish.",{'color':INK})],size=10.5,ls=1.3)
    ftc(s,13)
    return s

# ---------------------------------------------------------------- 14 QUESTIONS
def questions():
    s=slide()
    header(s,"05 · For tomorrow",
           [("Four questions ",False),("we bring to the table",True)],
           lead="Held loosely. The ranking of ideas is ours to lose; the starting point is yours to shape.",
           sect=(5,"THE QUESTIONS"))
    cards=[("database","Access and terms","Can a nonprofit safety layer sit on the feed — harsh braking and counts included — with terms that allow a public, attributed index? What does that take, and through whom: Google, Lepton, or both?"),
           ("gauge","The granularity truth","On the short urban streets outside school gates, what segment resolution and experimental-feed density should we honestly expect in Delhi and Jaipur?"),
           ("landmark","Which doors matter to you","Which cities and stories help TraffiCure and RMI most in the next two quarters? Our committee seats, police MoUs and NRSB engagement work best when pointed somewhere specific."),
           ("rocket","The working session","Can we fix a date for the Lepton + Maps + CFI technical session? Our corridors, zones and cloud project are ready this week.")]
    cwid=I(5.9); ch=I(1.95); gx=I(0.28); gy=I(0.22)
    for k,(ic,t,d) in enumerate(cards):
        r,c=divmod(k,2)
        x=ML+c*(cwid+gx); y=I(2.2)+r*(ch+gy)
        rect(s,x,y,cwid,ch,fill=PAPER if k%3 else MIST,line=LINE,rounded=True,radius=0.06)
        icon_disc(s,ic,x+I(0.22),y+I(0.22),I(0.46),BR_T,"#4A35FF")
        text(s,x+I(0.88),y+I(0.3),cwid-I(1.1),I(0.35),t,font=HEAD,size=12.5,bold=True,color=INK,ls=1.0)
        text(s,x+I(0.24),y+I(0.82),cwid-I(0.48),I(1.05),d,size=9.5,color=MUT,ls=1.28)
    ftc(s,14)
    return s

# ---------------------------------------------------------------- 15 CLOSE
def close():
    s=slide(bg=BRAND)
    s.shapes.add_picture(LOGO_WHITE,ML,I(0.62),width=I(2.0))
    text(s,ML,I(2.0),I(12.1),I(1.5),"Traffic data already runs the daily commute.\nTogether it can start preventing the crash.",font=HEAD,size=30,bold=True,color=PAPER,ls=1.2)
    text(s,ML,I(3.8),I(11.5),I(0.6),"Every idea in this deck is offered with its own kill-test attached. We would rather shelve three honestly than overclaim one.",size=12.5,color=LAVW,ls=1.4)
    hline(s,ML,I(4.6),I(5.2),color=RGBColor(0x6A,0x55,0xFF),wt=1.0)
    cols=[("Akhtar Hussain","Mission Lead","akhtar@crashfreeindia.org"),
          ("Deepanshu Gupta","Trustee","deepanshu@crashfreeindia.org"),
          ("Gajendra Jangid","Trustee","gajendra@crashfreeindia.org")]
    for j,(n,r,e) in enumerate(cols):
        x=ML+j*I(4.0)
        text(s,x,I(4.9),I(3.8),I(0.3),n,font=HEAD,size=13,bold=True,color=PAPER)
        text(s,x,I(5.24),I(3.8),I(0.28),f"{r} · {e}",size=9.5,color=LAVW)
    text(s,ML,I(7.05),I(12),I(0.3),"Crashfree India · A Cars24 commitment · Operated by Vision Zero Trust, New Delhi · crashfreeindia.org",size=9.5,color=LAVW)
    return s

if __name__=="__main__":
    new_deck()
    cover()          # 01
    whynow()         # 02
    agenda()         # 03
    ssi_built()      # 04
    ssi_speed()      # 05
    nearmiss()       # 06
    provefix()       # 07
    satark_feed()    # 08
    shelf()          # 09
    shape()          # 10
    about_tiles()    # 11
    satark_slide()   # 12
    plan90()         # 13
    questions()      # 14
    close()          # 15
    save(os.path.join(BUILD,"CFI_Trafficure_v1.pptx"))
