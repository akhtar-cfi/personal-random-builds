# CFI x TraffiCure / Google — discussion deck v3. 10 slides.
# One structured slide per idea (why / why CFI x TraffiCure / how together),
# SSI next steps specific, Rakshak with proof visuals, DRSC knowledge pack,
# final next-steps slide with the one-district demo-access ask.
import os
import build_nrsb3 as B
import nrsb2_style as ST
from nrsb2_style import *
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

I=Inches
RMI=os.path.join(SP,"assets/rmi")
def ar(n): return os.path.join(RMI,n)

TOTAL=10
AMBC=RGBColor(0xB2,0x6A,0x00)

def ftc(s,idx):
    hline(s,ML,I(7.06),CW)
    text(s,ML,I(7.14),I(9),I(0.25),"Crashfree India  ·  Discussion deck · Crashfree India × TraffiCure, Google Maps Platform",size=8.5,color=MUT,ls=1.0)
    text(s,SW-MR-I(1.2),I(7.14),I(1.2),I(0.25),f"{idx:02d} / {TOTAL:02d}",size=8.5,color=MUT,align=PP_ALIGN.RIGHT,ls=1.0)

B.footer=lambda s,wg,idx,total: ftc(s,idx)
def _sect_tab(s,cur,name):
    w=I(3.6)
    text(s,SW-MR-w,I(0.30),w,I(0.24),name,font=HEAD,size=8.5,bold=True,color=MUT,tracking=1.6,align=PP_ALIGN.RIGHT)
ST.sect_tab=_sect_tab

def patch_rail(s,old,new):
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text and old in r.text:
                        r.text=r.text.replace(old,new)

def stat_tile(s,x,y,w,h,val,lab,tint=MIST,col=BRAND,vs=20):
    rect(s,x,y,w,h,fill=tint,rounded=True,radius=0.08)
    text(s,x+I(0.16),y+I(0.1),w-I(0.32),I(0.5),val,font=HEAD,size=vs,bold=True,color=col)
    text(s,x+I(0.16),y+I(0.58),w-I(0.32),h-I(0.66),lab,size=8,color=MUT,ls=1.15)

def why_us_band(s,y,cfi,tc,w=I(7.5)):
    """two side-by-side chips: why CFI / why TraffiCure"""
    half=(w-I(0.14))/2
    rect(s,ML,I(y),half,I(1.02),fill=BR_T,rounded=True,radius=0.10)
    text(s,ML+I(0.16),I(y+0.08),half-I(0.32),I(0.22),"WHY CFI",font=HEAD,size=8,bold=True,color=BRAND,tracking=1.4)
    text(s,ML+I(0.16),I(y+0.32),half-I(0.32),I(0.64),cfi,size=8.5,color=INK,ls=1.2)
    x2=ML+half+I(0.14)
    rect(s,x2,I(y),half,I(1.02),fill=GR_T,rounded=True,radius=0.10)
    text(s,x2+I(0.16),I(y+0.08),half-I(0.32),I(0.22),"WHY TRAFFICURE",font=HEAD,size=8,bold=True,color=GREEN,tracking=1.4)
    text(s,x2+I(0.16),I(y+0.32),half-I(0.32),I(0.64),tc,size=8.5,color=INK,ls=1.2)

# ---------------------------------------------------------------- 01 COVER
def cover():
    s=slide()
    s.shapes.add_picture(LOGO_BLUE,ML,I(0.5),width=I(1.85))
    text(s,SW-MR-I(3.5),I(0.55),I(3.5),I(0.3),"AUGUST 2026  ·  NEW DELHI",font=HEAD,size=9,bold=True,color=MUT,tracking=1.8,align=PP_ALIGN.RIGHT)
    text(s,ML,I(1.72),I(8.4),I(0.3),"FOR DISCUSSION  ·  CRASHFREE INDIA × TRAFFICURE · GOOGLE MAPS PLATFORM",font=HEAD,size=10,bold=True,color=BRAND,tracking=2.0)
    text(s,ML,I(2.14),I(8.3),I(0.85),"The safety layer",font=HEAD,size=46,bold=True,color=INK,ls=1.0)
    text(s,ML,I(3.1),I(8.3),I(0.85),"Traffic data already manages congestion.\nTogether it can start preventing crashes.",font=HEAD,size=16.5,bold=True,color=BRAND,ls=1.25)
    rows=[("school","School Safety Index × speed — prototype built, launching with government"),
          ("map-pin","Live crash hotspots — found in the data, fixed through Project Rakshak"),
          ("landmark","The district knowledge pack — joint thought leadership, monthly, where decisions are made")]
    y=I(4.65)
    for i,(ic,t) in enumerate(rows):
        icon_disc(s,ic,ML,y+i*I(0.62),I(0.4),BR_T,"#4A35FF")
        text(s,ML+I(0.56),y+i*I(0.62)+I(0.05),I(8.0),I(0.4),t,font=HEAD,size=11.5,bold=True,color=INK,ls=1.1)
    text(s,ML,I(6.9),I(8.6),I(0.3),"Three ideas, one ask: one district's data, 90 days, MVPs built together.",size=9.5,color=MUT)
    text(s,SW-MR-I(2.6),I(6.9),I(2.6),I(0.3),"crashfreeindia.org",font=HEAD,size=9.5,bold=True,color=BRAND,align=PP_ALIGN.RIGHT)
    px=I(9.15); pw=I(3.55)
    photo_fit(s,ar("slide_ssi_index.png"),px,I(1.15),pw,I(2.9),caption="School Safety Index, Delhi NCR prototype, live.")
    photo(s,ar("defect_live.png"),px,I(4.55),pw,I(1.6),fill=True,anchor=0.0,
          caption="Road-Infrastructure Defect Repository, live.")
    return s

# ---------------------------------------------------------------- 02 TRAFFICURE, UNDERSTOOD
def tc_who():
    s=slide()
    header(s,"01 · Understanding TraffiCure",
           [("What TraffiCure is, ",False),("and what Pune proved",True)],
           lead="Software only. No cameras, no sensors. Google Maps probe data, turned into a working tool for the people who run the roads.",
           sect=(1,"TRAFFICURE"))
    rows=[("building-2","Built by Lepton Software","Google Maps Platform Premier Partner, working with Google since 2010. TraffiCure launched 2025.",BR_T,"#4A35FF"),
          ("gauge","Runs on Roads Management Insights","Google's public-sector traffic feed, launched Aug 2025. Every road, refreshed every 2 minutes.",TE_T,"#0E8A8A"),
          ("landmark","Live with Pune City Traffic Police","Deployed January 2026, activated in weeks. Corridors also measured in Delhi, Hyderabad, Ahmedabad.",GR_T,"#0A9955")]
    y=I(2.25)
    for i,(ic,t,d,tint,col) in enumerate(rows):
        irow(s,ML,y+i*I(0.92),I(7.5),ic,t,d,tint,col,tsize=11,dsize=9.5)
    text(s,ML,I(5.15),I(7.5),I(0.3),"THE PUNE RESULT, POLICE-MEASURED",font=HEAD,size=9,bold=True,color=MUT,tracking=1.8)
    stat_tile(s,ML,I(5.5),I(2.4),I(1.15),"964","roads watched live, every 2 minutes")
    stat_tile(s,ML+I(2.55),I(5.5),I(2.4),I(1.15),"20 → 26.8","km/h average speed, in a year")
    stat_tile(s,ML+I(5.1),I(5.5),I(2.4),I(1.15),"22","worst junctions ranked and briefed")
    px=I(8.55); pw=I(4.15)
    rect(s,px,I(2.25),pw,I(4.35),fill=PAPER,line=LINE,rounded=True,radius=0.06)
    text(s,px+I(0.22),I(2.45),pw-I(0.44),I(0.3),"WHY PUNE MATTERS TO US",font=HEAD,size=9,bold=True,color=BRAND,tracking=1.6)
    text(s,px+I(0.22),I(2.85),pw-I(0.44),I(3.5),
         "The police did not stop at watching traffic.\n\nThey ranked their worst junctions, diagnosed why each fails, and handed the municipal corporation an engineering brief. Fixes got built.\n\nData → ranked list → brief to authority → built fix.\n\nThat is exactly how Project Rakshak works. Same operating theory, different data.",size=10,color=INK,ls=1.35)
    ftc(s,2)
    return s

# ---------------------------------------------------------------- 03 THE DATA, SAFETY LENS
def tc_data():
    s=slide()
    header(s,"01 · Understanding the data",
           [("Five signals in the feed, ",False),("read through a safety lens",True)],
           lead="The product was built to manage traffic. Every signal in it has a second life in safety.",
           sect=(1,"THE DATA"))
    rows=[("gauge","Live speeds","Every road, every 2 minutes","How fast traffic really moves past a school gate at 8 am.",BR_T,"#4A35FF",""),
          ("database","Speed history","Hourly, by road, in BigQuery","Baselines: what is normal, what changed after a fix was built.",TE_T,"#0E8A8A",""),
          ("users","Vehicle counts","Per road, per hour","Exposure: how many vehicles pass the danger. Priority = risk × exposure.",GR_T,"#0A9955","NEW · MAY 2026"),
          ("zap","Harsh braking","Anonymised points, junction scale","Near-misses, visible before the crash. Their guide calls it the safety signal nobody else has.",AM_T,"#B26A00","NEW · MAY 2026"),
          ("bot","User reports","Crashes and hazards, from Maps users","Incidents the official record may never see.",BR_T,"#4A35FF","NEW · MAY 2026")]
    y=I(2.2)
    for i,(ic,t,what,use,tint,col,tag) in enumerate(rows):
        yy=y+i*I(0.88)
        rect(s,ML,yy,CW,I(0.76),fill=PAPER if i%2 else MIST,rounded=True,radius=0.10)
        icon_disc(s,ic,ML+I(0.18),yy+I(0.15),I(0.46),tint,col)
        text(s,ML+I(0.82),yy+I(0.12),I(1.9),I(0.5),t,font=HEAD,size=11.5,bold=True,color=INK,ls=1.05)
        text(s,ML+I(0.82),yy+I(0.45),I(2.6),I(0.26),what,size=8.5,color=MUT,ls=1.1)
        text(s,ML+I(3.7),yy+I(0.14),I(6.2),I(0.55),use,size=10,color=INK,ls=1.2)
        if tag:
            chip(s,SW-MR-I(1.75),yy+I(0.21),I(1.6),I(0.34),tag,fill=AM_T,color=AMBC,size=7.5)
    ftc(s,3)
    return s

# ---------------------------------------------------------------- 04 IDEA 1 · SSI (one slide)
def idea_ssi():
    s=slide()
    header(s,"02 · Idea 01 · School Safety Index × speed",
           [("Grade every school gate by ",False),("measured danger",True)],
           lead="A public 0–100 grade for the road outside every school. Built, live for 3,814 Delhi NCR schools, launching with government. Speed past the gate is the one input it still lacks.",
           sect=(2,"IDEA 01 · SSI"))
    # WHY strip
    text(s,ML,I(2.08),I(7.5),I(0.24),"WHY",font=HEAD,size=8.5,bold=True,color=MUT,tracking=1.8)
    stat_tile(s,ML,I(2.34),I(2.42),I(1.08),"~10,000","minors die on Indian roads every year (MoRTH)",vs=17)
    stat_tile(s,ML+I(2.54),I(2.34),I(2.42),I(1.08),"4.6%","of crashes happen in school & college zones (e-DAR)",vs=17)
    stat_tile(s,ML+I(5.08),I(2.34),I(2.42),I(1.08),"0","Indian cities rank school-zone danger today",vs=17)
    # WHY US band
    text(s,ML,I(3.62),I(7.5),I(0.24),"WHY CFI × TRAFFICURE",font=HEAD,size=8.5,bold=True,color=MUT,tracking=1.8)
    why_us_band(s,3.88,
        "The index exists: IRC:SP:32 method, verification, public report cards — and DRSC seats to launch it with government.",
        "The one input a photo cannot give: speed past the gate at school-run hours, plus counts as exposure. No one else has it.")
    # NEXT STEPS
    text(s,ML,I(5.12),I(7.5),I(0.24),"NEXT STEPS, SPECIFICALLY",font=HEAD,size=8.5,bold=True,color=MUT,tracking=1.8)
    steps=[("1","Week 1","Demo access + technical session on 10 gate corridors, already defined on our cloud project."),
           ("2","Weeks 2–6","Speed-exposure scoring on the ~100-zone government-launch batch; thresholds agreed up front."),
           ("3","Weeks 6–10","Speed layer into the public index; student teams field-verify the worst 20 zones."),
           ("4","This quarter","Government launch of the index — speed layer in, both names on it.")]
    for k,(n,w,d) in enumerate(steps):
        yy=I(5.38+k*0.42)
        num_disc(s,ML,yy,I(0.3),int(n))
        text(s,ML+I(0.42),yy+I(0.0),I(1.25),I(0.3),w,font=HEAD,size=9,bold=True,color=BRAND)
        text(s,ML+I(1.75),yy+I(0.0),I(5.8),I(0.4),d,size=9,color=INK,ls=1.12)
    px=I(8.55); pw=I(4.15)
    photo_fit(s,ar("slide_ssi_index.png"),px,I(2.15),pw,I(2.5),caption="Live: the index, worst grades first.")
    photo_fit(s,ar("slide_ssi_school.png"),px,I(4.9),pw,I(1.85),caption="A school's report card, evidence pinned.")
    ftc(s,4)
    return s

# ---------------------------------------------------------------- 05 IDEA 2 · RAKSHAK (one slide)
def idea_rakshak():
    s=slide()
    header(s,"02 · Idea 02 · Live hotspots × Project Rakshak",
           [("Find danger in the data, ",False),("fix it on the ground",True)],
           lead="Official blackspot lists go stale for years. The feed refreshes the danger map monthly — and Rakshak already turns danger maps into audits, joint plans and built fixes.",
           sect=(2,"IDEA 02 · RAKSHAK"))
    # flow strip (one row, 5 chips)
    flow=[("zap","Signals + crashes"),("layers","Live hotspot list"),("search-check","Rakshak audit"),("landmark","Joint plan to authority"),("git-compare","Fix built & measured")]
    cw2=I(1.41); gap=I(0.16)
    for k,(ic,t) in enumerate(flow):
        x=ML+k*(cw2+gap)
        rect(s,x,I(2.12),cw2,I(0.92),fill=BR_T if k in (1,3) else PAPER,line=LINE,rounded=True,radius=0.12)
        icon(s,ic,x+cw2/2-I(0.12),I(2.24),I(0.24),"#4A35FF")
        text(s,x+I(0.05),I(2.52),cw2-I(0.10),I(0.48),t,font=HEAD,size=7.5,bold=True,color=INK,align=PP_ALIGN.CENTER,ls=1.05)
        if k<4:
            text(s,x+cw2-I(0.01),I(2.44),gap+I(0.02),I(0.3),"›",font=HEAD,size=11,bold=True,color=BRAND,align=PP_ALIGN.CENTER)
    # WHAT THE DATA DIAGNOSES
    text(s,ML,I(3.3),I(7.5),I(0.24),"THE DESK DIAGNOSIS — ANSWERS THAT TODAY NEED A PHYSICAL AUDIT",font=HEAD,size=8.5,bold=True,color=MUT,tracking=1.4)
    diag=[("timer","When","Speeding windows: 11 pm chaos or 8 am school-run?"),
          ("zap","Where","Braking clusters mark the exact conflict point"),
          ("gauge","Why","Stop-go congestion, free-flow speeding, or design failure"),
          ("users","Who","Counts say how many people the danger exposes")]
    for k,(ic,t,d) in enumerate(diag):
        x=ML+(k%2)*I(3.8); yy=I(3.6)+ (k//2)*I(0.62)
        icon(s,ic,x,yy+I(0.02),I(0.24),"#B26A00")
        text(s,x+I(0.36),yy,I(0.62),I(0.28),t,font=HEAD,size=9.5,bold=True,color=INK)
        text(s,x+I(1.02),yy,I(2.75),I(0.55),d,size=8.5,color=MUT,ls=1.1)
    text(s,ML,I(4.95),I(7.5),I(0.35),
         [("Field teams arrive knowing the corridor's story — ",{'color':MUT,'size':9.5}),
          ("the visit confirms it and designs the fix.",{'bold':True,'color':INK,'size':9.5})],ls=1.2)
    # WHY US + HOW
    why_us_band(s,5.42,
        "150+ audits a year, 25+ authority approvals, 4 DRSC seats to table joint plans — and a public tracker for every step.",
        "The only live all-roads feed, plus the Pune playbook: ranked list → engineering brief → built fix.")
    text(s,ML,I(6.6),I(7.5),I(0.35),
         [("First 90 days:  ",{'bold':True,'color':BRAND,'size':9.5}),
          ("overlay braking on Jaipur's 96 crash clusters · baseline 2 sites entering construction · live speed panel on the Rakshak dashboard.",{'color':INK,'size':9.5})],ls=1.2)
    # right column: proof visuals
    px=I(8.55); pw=I(4.15)
    photo(s,a5("implementation_work.jpg"),px,I(2.12),pw,I(2.2),fill=True,anchor=0.3,
          caption="A Rakshak fix under construction — one of 15 in implementation.")
    photo_fit(s,ar("rakshak_approval_letter.jpg"),px,I(4.75),pw,I(1.75),
          caption="An authority approval letter. Every site tracked in public: crashfreeindia.org/rakshak/dashboard")
    ftc(s,5)
    return s

# ---------------------------------------------------------------- 06 IDEA 3 · DRSC PACK (one slide)
def idea_drsc():
    s=slide()
    header(s,"02 · Idea 03 · The district knowledge pack",
           [("Joint thought leadership, ",False),("on the table where districts decide",True)],
           lead="A co-branded CFI × TraffiCure pack on the district's key traffic problems — monthly or bi-monthly, tabled at the District Road Safety Committee.",
           sect=(2,"IDEA 03 · DRSC"))
    text(s,ML,I(2.1),I(7.5),I(0.24),"HOW THE FEED BECOMES THE PACK",font=HEAD,size=8.5,bold=True,color=MUT,tracking=1.8)
    rows=[("gauge","Raw speeds & delay","→","Worst five corridors, ranked by measured speed and delay"),
          ("school","School-hour speeds","→","School gates above safe speed at 8 am, named"),
          ("zap","Harsh-braking points","→","The district's top conflict junctions, mapped"),
          ("git-compare","Before / after","→","Did last meeting's action work? Measured, not claimed")]
    y=I(2.42)
    for k,(ic,a,arrow,b) in enumerate(rows):
        yy=y+k*I(0.66)
        rect(s,ML,yy,I(7.5),I(0.56),fill=MIST if k%2==0 else PAPER,rounded=True,radius=0.12)
        icon(s,ic,ML+I(0.16),yy+I(0.15),I(0.26),"#4A35FF")
        text(s,ML+I(0.56),yy+I(0.14),I(2.15),I(0.3),a,font=HEAD,size=9.5,bold=True,color=INK)
        text(s,ML+I(2.78),yy+I(0.12),I(0.3),I(0.3),arrow,font=HEAD,size=12,bold=True,color=BRAND)
        text(s,ML+I(3.15),yy+I(0.14),I(4.2),I(0.3),b,size=9,color=INK,ls=1.1)
    why_us_band(s,5.25,
        "We sit on four DRSCs, and we write evidence the way government reads it. The pack enters the statutory record.",
        "The data no district has ever seen about itself — attributed, every month, in the room.")
    text(s,ML,I(6.34),I(7.5),I(0.55),
         [("The quiet door:  ",{'bold':True,'color':BRAND,'size':9.5}),
          ("around that table sit the district's decision-makers — administration, police, PWD, transport. A page they rely on every month is the most natural introduction a data partnership can have. District trust is where state conversations begin.",{'color':MUT,'size':9.5})],ls=1.25)
    # right: pack mock
    px=I(8.55); pw=I(4.15)
    rect(s,px,I(2.1),pw,I(4.5),fill=PAPER,line=LINE,rounded=True,radius=0.06)
    rect(s,px,I(2.1),pw,I(0.5),fill=BRAND)
    text(s,px+I(0.2),I(2.19),pw-I(0.4),I(0.32),"GURUGRAM ROAD SAFETY BRIEF · SEPT 2026",font=HEAD,size=8.5,bold=True,color=PAPER,tracking=0.8)
    items=[("01","Corridor watch","NH-48 service road slowest at 6 pm; speeds up 12% at night"),
           ("02","School gates","3 gates above 40 km/h during morning school-run"),
           ("03","Conflict points","Iffco Chowk ramp: highest braking density this month"),
           ("04","Action check","Sector 29 signal retiming: -18% delay, held for 6 weeks"),
           ("05","For the record","One recommendation for this meeting, with the data behind it")]
    yy=I(2.78)
    for n,t,d in items:
        text(s,px+I(0.2),yy,I(0.4),I(0.3),n,font=HEAD,size=9.5,bold=True,color=BRAND)
        text(s,px+I(0.62),yy,pw-I(0.85),I(0.3),t,font=HEAD,size=9.5,bold=True,color=INK)
        text(s,px+I(0.62),yy+I(0.24),pw-I(0.85),I(0.4),d,size=8,color=MUT,ls=1.12)
        yy=yy+I(0.72)
    text(s,px+I(0.2),I(6.38),pw-I(0.4),I(0.24),"Illustrative contents · co-branded · one page",size=7.5,color=MUT)
    ftc(s,6)
    return s

# ---------------------------------------------------------------- 07 ABOUT
def about_tiles():
    s=slide()
    header(s,"03 · About us",[("Six lines of work, ",False),("one connected system",True)],
           lead="A nonprofit registered as the Vision Zero Trust · Co-founded by Cars24 and the Indian Road Safety Council · 18 cities · everything below is live.",
           sect=(3,"ABOUT US"))
    tiles=[("Project Rakshak","Student-led road audits, fixed with authorities","31 sites · 15 in implementation",a5("implementation_work.jpg"),0.30),
           ("SATARK enforcement","AI support that helps police find high-risk vehicles","3,00,000+ vehicles screened",a5("satark_billboard.jpg"),0.42),
           ("Victim support","Aasha chatbot, calculator, hospital helpdesks","100+ victims helped in first 4 days",a5("ph_helpdesk_desk.jpg"),0.25),
           ("Open public data","Four self-updating data platforms, free to use","36 states & 50 cities covered",a5("sctracker.png"),0.05),
           ("School Safety Index","Public grading of school-zone road safety","3,814 Delhi NCR schools mapped",ar("slide_ssi_school.png"),0.0),
           ("Youth & policy","NextMile ideathon and national dialogue series","2,000+ participants engaged",a5("ph_ideathon_winners.jpg"),0.25)]
    tw=I(3.95); th=I(2.32); gx=I(0.12); gy=I(0.16)
    for i,(t,d,st,img,anc) in enumerate(tiles):
        r,c=divmod(i,3)
        x=ML+c*(tw+gx); y=I(2.12)+r*(th+gy)
        rect(s,x,y,tw,th,fill=PAPER,line=LINE,rounded=True,radius=0.06)
        photo(s,img,x,y,tw,I(1.28),fill=True,border=False,anchor=anc)
        rect(s,x,y,tw,I(1.28),fill=None,line=LINE,lw=1.0)
        text(s,x+I(0.14),y+I(1.36),tw-I(0.28),I(0.3),t,font=HEAD,size=11.5,bold=True,color=INK,ls=1.0)
        text(s,x+I(0.14),y+I(1.62),tw-I(0.28),I(0.35),d,size=9,color=MUT,ls=1.12)
        text(s,x+I(0.14),y+I(2.02),tw-I(0.28),I(0.26),st,font=HEAD,size=9,bold=True,color=BRAND,ls=1.0)
    ftc(s,7)
    return s

# ---------------------------------------------------------------- 08 SATARK
def satark_slide():
    s=B.satark(8,TOTAL,None)
    patch_rail(s,"02 · WHAT WE HAVE BUILT, ENFORCEMENT TECHNOLOGY","03 · ABOUT US, SATARK IN THE FIELD")
    patch_rail(s,"02 · What we have built, enforcement technology","03 · About us, SATARK in the field")
    patch_rail(s,"WHAT WE HAVE BUILT","ABOUT US")
    return s

# ---------------------------------------------------------------- 09 NEXT STEPS, BOTH SIDES
def nextsteps():
    s=slide()
    header(s,"04 · Next steps",
           [("One district, 90 days, ",False),("three MVPs built together",True)],
           lead="Everything on our side is funded and ready. The ask is small and specific.",
           sect=(4,"NEXT STEPS"))
    x1=ML; w=I(5.9)
    rect(s,x1,I(2.2),w,I(3.05),fill=BR_T,rounded=True,radius=0.06)
    text(s,x1+I(0.22),I(2.38),w-I(0.44),I(0.28),"WHAT WE WILL DO — AT OUR COST",font=HEAD,size=9.5,bold=True,color=BRAND,tracking=1.4)
    left=[("school","SSI MVP: speed-exposure scores on the ~100-zone launch batch; field-verify worst 20"),
          ("map-pin","Rakshak MVP: braking overlay on Jaipur's 96 clusters; live speed panel on the public dashboard"),
          ("landmark","Knowledge pack MVP: first co-branded district brief, tabled at a DRSC we sit on"),
          ("users","Student field teams, verification, government formatting, all follow-through")]
    yy=I(2.78)
    for ic,t in left:
        icon(s,ic,x1+I(0.24),yy+I(0.02),I(0.26),"#4A35FF")
        text(s,x1+I(0.62),yy,w-I(0.84),I(0.55),t,size=9.5,color=INK,ls=1.2)
        yy=yy+I(0.6)
    x2=ML+w+I(0.28)
    rect(s,x2,I(2.2),w,I(3.05),fill=GR_T,rounded=True,radius=0.06)
    text(s,x2+I(0.22),I(2.38),w-I(0.44),I(0.28),"WHAT WE ASK OF TRAFFICURE + GOOGLE",font=HEAD,size=9.5,bold=True,color=GREEN,tracking=1.4)
    right=[("database","Demo access for ONE district — Gurugram or Jaipur: speeds, counts, harsh braking"),
          ("cpu","One technical session: Lepton + Maps + CFI, on our corridors, this week"),
          ("badge-check","Terms for public, attributed use in the index, the dashboard and the pack"),
          ("handshake","A named counterpart to review the MVPs with us at day 90")]
    yy=I(2.78)
    for ic,t in right:
        icon(s,ic,x2+I(0.24),yy+I(0.02),I(0.26),"#0A9955")
        text(s,x2+I(0.62),yy,w-I(0.84),I(0.55),t,size=9.5,color=INK,ls=1.2)
        yy=yy+I(0.6)
    rect(s,ML,I(5.5),CW,I(1.1),fill=MIST,rounded=True,radius=0.10)
    text(s,ML+I(0.25),I(5.64),CW-I(0.5),I(0.85),
         [("The shape of the next 90 days:  ",{'bold':True,'color':INK,'size':11}),
          ("demo access on one district → three MVPs built on real data → reviewed together at day 90 → then we decide, with evidence, how far to go. If an MVP fails its test, we say so and shelve it.",{'color':MUT,'size':10.5})],ls=1.35)
    ftc(s,9)
    return s

# ---------------------------------------------------------------- 10 CLOSE
def close():
    s=slide(bg=BRAND)
    s.shapes.add_picture(LOGO_WHITE,ML,I(0.62),width=I(2.0))
    text(s,ML,I(2.0),I(12.1),I(1.5),"Traffic data already runs the daily commute.\nTogether it can start preventing the crash.",font=HEAD,size=30,bold=True,color=PAPER,ls=1.2)
    text(s,ML,I(3.8),I(11.5),I(0.6),"Three ideas, one district, ninety days. We bring the method, the field force and the government seats; the data does the rest.",size=12.5,color=LAVW,ls=1.4)
    hline(s,ML,I(4.6),I(5.2),color=RGBColor(0x6A,0x55,0xFF),wt=1.0)
    cols=[("Akhtar Hussain","Mission Lead","akhtar@crashfreeindia.org"),
          ("Deepanshu Gupta","Trustee","deepanshu@crashfreeindia.org"),
          ("Gajendra Jangid","Trustee","gajendra@crashfreeindia.org")]
    for j,(n,r,e) in enumerate(cols):
        x=ML+j*I(4.0)
        text(s,x,I(4.9),I(3.8),I(0.3),n,font=HEAD,size=13,bold=True,color=PAPER)
        text(s,x,I(5.24),I(3.8),I(0.28),f"{r} · {e}",size=9.5,color=LAVW)
    text(s,ML,I(7.05),I(12),I(0.3),"Crashfree India · A Cars24 commitment · Operated by Vision Zero Trust, New Delhi · crashfreeindia.org",size=9.5,color=LAVW)
    return s

if __name__=="__main__":
    new_deck()
    cover()          # 01
    tc_who()         # 02
    tc_data()        # 03
    idea_ssi()       # 04
    idea_rakshak()   # 05
    idea_drsc()      # 06
    about_tiles()    # 07
    satark_slide()   # 08
    nextsteps()      # 09
    close()          # 10
    save(os.path.join(BUILD,"CFI_Trafficure_v3.pptx"))
