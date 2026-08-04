# Crashfree India — generic organisation overview deck (shareable), 17 slides.
# Reuses the NRSB white-paper design system and initiative slides, with a
# neutral footer/section tab, fresh cover/close, and the expert bench minus
# Abhay Damle.
import os
import nrsb2_style as ST
import build_nrsb3 as B
from nrsb2_style import *
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

I=Inches
V6=os.path.join(SP,"assets/v6/ready")
def a6(n): return os.path.join(V6,n)

TOTAL=17

# --- neutral footer + section tab (override the EoI-specific ones) ---
def _footer(s,wg,idx,total):
    hline(s,ML,I(7.06),CW)
    text(s,ML,I(7.14),I(9),I(0.25),"Crashfree India  ·  Organisation Overview  ·  2026",size=8.5,color=MUT,ls=1.0)
    text(s,SW-MR-I(1.2),I(7.14),I(1.2),I(0.25),f"{idx:02d} / {total:02d}",size=8.5,color=MUT,align=PP_ALIGN.RIGHT,ls=1.0)
def _sect_tab(s,cur,name):
    w=I(3.6)
    text(s,SW-MR-w,I(0.30),w,I(0.24),name,font=HEAD,size=8.5,bold=True,color=MUT,tracking=1.6,align=PP_ALIGN.RIGHT)
B.footer=_footer
ST.sect_tab=_sect_tab
circp=B.circp  # circular-photo helper lives in build_nrsb3

def ov_footer(s,idx): _footer(s,None,idx,TOTAL)

# ---------------------------------------------------------------- 01 COVER
def cover():
    s=slide()
    s.shapes.add_picture(LOGO_BLUE,ML,I(0.5),width=I(1.85))
    text(s,SW-MR-I(3.5),I(0.55),I(3.5),I(0.3),"ORGANISATION OVERVIEW  ·  2026",font=HEAD,size=9,bold=True,color=MUT,tracking=1.8,align=PP_ALIGN.RIGHT)
    text(s,ML,I(1.72),I(8.2),I(0.3),"A ROAD SAFETY NONPROFIT  ·  NEW DELHI",font=HEAD,size=10.5,bold=True,color=BRAND,tracking=2.2)
    text(s,ML,I(2.14),I(8.3),I(0.85),"Crashfree India",font=HEAD,size=46,bold=True,color=INK,ls=1.0)
    text(s,ML,I(3.1),I(8.3),I(0.85),"Adding capacity to India's road safety system:\nevidence, institutions, and technology that follows through",font=HEAD,size=16.5,bold=True,color=BRAND,ls=1.25)
    rows=[("search-check","Field audits and research to official standards, fixed with the authorities that own the road"),
          ("cpu","Technology for police, victims and the public: enforcement, open data, and a chatbot on WhatsApp"),
          ("landmark","Working inside the machinery: 4 district committees, police MoUs, 25+ written approvals")]
    y=I(4.65)
    icons=["search-check","cpu","landmark"]
    for i,(ic,t) in enumerate(rows):
        icon_disc(s,ic,ML,y+i*I(0.62),I(0.4),BR_T,"#4A35FF")
        text(s,ML+I(0.56),y+i*I(0.62)+I(0.05),I(8.0),I(0.4),t,font=HEAD,size=11.5,bold=True,color=INK,ls=1.1)
    text(s,ML,I(6.9),I(8.6),I(0.3),"Registered as the Vision Zero Trust · A Cars24 commitment · Co-founded with the Indian Road Safety Council",size=9.5,color=MUT)
    text(s,SW-MR-I(2.6),I(6.9),I(2.6),I(0.3),"crashfreeindia.org",font=HEAD,size=9.5,bold=True,color=BRAND,align=PP_ALIGN.RIGHT)
    px,pyy,pw,ph_=I(9.1),I(1.15),I(3.6),I(4.9)
    photo(s,a5("ph_dhoni_white.jpg"),px,pyy,pw,ph_,fill=True,anchor=0.12)
    rect(s,px,pyy+ph_-I(0.56),pw,I(0.56),fill=BRAND)
    text(s,px,pyy+ph_-I(0.52),pw,I(0.5),"MS DHONI · GOODWILL AMBASSADOR,\nCRASHFREE INDIA",font=HEAD,size=8.5,bold=True,color=PAPER,align=PP_ALIGN.CENTER,ls=1.15,tracking=1.0)
    return s

# ---------------------------------------------------------------- 05 RAKSHAK (fresh)
def rakshak_ov(idx):
    s=slide()
    header(s,"02 · What we have built, road infrastructure",
           [("Project Rakshak: student audits ",False),("that end in built fixes",True)],
           lead="Trained student teams, guided by senior experts, audit dangerous locations to IRC codes, and we walk each fix through the road-owning authority until it is built.",
           sect=(2,"WHAT WE HAVE BUILT"))
    fun=[("120+","locations screened",MIST),("31","sites audited, 11 official black spots",MIST),
         ("25+","written authority approvals",MIST),("15","sites in implementation, 7 fully built",None)]
    y=I(2.2)
    for j,(v,l,fill) in enumerate(fun):
        x=ML+j*I(2.1)
        if fill is None:
            rect(s,x,y,I(1.95),I(1.25),fill=BRAND,rounded=True,radius=0.12)
            text(s,x+I(0.14),y+I(0.12),I(1.7),I(0.55),v,font=HEAD,size=24,bold=True,color=PAPER,ls=1.0)
            text(s,x+I(0.14),y+I(0.62),I(1.7),I(0.55),l,size=8.5,color=PAPER,ls=1.1)
        else:
            rect(s,x,y,I(1.95),I(1.25),fill=MIST,rounded=True,radius=0.12)
            text(s,x+I(0.14),y+I(0.12),I(1.7),I(0.55),v,font=HEAD,size=24,bold=True,color=BRAND,ls=1.0)
            text(s,x+I(0.14),y+I(0.62),I(1.7),I(0.55),l,size=8.5,color=MUT,ls=1.1)
        if j<3: text(s,x+I(1.93),y+I(0.4),I(0.2),I(0.4),"→",size=14,color=MUT)
    text(s,ML,I(3.7),I(8.2),I(0.3),"31 student teams · 20 institutions including nine IITs · 900+ stakeholder surveys · 17 subject-matter experts guiding every audit",size=10,color=MUT,ls=1.3)
    rect(s,ML,I(4.25),I(8.2),I(1.35),fill=GR_T,rounded=True,radius=0.10)
    text(s,ML+I(0.22),I(4.38),I(7.8),I(0.3),"MARQUEE CASE, OUTER RING ROAD, SOUTH DELHI",font=HEAD,size=8.5,bold=True,color=GREEN,tracking=1.2)
    text(s,ML+I(0.22),I(4.7),I(7.8),I(0.8),
         "On our audit, PWD built a continuous 1.8 m accessible pedestrian corridor, upgraded drainage at ponding sites, installed bollards against footpath encroachment, and deployed camera-backed speed-feedback boards.",size=10,color=INK,ls=1.3)
    text(s,ML,I(5.8),I(8.2),I(0.7),
         [("Every site is tracked on a public dashboard ",{'bold':True}),
          ("from submission through approval to completed construction, an honest read on how quickly agencies can act.",{'color':MUT})],size=10,ls=1.3)
    photo(s,a5("implementation_work.jpg"),I(9.1),I(2.15),I(3.6),I(4.35),fill=True,anchor=0.3,
          caption="Construction underway after a student audit, tracked publicly to completion.")
    ov_footer(s,idx)
    return s

# ---------------------------------------------------------------- 09 GIG (fresh)
def gig_ov(idx):
    s=slide()
    header(s,"02 · What we have built, vulnerable road users",
           [("The fastest-growing risk on Indian roads ",False),("rides a two-wheeler",True)],
           lead="Two-wheeler riders carry the largest share of road deaths, and gig and delivery work concentrates that exposure. We study it, and we work it.",
           sect=(2,"WHAT WE HAVE BUILT"))
    tiles=[("~45%","of India's crash fatalities are two-wheeler users, rising year on year"),
           ("23.5 M","gig and platform workers projected by 2029-30 (NITI Aayog)"),
           ("1 in 2","delivery riders in a 431-rider Mumbai crash study had already crashed")]
    tw=I(2.6)
    for j,(v,l) in enumerate(tiles):
        x=ML+j*(tw+I(0.12)); y=I(2.35)
        rect(s,x,y,tw,I(1.5),fill=MIST,rounded=True,radius=0.08)
        text(s,x+I(0.16),y+I(0.12),tw-I(0.32),I(0.45),v,font=HEAD,size=20,bold=True,color=BRAND)
        text(s,x+I(0.16),y+I(0.62),tw-I(0.32),I(0.8),l,size=8.5,color=MUT,ls=1.2)
    rows=[("file-search","A published study","'When Risks Become Routine' (Feb 2026): 40 pages, 87 references, from incentive design and helmet compliance to the e-bike regulatory blind spot.",AM_T,"#E58900"),
          ("gauge","A 300-rider study, now in the field","Quantifying how often riders face near-misses, which risk factors matter most, and what rider profiles exist, so interventions can be targeted.",TE_T,"#0E8A8A"),
          ("bike","Field, not only desk","Fieldwork included shadowing a delivery rider through a full nine-hour shift, documented on video.",BR_T,"#4A35FF")]
    y=I(4.15)
    for j,(ic,t,d,tint,col) in enumerate(rows):
        irow(s,ML,y+j*I(0.82),I(7.7),ic,t,d,tint,col,tsize=11,dsize=9)
    photo(s,a5("ph_gig_riders.jpg"),I(8.95),I(4.05),I(3.75),I(2.0),fill=True,anchor=0.35,
          caption="Delivery-rider outreach, Road Safety Month 2026.")
    ov_footer(s,idx)
    return s

# ---------------------------------------------------------------- 10 RESEARCH (fresh)
def research_ov(idx):
    s=slide()
    header(s,"02 · What we have built, our research",
           [("Published evidence ",False),("the sector cites",True)],
           lead="Original research, launched publicly and covered by national press. Each report turns a policy gap into something authorities can act on.",
           sect=(2,"WHAT WE HAVE BUILT"))
    reports=[("cov_justice_unserved.png","'Justice Unserved'","Jan 2026","Why so few hit-and-run and third-party claims reach the families entitled to them. Launched at IIT Delhi; findings presented to MoRTH."),
             ("cov_deposited.png","'Deposited But Not Delivered'","2026","~₹986 crore of court-awarded compensation lies unclaimed across five High Court jurisdictions, a bottleneck the Supreme Court noted in 2025."),
             ("primer_cover-01.jpg","Pre-Hospital Care Primer","Sep 2025","Benchmarks ERSS-112 response and documents state EMS innovations, with contributions from the AIIMS Trauma Centre.")]
    tw=I(3.95)
    for j,(img,t,dt,d) in enumerate(reports):
        x=ML+j*(tw+I(0.12)); y=I(2.35)
        rect(s,x,y,tw,I(4.1),fill=PAPER,line=LINE,rounded=True,radius=0.06)
        photo_fit(s,a6(img),x+I(0.2),y+I(0.2),tw-I(0.4),I(2.2),border=False,align='center')
        text(s,x+I(0.2),y+I(2.55),tw-I(0.4),I(0.3),t,font=HEAD,size=12,bold=True,color=INK,ls=1.0)
        chip(s,x+I(0.2),y+I(2.9),I(1.3),I(0.32),dt,fill=BR_T,color=BRAND,size=8)
        text(s,x+I(0.2),y+I(3.32),tw-I(0.4),I(0.7),d,size=8.5,color=MUT,ls=1.2)
    text(s,ML,I(6.6),CW,I(0.3),
         [("Plus: ",{'bold':True}),("the gig-rider brief 'When Risks Become Routine', and four open public-data platforms, all free to use and openly licensed.",{'color':MUT})],size=9.5,ls=1.2)
    ov_footer(s,idx)
    return s

# ---------------------------------------------------------------- 16 BENCH (fresh, minus Damle)
def bench_ov(idx):
    s=slide()
    header(s,"03 · Our people, the expert bench",
           [("Advised by the people who ",False),("built India's road safety practice",True)],
           sect=(3,"OUR PEOPLE"))
    y=I(1.7)
    rect(s,ML,y,I(6.0),I(1.02),fill=MIST,rounded=True,radius=0.12)
    circp(s,"x_gtiwari.png",ML+I(0.14),y+I(0.14),I(0.74))
    text(s,ML+I(1.02),y+I(0.13),I(4.9),I(0.35),"Prof. Geetam Tiwari, Board of Advisors",font=HEAD,size=11,bold=True,color=INK)
    text(s,ML+I(1.02),y+I(0.45),I(4.9),I(0.55),"Professor Emerita & Chair, TRIP Centre, IIT Delhi · four decades of transport safety research",size=9,color=MUT,ls=1.2)
    x2=I(6.85)
    rect(s,x2,y,I(5.85),I(1.02),fill=MIST,rounded=True,radius=0.12)
    circp(s,"x_tewari.png",x2+I(0.14),y+I(0.14),I(0.74))
    text(s,x2+I(1.02),y+I(0.13),I(4.7),I(0.35),"Piyush Tewari, Board of Advisors",font=HEAD,size=11,bold=True,color=INK)
    text(s,x2+I(1.02),y+I(0.45),I(4.7),I(0.55),"Founder & CEO, SaveLIFE Foundation · Architect of the Good Samaritan Law · Zero-Fatality Corridors",size=9,color=MUT,ls=1.2)
    text(s,ML,I(2.95),I(11),I(0.26),"JURY, MENTORS & REVIEWERS AROUND OUR WORK",font=HEAD,size=8.5,bold=True,color=MUT,tracking=1.6)
    heads=[("midha","Justice J.R. Midha","Delhi High Court (Retd.)"),
           ("dange","Vaibhav Dange","Ex-Advisor, MoRTH"),("tandon","Dr. Maya Tandon","Sahayta Trust"),
           ("ratnam","R. Venkat Ratnam","Government of Punjab"),("asrivastava","Akhilesh Srivastava","ITS India · IRF"),
           ("raman","Aishwarya Raman","OMI Foundation"),("nkumar","Neeraj Kumar","Tata AIG"),
           ("pandey","R.S. Pandey","FICCI Road Safety"),("dash","Dipak K. Dash","Times of India")]
    cw_=I(2.42); y0=I(3.28)
    for i,(f,n,o) in enumerate(heads):
        r,c=divmod(i,5)
        x=ML+c*cw_; yy=y0+r*I(1.32)
        circp(s,"x_"+f+".png",x+I(0.86),yy,I(0.7))
        text(s,x,yy+I(0.74),cw_-I(0.1),I(0.3),n,font=HEAD,size=8.5,bold=True,color=INK,align=PP_ALIGN.CENTER,ls=1.0)
        text(s,x,yy+I(0.95),cw_-I(0.1),I(0.3),o,size=7.5,color=MUT,align=PP_ALIGN.CENTER,ls=1.0)
    rect(s,ML,I(6.08),CW,I(0.68),fill=MIST,rounded=True,radius=0.14)
    text(s,ML+I(0.22),I(6.15),CW-I(0.45),I(0.55),
         [("Engaged through the National Road Safety Dialogue Series: ",{'bold':True}),
          ("Balraj Bhanot (ex-Chair, CMVR-TSC) · I.V. Rao (TERI; ex-Maruti Suzuki) · D.P. Saste (CIRT / ex-MoRTH) · Dr. Sushma Sagar (AIIMS Trauma) · Dr. S. Velmurugan (CSIR-CRRI), and jurors from OMI, NIUA, WRI, CII and Tata AIG.",{'color':MUT})],size=9,ls=1.25)
    ov_footer(s,idx)
    return s

# ---------------------------------------------------------------- 17 CLOSE (fresh)
def close_ov():
    s=slide(bg=BRAND)
    s.shapes.add_picture(LOGO_WHITE,ML,I(0.62),width=I(2.0))
    text(s,ML,I(2.1),I(12.1),I(1.4),"No life lost on an Indian road\nis acceptable.",font=HEAD,size=32,bold=True,color=PAPER,ls=1.2)
    text(s,ML,I(3.85),I(11.5),I(0.6),"We add the hands, the evidence and the tools that turn that goal into work the system can carry, and we track it in public.",size=12.5,color=LAVW,ls=1.4)
    hline(s,ML,I(4.72),I(5.2),color=RGBColor(0x6A,0x55,0xFF),wt=1.0)
    text(s,ML,I(4.94),I(10.8),I(0.85),
         '"We already know what is going wrong. What we need now is more responsibility, more discipline, more respect for life."',
         font=HEAD,size=14,bold=True,color=PAPER,ls=1.35,italic=True)
    text(s,ML,I(5.66),I(8),I(0.3),"MS Dhoni · Goodwill Ambassador, Crashfree India",size=10,color=LAVW)
    rows=[("globe","crashfreeindia.org"),("mail","akhtar@crashfreeindia.org  ·  helpdesk@crashfreeindia.org")]
    y=I(6.14)
    for ic,t in rows:
        icon(s,ic,ML,y,I(0.28),"#C8C0FF")
        text(s,ML+I(0.45),y-I(0.02),I(11),I(0.35),t,size=11,color=PAPER,ls=1.2)
        y=y+I(0.46)
    text(s,ML,I(7.05),I(12),I(0.3),"Crashfree India · A Cars24 commitment · Operated by Vision Zero Trust, New Delhi",size=9.5,color=LAVW)
    return s

if __name__=="__main__":
    new_deck()
    cover()                         # 01
    B.about(2,TOTAL,None)           # 02
    B.approach(3,TOTAL,None)        # 03
    B.portfolio(4,TOTAL,None)       # 04
    rakshak_ov(5)                   # 05
    B.satark(6,TOTAL,None)          # 06
    B.opendata(7,TOTAL,None)        # 07
    B.victim_suite(8,TOTAL,None)    # 08
    gig_ov(9)                       # 09
    research_ov(10)                 # 10
    B.youth(11,TOTAL,None)          # 11
    B.ground(12,TOTAL,None)         # 12
    B.dhoni(13,TOTAL,None)          # 13
    B.partners(14,TOTAL,None)       # 14
    B.team(15,TOTAL,None)           # 15
    bench_ov(16)                    # 16
    close_ov()                      # 17
    save(os.path.join(BUILD,"CFI_Overview_v1.pptx"))
