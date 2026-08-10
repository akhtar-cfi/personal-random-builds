# Crashfree India — mentor briefing deck v2, 16 slides.
# v2 changes: tech+governance cover with focus-area line, tech+policy about
# slide, no funding amounts, gig-rider slide restored, new School Safety Index
# and press-coverage slides.
import os
import nrsb2_style as ST
import build_nrsb3 as B
import build_overview as OV
from build_overview import rakshak_ov, research_ov, close_ov
from nrsb2_style import *
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

I = Inches
TOTAL = 16
RMI = os.path.join(SP, "assets/rmi")
PRESS = os.path.join(SP, "assets/press")
def ap(n): return os.path.join(PRESS, n)

def m_footer(s, wg, idx, total):
    hline(s, ML, I(7.06), CW)
    text(s, ML, I(7.14), I(9), I(0.25), "Crashfree India  ·  A briefing for mentors  ·  2026", size=8.5, color=MUT, ls=1.0)
    text(s, SW-MR-I(1.2), I(7.14), I(1.2), I(0.25), f"{idx:02d} / {total:02d}", size=8.5, color=MUT, align=PP_ALIGN.RIGHT, ls=1.0)

B.footer = m_footer
OV._footer = m_footer
OV.TOTAL = TOTAL

# ---------------------------------------------------------------- 01 COVER
def cover_m():
    s = slide()
    s.shapes.add_picture(LOGO_BLUE, ML, I(0.5), width=I(1.85))
    text(s, SW-MR-I(3.5), I(0.55), I(3.5), I(0.3), "A BRIEFING FOR MENTORS  ·  2026", font=HEAD, size=9, bold=True, color=MUT, tracking=1.8, align=PP_ALIGN.RIGHT)
    text(s, ML, I(1.66), I(8.2), I(0.3), "A ROAD SAFETY NONPROFIT  ·  NEW DELHI", font=HEAD, size=10.5, bold=True, color=BRAND, tracking=2.2)
    text(s, ML, I(2.08), I(8.3), I(0.85), "Crashfree India", font=HEAD, size=46, bold=True, color=INK, ls=1.0)
    text(s, ML, I(3.04), I(8.3), I(0.85), "Technology and governance for road safety,\nbuilt into an institution designed to make itself redundant", font=HEAD, size=16, bold=True, color=BRAND, ls=1.25)
    rows = [("cpu", "Digital public goods: enforcement AI, open data platforms, victim tools on WhatsApp"),
            ("landmark", "Governance from inside: 4 district committee seats, police MoUs, 25+ written approvals"),
            ("building-2", "Standards, manuals and dashboards built so the system can run this work without us")]
    y = I(4.52)
    for i, (ic, t) in enumerate(rows):
        icon_disc(s, ic, ML, y+i*I(0.60), I(0.4), BR_T, "#4A35FF")
        text(s, ML+I(0.56), y+i*I(0.60)+I(0.05), I(7.85), I(0.4), t, font=HEAD, size=11.5, bold=True, color=INK, ls=1.1)
    text(s, ML, I(6.42), I(8.4), I(0.3),
         [("FOCUS AREAS   ", {'bold': True, 'color': BRAND, 'tracking': 1.6, 'size': 9}),
          ("Road infrastructure · Enforcement · Crash compensation · Gig rider safety · School zones · Open data", {'color': INK, 'size': 9.5})],
         font=HEAD, bold=True)
    text(s, ML, I(6.92), I(8.6), I(0.3), "Registered as the Vision Zero Trust · A Cars24 commitment · Co-founded with the Indian Road Safety Council", size=9.5, color=MUT)
    text(s, SW-MR-I(2.6), I(6.92), I(2.6), I(0.3), "crashfreeindia.org", font=HEAD, size=9.5, bold=True, color=BRAND, align=PP_ALIGN.RIGHT)
    px, pyy, pw, ph_ = I(9.1), I(1.15), I(3.6), I(4.9)
    photo(s, a5("ph_dhoni_white.jpg"), px, pyy, pw, ph_, fill=True, anchor=0.12)
    rect(s, px, pyy+ph_-I(0.56), pw, I(0.56), fill=BRAND)
    text(s, px, pyy+ph_-I(0.52), pw, I(0.5), "MS DHONI · GOODWILL AMBASSADOR,\nCRASHFREE INDIA", font=HEAD, size=8.5, bold=True, color=PAPER, align=PP_ALIGN.CENTER, ls=1.15, tracking=1.0)
    return s

# ---------------------------------------------------------------- 02 PERSONAL NOTE
def note(idx):
    s = slide()
    text(s, ML, I(0.30), I(8), I(0.24), "BEFORE THE SLIDES", font=HEAD, size=9, bold=True, color=BRAND, tracking=2.2)
    hline(s, ML, I(0.62), I(0.55), color=BRAND, wt=2.0)
    text(s, ML, I(0.76), I(12), I(0.5), [("A personal ", {'color': INK}), ("note", {'color': BRAND})], font=HEAD, size=22.5, bold=True, ls=1.1)
    x = ML+I(0.95); w = I(10.2); y = I(1.5)
    rect(s, x, y, w, I(5.42), fill=PAPER, line=LINE, lw=1.2)
    rect(s, x, y, w, I(0.06), fill=BRAND)
    s.shapes.add_picture(LOGO_BLUE, x+I(0.5), y+I(0.32), width=I(1.35))
    text(s, x+w-I(3.2), y+I(0.36), I(2.7), I(0.3), "New Delhi · August 2026", size=9.5, color=MUT, align=PP_ALIGN.RIGHT)
    body = ("Dear friend,\n"
            "This deck goes to a small set of people whose judgement I deeply admire. I am writing to ask for something specific: your counsel, from time to time, as Crashfree India grows.\n"
            "Crashfree India is a road safety nonprofit, registered as the Vision Zero Trust, co-founded with the Indian Road Safety Council and backed by Cars24. India loses more than 1.7 lakh lives on its roads every year. In our first year we built an audit-to-implementation engine that has moved 15 dangerous sites into rectification, an AI enforcement platform running with three police forces, a compensation practice that follows crash victims' families from FIR to disbursal, and open data platforms that turn scattered public records into citable evidence. MS Dhoni carries this message to the country as our Goodwill Ambassador.\n"
            "The work has reached the stage where ambition can outrun experience. We are deciding which programmes become institutions, how to work with government at scale, and how to build an organisation that can hold this mission for a decade. On questions like these, we have more energy than experience, and we know it.\n"
            "The pages that follow are a short tour of what we have built. If the work speaks to you, I would be grateful for an hour of your time, and for the licence to write to you when we face a decision where your experience would show us the way.")
    tb = text(s, x+I(0.5), y+I(0.92), w-I(1.0), I(3.7), body, size=10.4, color=INK, ls=1.28)
    for p in tb.text_frame.paragraphs: p.space_after = Pt(6.5)
    text(s, x+I(0.5), y+I(4.56), w-I(1.0), I(0.3), "With respect,", size=10.4, color=INK)
    text(s, x+I(0.5), y+I(4.84), w-I(1.0), I(0.3), "Akhtar Hussain", font=HEAD, size=11.5, bold=True, color=INK)
    text(s, x+I(0.5), y+I(5.08), w-I(1.0), I(0.26), "Mission Lead, Crashfree India  ·  akhtar@crashfreeindia.org  ·  crashfreeindia.org", size=9.5, color=MUT)
    m_footer(s, None, idx, TOTAL)
    return s

# ---------------------------------------------------------------- 03 ABOUT (tech + policy)
def about_m(idx):
    s = slide()
    header(s, "01 · About us", [("Crashfree India: technology and policy, ", False), ("built to be handed over", True)],
           lead="Registered as the Vision Zero Trust · Co-founded by Cars24 and the Indian Road Safety Council · Working since June 2025 · New Delhi",
           sect=(1, "ABOUT US"))
    rows = [("search-check", "Evidence first", "Structured field audits, file reviews and research to official standards, not opinion.", GR_T, "#0A9955"),
            ("cpu", "Technology as public goods", "Enforcement AI, open data platforms and WhatsApp tools, built with authorities and free for any of them to adopt.", TE_T, "#0E8A8A"),
            ("scale", "Policy that lands", "Research written so officials can act on it: briefed to ministries and district committees, cited by national newsrooms.", BR_T, "#4A35FF"),
            ("building-2", "Built to make ourselves redundant", "Standards, audit manuals and public dashboards designed so institutions can run the work without us.", AM_T, "#E58900")]
    y = I(2.2)
    for i, (ic, t, d, tint, col) in enumerate(rows):
        irow(s, ML, y+i*I(0.82), I(7.6), ic, t, d, tint, col)
    rect(s, ML, I(5.65), I(7.55), I(0.92), fill=BR_T, rounded=True, radius=0.16)
    text(s, ML+I(0.25), I(5.72), I(7.1), I(0.8),
         [("Proof in year one: ", {'bold': True, 'color': BRAND}),
          ("31 sites audited · 25+ written authority approvals · implementation started at 15 sites · 3,00,000+ vehicles screened · seats on 4 district committees", {'color': BRAND})],
         size=11, ls=1.3)
    photo(s, a5("cfi_map_dots.png"), I(8.85), I(2.15), I(3.6), caption="Where we work today, 18 cities, north to south.")
    m_footer(s, None, idx, TOTAL)
    return s

# ---------------------------------------------------------------- 04 PORTFOLIO (focus areas)
def portfolio_m(idx):
    s = slide()
    header(s, "02 · What we have built", [("Six lines of work, ", False), ("one connected system", True)],
           lead="Audits create data · data supports policy · policy guides the field. Six focus areas, one system.",
           sect=(2, "WHAT WE HAVE BUILT"))
    tiles = [("Project Rakshak", "Student-led road audits, fixed with authorities", "31 sites · 7 fixes fully built", a5("implementation_work.jpg"), 0.30),
             ("SATARK enforcement", "AI support that helps police find high-risk vehicles", "3,00,000+ vehicles screened", a5("satark_billboard.jpg"), 0.42),
             ("Crash compensation", "A policy implementation gap we close with tech, embedded in districts and DRSCs", "Rajasthan portal · 4 DRSC seats", a5("ph_helpdesk_desk.jpg"), 0.25),
             ("Open public data", "Four self-updating data platforms, free to use", "36 states & 50 cities covered", a5("sctracker.png"), 0.05),
             ("Gig rider safety", "Research-led: published brief, 300-rider study, platform engagement", "1 in 2 of 431 riders had crashed", a5("ph_gig_riders.jpg"), 0.35),
             ("School Safety Index", "A public 0–100 score for the road outside every school gate", "Delhi & Jaipur · 5,386 schools mapped", os.path.join(RMI, "slide_ssi_index.png"), 0.0)]
    tw = I(3.95); th = I(2.32); gx = I(0.12); gy = I(0.16)
    for i, (t, d, st, img, anc) in enumerate(tiles):
        r, c = divmod(i, 3)
        x = ML + c*(tw+gx); y = I(2.12) + r*(th+gy)
        rect(s, x, y, tw, th, fill=PAPER, line=LINE, rounded=True, radius=0.06)
        photo(s, img, x, y, tw, I(1.28), fill=True, border=False, anchor=anc)
        rect(s, x, y, tw, I(1.28), fill=None, line=LINE, lw=1.0)
        text(s, x+I(0.14), y+I(1.36), tw-I(0.28), I(0.3), t, font=HEAD, size=11.5, bold=True, color=INK, ls=1.0)
        text(s, x+I(0.14), y+I(1.62), tw-I(0.28), I(0.35), d, size=9, color=MUT, ls=1.12)
        text(s, x+I(0.14), y+I(2.02), tw-I(0.28), I(0.26), st, font=HEAD, size=9, bold=True, color=BRAND, ls=1.0)
    m_footer(s, None, idx, TOTAL)
    return s

# ---------------------------------------------------------------- 10 GIG RIDERS (with report snippet)
def gig_m(idx):
    s = slide()
    header(s, "02 · What we have built, vulnerable road users",
           [("The fastest-growing risk on Indian roads ", False), ("rides a two-wheeler", True)],
           lead="Two-wheeler riders carry the largest share of road deaths, and gig and delivery work concentrates that exposure. We study it, and we work it.",
           sect=(2, "WHAT WE HAVE BUILT"))
    tiles = [("~45%", "of India's crash fatalities are two-wheeler users, rising year on year"),
             ("23.5 M", "gig and platform workers projected by 2029-30 (NITI Aayog)"),
             ("1 in 2", "delivery riders in a 431-rider Mumbai crash study had already crashed")]
    tw = I(2.6)
    for j, (v, l) in enumerate(tiles):
        x = ML+j*(tw+I(0.12)); y = I(2.35)
        rect(s, x, y, tw, I(1.5), fill=MIST, rounded=True, radius=0.08)
        text(s, x+I(0.16), y+I(0.12), tw-I(0.32), I(0.45), v, font=HEAD, size=20, bold=True, color=BRAND)
        text(s, x+I(0.16), y+I(0.62), tw-I(0.32), I(0.8), l, size=8.5, color=MUT, ls=1.2)
    rows = [("file-search", "A published study", "'When Risks Become Routine' (Feb 2026): 40 pages, 87 references, from incentive design and helmet compliance to the e-bike regulatory blind spot.", AM_T, "#E58900"),
            ("gauge", "A 300-rider study, now in the field", "Quantifying how often riders face near-misses, which risk factors matter most, and what rider profiles exist, so interventions can be targeted.", TE_T, "#0E8A8A"),
            ("bike", "Field, not only desk", "Fieldwork included shadowing a delivery rider through a full nine-hour shift, documented on video.", BR_T, "#4A35FF")]
    y = I(4.15)
    for j, (ic, t, d, tint, col) in enumerate(rows):
        irow(s, ML, y+j*I(0.82), I(7.7), ic, t, d, tint, col, tsize=11, dsize=9)
    photo_fit(s, ap("cov_gig.png"), I(8.95), I(2.15), I(3.75), I(4.15), align='center',
              caption="'When Risks Become Routine' · Crashfree India research brief, February 2026.")
    m_footer(s, None, idx, TOTAL)
    return s

# ---------------------------------------------------------------- 11 SCHOOL SAFETY INDEX
def ssi(idx):
    s = slide()
    header(s, "02 · What we have built, school zones",
           [("School Safety Index: a public score ", False), ("for every school's road", True)],
           lead="A child's walk to school can be graded. We grade it against India's own standard, publish the evidence, and hand the worst zones to field teams and the road-owning authority.",
           sect=(2, "WHAT WE HAVE BUILT"))
    rows = [("scan-line", "Screening at city scale", "Dated Street View imagery, scored 0–100 by an AI auditor against IRC:SP:32-2023, India's school-zone standard, and independently verified before publication.", BR_T, "#4A35FF"),
            ("school", "Delhi and Jaipur first", "5,386 schools mapped across the two cities. Named schools and per-school photo evidence published from the first batch, re-audited every six months.", TE_T, "#0E8A8A"),
            ("footprints", "It measures the road, never the school", "A low score is a demand on the road-owning authority, not a criticism of the school. The worst zones feed our student field audits, the 41-check ground tier.", GR_T, "#0A9955")]
    y = I(2.42)
    for i, (ic, t, d, tint, col) in enumerate(rows):
        irow(s, ML, y+i*I(0.98), I(7.55), ic, t, d, tint, col, tsize=11, dsize=9)
    tiles = [("~10,000", "minors died on Indian roads in 2023 (MoRTH)"),
             ("4.6%", "of crashes occur in school and college zones (e-DAR 2024)"),
             ("730+", "schools graded from imagery worldwide, validating the method (iRAP)")]
    tw = I(2.42)
    for j, (v, l) in enumerate(tiles):
        x = ML+j*(tw+I(0.14)); ty = I(5.52)
        rect(s, x, ty, tw, I(1.24), fill=MIST, rounded=True, radius=0.10)
        text(s, x+I(0.15), ty+I(0.10), tw-I(0.3), I(0.4), v, font=HEAD, size=17, bold=True, color=BRAND)
        text(s, x+I(0.15), ty+I(0.52), tw-I(0.3), I(0.68), l, size=8, color=MUT, ls=1.15)
    photo(s, os.path.join(RMI, "slide_ssi_index.png"), I(8.55), I(2.42), I(4.15), I(3.62), fill=True, anchor=0.0,
          caption="The pilot dashboard: score-first, evidence-pinned. Delhi and Jaipur batches are in build now.")
    m_footer(s, None, idx, TOTAL)
    return s

# ---------------------------------------------------------------- 12 IN THE NEWS
def media(idx):
    s = slide()
    header(s, "02 · What we have built, in the news",
           [("The work, ", False), ("as the press tells it", True)],
           lead="National newsrooms cite our research, cover our deployments, and quote our researchers by name.",
           sect=(2, "WHAT WE HAVE BUILT"))
    cards = [
        ("strip_dh_ju.jpg", "DECCAN HERALD  ·  MARCH 2026",
         "Road crash victims struggle as ₹80,000 crore compensation lies unpaid",
         "National coverage of our 'Justice Unserved' brief; DH hosts the full report for readers."),
        ("strip_wire_survivors.jpg", "THE WIRE  ·  INDIASPEND  ·  MAY 2026",
         "For India's road accident survivors, legal compensation is elusive",
         "Quotes our research lead Kesar Kanjhlia on where claims stall."),
        ("strip_indiaspend_crash.jpg", "INDIASPEND  ·  ISIGNAL  ·  MARCH 2026",
         "One crash, years of crisis: the hidden scale of serious road injuries",
         "Cites our gig-rider research; quotes Aastha Shreeharsh on missing risk data."),
        ("strip_dh_billboard.jpg", "DECCAN HERALD  ·  SEPT 2025",
         "Big boss watching! Billboard at Trinity Circle displays violations in real time",
         "Our live challan billboard with Bengaluru Traffic Police and Cars24."),
        ("strip_tli_billboard.jpg", "THE LOGICAL INDIAN  ·  SEPT 2025",
         "India's first AI-powered billboard to display pending challans in real time",
         "Also covered by Siasat, CarToq, Media Infoline and Daily Jagran."),
        (None, "FROM OUR POLICY DESK",
         "A steady line of published argument",
         "Articles and explainers by Shubham Kumar and our policy team, from claim pipelines to enforcement design."),
    ]
    cw, chh = I(3.95), I(2.14)
    gx, gy = I(0.135), I(0.14)
    y0 = I(2.26)
    for k, (img, outlet, head, desc) in enumerate(cards):
        r, c = divmod(k, 3)
        x = ML + c*(cw+gx); y = y0 + r*(chh+gy)
        rect(s, x, y, cw, chh, fill=PAPER, line=LINE, lw=1.0, rounded=True, radius=0.06)
        if img:
            photo(s, ap(img), x, y, cw, I(0.86), fill=True, border=False, anchor=0.0)
            rect(s, x, y, cw, I(0.86), fill=None, line=LINE, lw=1.0)
            text(s, x+I(0.18), y+I(0.98), cw-I(0.36), I(0.22), outlet, font=HEAD, size=7.5, bold=True, color=BRAND, tracking=1.2)
            text(s, x+I(0.18), y+I(1.22), cw-I(0.36), I(0.5), head, font=HEAD, size=9.5, bold=True, color=INK, ls=1.12)
            text(s, x+I(0.18), y+I(1.76), cw-I(0.36), I(0.34), desc, size=7.5, color=MUT, ls=1.15)
        else:
            rect(s, x, y, cw, chh, fill=BR_T, rounded=True, radius=0.06)
            text(s, x+I(0.2), y+I(0.2), cw-I(0.4), I(0.22), outlet, font=HEAD, size=7.5, bold=True, color=BRAND, tracking=1.2)
            text(s, x+I(0.2), y+I(0.5), cw-I(0.4), I(0.6), head, font=HEAD, size=11.5, bold=True, color=INK, ls=1.15)
            text(s, x+I(0.2), y+I(1.14), cw-I(0.4), I(0.85), desc, size=8.5, color=MUT, ls=1.25)
    text(s, ML, I(6.72), CW, I(0.3),
         [("Also in the record:  ", {'bold': True, 'color': INK}),
          ("'Justice Unserved' launched at IIT Delhi and presented to MoRTH · launch coverage in Team-BHP, Adgully and The CSR Universe.", {'color': MUT})],
         size=8.5, ls=1.2)
    m_footer(s, None, idx, TOTAL)
    return s

# ---------------------------------------------------------------- 15 YEAR AHEAD + GUIDANCE
def guidance(idx):
    s = slide()
    header(s, "04 · The year ahead",
           [("Where we are going, and ", False), ("where your counsel would help", True)],
           lead="Year 2 moves from programmes to institutions. These are the goals, and what we hope to learn from your experience.",
           sect=(4, "THE YEAR AHEAD"))
    colw = I(5.85)
    text(s, ML, I(2.02), colw, I(0.26), "THE YEAR AHEAD", font=HEAD, size=9, bold=True, color=MUT, tracking=1.8)
    rows_l = [
        ("hard-hat", "Project Rakshak", "Audits running inside institutions, with 75+ implementations initiated and every one tracked publicly."),
        ("scale", "Crash compensation", "Claim pipeline fixed in two districts, mapped in eight, with a playbook states can adopt."),
        ("cpu", "SATARK", "From 3 cities toward 30, as shared public infrastructure for traffic police."),
        ("rocket", "Two exploratory bets", "School Safety Index and gig rider safety, each scaled only if a small, honest pilot earns it."),
        ("building-2", "Institutions that outlive us", "An expert consortium, an audit manual and a publishing cadence that outlive any single cohort."),
    ]
    y = I(2.38)
    for i, (ic, t, d) in enumerate(rows_l):
        irow(s, ML, y+i*I(0.86), colw, ic, t, d, BR_T, "#4A35FF")
    x2 = ML+colw+I(0.42)
    w2 = SW-MR-x2
    text(s, x2, I(2.02), w2, I(0.26), "WHERE WE SEEK YOUR GUIDANCE", font=HEAD, size=9, bold=True, color=MUT, tracking=1.8)
    rows_r = [
        ("building-2", "How you built institutions", "Organisations that outlast their founders: the standards, governance and culture that takes."),
        ("zap", "How you solved the hardest problems", "The judgement earned on ambiguous, high-stakes problems. We would bring you ours early."),
        ("landmark", "Working with government at scale", "Partnerships that deepen while we keep our independence and publish everything."),
        ("users", "People and leadership", "Growing a young team into leaders, and drawing senior operators to a nonprofit mission."),
        ("target", "Focus", "What to stop, what to scale, and when. The discipline of a small organisation with a national mission."),
    ]
    for i, (ic, t, d) in enumerate(rows_r):
        irow(s, x2, y+i*I(0.86), w2, ic, t, d, GR_T, "#0A9955")
    m_footer(s, None, idx, TOTAL)
    return s

if __name__ == "__main__":
    new_deck()
    cover_m()                      # 01
    note(2)                        # 02
    about_m(3)                     # 03
    portfolio_m(4)                 # 04
    rakshak_ov(5)                  # 05
    B.satark(6, TOTAL, None)       # 06
    B.opendata(7, TOTAL, None)     # 07
    B.victim_suite(8, TOTAL, None) # 08
    research_ov(9)                 # 09
    gig_m(10)                      # 10
    ssi(11)                        # 11
    media(12)                      # 12
    B.youth(13, TOTAL, None)       # 13
    B.team(14, TOTAL, None)        # 14
    guidance(15)                   # 15
    close_ov()                     # 16
    save(os.path.join(BUILD, "CFI_Mentor_v3.pptx"))
