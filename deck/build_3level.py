#!/usr/bin/env python3
"""CFI AOP 2026-27 — 3-level bifurcation (Vertical > Activity > Sub-Activity) from the actual Rs 6 Cr budget."""
import json
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INDIGO="4A35FF"; INK="1A1C1C"; MUT="777589"; MIST="EDEEF2"; TINT="F3F1FF"
TINT2="E9E6FF"; BAND="FAFAFB"; AMBER="E58900"; WHITE="FFFFFF"
RUPEE='"₹"#,##0'; LACS='0.0'; PCT='0.0%'

thin=Side(style="thin", color="E2E3EC")
border=Border(left=thin,right=thin,top=thin,bottom=thin)

items=json.load(open('items.json'))

# ordered tree V > A > S
tree=OrderedDict()
for it in items:
    tree.setdefault(it['v'],OrderedDict()).setdefault(it['a'],OrderedDict()).setdefault(it['s'],0.0)
    tree[it['v']][it['a']][it['s']]+=it['b']

wb=Workbook()

# ================= Sheet 3 (built first for references): Line Items =================
wl=wb.active; wl.title="Line Items"
wl.sheet_view.showGridLines=False
def H(ws,c,txt,fill=INK,color=WHITE,size=10,align="center"):
    cc=ws.cell(*c) if isinstance(c,tuple) else ws[c]
    cc.value=txt; cc.font=Font(name="Calibri",size=size,bold=True,color=color)
    cc.fill=PatternFill("solid",fgColor=fill); cc.alignment=Alignment(horizontal=align,vertical="center",wrap_text=True); cc.border=border
    return cc
wl.merge_cells("A1:G1")
H(wl,"A1","CFI · AOP 2026–27 · Budget line items (source of truth · 173 rows · ₹6.0 Cr)",fill=INDIGO,align="left",size=11)
wl.row_dimensions[1].height=26
heads=["Vertical","Activity","Sub-Activity","Budget Header","Description","Budget (₹)","In lacs"]
for j,h in enumerate(heads,1): H(wl,(2,j),h)
wl.row_dimensions[2].height=26
r=3
for it in items:
    b = r%2==0
    fill = WHITE if b else BAND
    def put(col,val,align="left",fmt=None,size=9,color=INK,bold=False,wrap=True):
        cc=wl.cell(row=r,column=col,value=val)
        cc.font=Font(name="Calibri",size=size,bold=bold,color=color)
        cc.alignment=Alignment(horizontal=align,vertical="center",wrap_text=wrap)
        cc.fill=PatternFill("solid",fgColor=fill); cc.border=border
        if fmt: cc.number_format=fmt
    put(1,it['v']); put(2,it['a']); put(3,it['s']); put(4,it['h'])
    put(5,it['desc'],size=8,color=MUT)
    put(6,it['b'],align="right",fmt=RUPEE,wrap=False)
    put(7,f"=F{r}/100000",align="right",fmt=LACS,wrap=False)
    r+=1
Ndata_last=r-1
# total row
tr=r
tc=wl.cell(row=tr,column=1,value="Grand Total"); tc.font=Font(name="Calibri",size=10,bold=True); tc.fill=PatternFill("solid",fgColor=MIST); tc.border=border
for col in (2,3,4,5):
    c=wl.cell(row=tr,column=col); c.fill=PatternFill("solid",fgColor=MIST); c.border=border
c6=wl.cell(row=tr,column=6,value=f"=SUM(F3:F{Ndata_last})"); c6.font=Font(bold=True); c6.number_format=RUPEE; c6.alignment=Alignment(horizontal="right"); c6.fill=PatternFill("solid",fgColor=MIST); c6.border=border
c7=wl.cell(row=tr,column=7,value=f"=F{tr}/100000"); c7.font=Font(bold=True); c7.number_format=LACS; c7.alignment=Alignment(horizontal="right"); c7.fill=PatternFill("solid",fgColor=MIST); c7.border=border
for col,w in zip("ABCDEFG",[20,30,30,30,46,14,9]): wl.column_dimensions[col].width=w
wl.freeze_panes="A3"
LI_A=f"'Line Items'!$A$3:$A${Ndata_last}"
LI_B=f"'Line Items'!$B$3:$B${Ndata_last}"
LI_C=f"'Line Items'!$C$3:$C${Ndata_last}"
LI_F=f"'Line Items'!$F$3:$F${Ndata_last}"
GRAND=f"SUM({LI_F})"

# ================= Sheet 1: 3-Level Bifurcation =================
wb3=wb.create_sheet("3-Level Bifurcation",0)
wb3.sheet_view.showGridLines=False
wb3.sheet_properties.outlinePr.summaryBelow=False
wb3.merge_cells("A1:E1")
H(wb3,"A1","CFI · AOP 2026–27 · 3-level bifurcation   (Vertical › Activity › Sub-Activity · ₹6.0 Cr)",fill=INDIGO,align="left",size=12)
wb3.row_dimensions[1].height=30
for j,h in enumerate(["Particulars","Level","Budget (₹)","In lacs","% of AOP"],1): H(wb3,(2,j),h)
wb3.row_dimensions[2].height=24

def line(row, particular, level, fill, indent, budget_formula, pct=True, bold=False, size=10, color=INK, olevel=0, hidden=False):
    a=wb3.cell(row=row,column=1,value=particular)
    a.font=Font(name="Calibri",size=size,bold=bold,color=color)
    a.alignment=Alignment(horizontal="left",vertical="center",indent=indent,wrap_text=True)
    a.fill=PatternFill("solid",fgColor=fill); a.border=border
    lc=wb3.cell(row=row,column=2,value=level)
    lc.font=Font(name="Calibri",size=8,bold=False,color=MUT); lc.alignment=Alignment(horizontal="center",vertical="center")
    lc.fill=PatternFill("solid",fgColor=fill); lc.border=border
    bc=wb3.cell(row=row,column=3,value=budget_formula)
    bc.font=Font(name="Calibri",size=size,bold=bold,color=color); bc.number_format=RUPEE
    bc.alignment=Alignment(horizontal="right",vertical="center"); bc.fill=PatternFill("solid",fgColor=fill); bc.border=border
    lac=wb3.cell(row=row,column=4,value=f"=C{row}/100000")
    lac.font=Font(name="Calibri",size=size,bold=bold,color=color); lac.number_format=LACS
    lac.alignment=Alignment(horizontal="right",vertical="center"); lac.fill=PatternFill("solid",fgColor=fill); lac.border=border
    pc=wb3.cell(row=row,column=5,value=(f"=C{row}/$C${GRAND_ROW}" if pct else None))
    pc.font=Font(name="Calibri",size=size,bold=bold,color=(INDIGO if bold else MUT)); pc.number_format=PCT
    pc.alignment=Alignment(horizontal="right",vertical="center"); pc.fill=PatternFill("solid",fgColor=fill); pc.border=border
    if olevel: wb3.row_dimensions[row].outline_level=olevel
    if hidden: wb3.row_dimensions[row].hidden=False

# grand total row placed at top (row 3) for quick read
GRAND_ROW=3
line(3,"TOTAL — Annual Operating Plan 2026–27","",TINT2,0,f"={GRAND}",pct=True,bold=True,size=11)
r=4
# helper keys in hidden cols G/H/I
def keys(row,v=None,a=None,s=None):
    if v is not None: wb3.cell(row=row,column=7,value=v)
    if a is not None: wb3.cell(row=row,column=8,value=a)
    if s is not None: wb3.cell(row=row,column=9,value=s)

for v,acts in tree.items():
    keys(r,v=v)
    line(r,v,"Vertical",TINT,0,f"=SUMIF({LI_A},$G{r},{LI_F})",pct=True,bold=True,size=10.5,color=INDIGO)
    r+=1
    for a,subs in acts.items():
        keys(r,v=v,a=a)
        line(r,a,"Activity",MIST,1,f"=SUMIFS({LI_F},{LI_A},$G{r},{LI_B},$H{r})",pct=True,bold=True,size=10,olevel=1)
        r+=1
        for s,amt in subs.items():
            keys(r,v=v,a=a,s=s)
            label=s if s else "(unspecified)"
            line(r,label,"Sub-activity",WHITE,2,f"=SUMIFS({LI_F},{LI_A},$G{r},{LI_B},$H{r},{LI_C},$I{r})",pct=False,bold=False,size=9,color=INK,olevel=2)
            r+=1
for col,w in zip("ABCDE",[58,12,15,9,10]): wb3.column_dimensions[col].width=w
for col in ("G","H","I"): wb3.column_dimensions[col].hidden=True
wb3.freeze_panes="A3"

# ================= Sheet 2: By Vertical =================
wv=wb.create_sheet("By Vertical",1)
wv.sheet_view.showGridLines=False
wv.merge_cells("A1:D1")
H(wv,"A1","AOP 2026–27 · Summary by vertical",fill=INDIGO,align="left",size=12)
wv.row_dimensions[1].height=28
for j,h in enumerate(["Vertical","Budget (₹)","In lacs","% of AOP"],1): H(wv,(2,j),h)
wv.row_dimensions[2].height=22
verts=list(tree.keys())
vr=3
gr=vr+len(verts)
for i,v in enumerate(verts):
    row=vr+i; fill=WHITE if i%2==0 else BAND
    a=wv.cell(row=row,column=1,value=v); a.font=Font(bold=True,size=10); a.alignment=Alignment(horizontal="left",vertical="center"); a.fill=PatternFill("solid",fgColor=fill); a.border=border
    b=wv.cell(row=row,column=2,value=f"=SUMIF({LI_A},A{row},{LI_F})"); b.number_format=RUPEE; b.alignment=Alignment(horizontal="right"); b.fill=PatternFill("solid",fgColor=fill); b.border=border
    c=wv.cell(row=row,column=3,value=f"=B{row}/100000"); c.number_format=LACS; c.alignment=Alignment(horizontal="right"); c.fill=PatternFill("solid",fgColor=fill); c.border=border
    d=wv.cell(row=row,column=4,value=f"=B{row}/$B${gr}"); d.number_format=PCT; d.alignment=Alignment(horizontal="right"); d.fill=PatternFill("solid",fgColor=fill); d.border=border
a=wv.cell(row=gr,column=1,value="Grand Total"); a.font=Font(bold=True); a.fill=PatternFill("solid",fgColor=MIST); a.border=border
b=wv.cell(row=gr,column=2,value=f"=SUM(B{vr}:B{gr-1})"); b.font=Font(bold=True); b.number_format=RUPEE; b.alignment=Alignment(horizontal="right"); b.fill=PatternFill("solid",fgColor=MIST); b.border=border
c=wv.cell(row=gr,column=3,value=f"=B{gr}/100000"); c.font=Font(bold=True); c.number_format=LACS; c.alignment=Alignment(horizontal="right"); c.fill=PatternFill("solid",fgColor=MIST); c.border=border
d=wv.cell(row=gr,column=4,value=f"=B{gr}/$B${gr}"); d.font=Font(bold=True); d.number_format=PCT; d.alignment=Alignment(horizontal="right"); d.fill=PatternFill("solid",fgColor=MIST); d.border=border
for col,w in zip("ABCD",[42,16,10,10]): wv.column_dimensions[col].width=w

wb.move_sheet("3-Level Bifurcation", -wb.sheetnames.index("3-Level Bifurcation"))
out="CFI_AOP_3-Level-Bifurcation_v1_2026-08-19.xlsx"
wb.save(out)
print("saved",out,"| sheets:",wb.sheetnames,"| line-item last row:",Ndata_last)
