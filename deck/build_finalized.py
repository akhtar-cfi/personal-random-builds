#!/usr/bin/env python3
"""CFI AOP 2026-27 — 3-level bifurcation from the FINALIZED Rs 6.00 Cr AOP (programme verticals)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INDIGO="4A35FF"; INK="1A1C1C"; MUT="777589"; MIST="EDEEF2"; TINT="F3F1FF"
TINT2="E9E6FF"; BAND="FAFAFB"; AMBER="E58900"; WHITE="FFFFFF"; GREEN="0A9955"
RUPEE='"₹"#,##0'; LACS='0.0'; PCT='0.0%'
thin=Side(style="thin", color="E2E3EC"); border=Border(left=thin,right=thin,top=thin,bottom=thin)

# ---------------------------------------------------------------------------
# DATA: vertical -> (vertical_note) -> list of (activity, [ (line, amount, description) ])
# Level-1 verticals + totals are EXACT from the board-approved AOP budget slide.
# Trust Admin's 5 headers are EXACT from the AOP org-backbone slide.
# Level-2/3 splits within programmes are a first-pass (flagged in Description / Notes).
# ---------------------------------------------------------------------------
DATA = [
 ("Project Rakshak", "Core programme · audit-to-implementation pipeline", [
    ("Institutional track (IIT / NIT / SPA teams)", [
       ("Student team stipends", 3240000, "₹36,000 per team on completion · ~90 institutional sites"),
       ("Expert mentor honoraria", 1300000, "Faculty & road-safety expert mentors across the cohort"),
       ("Campus outreach, onboarding & merchandise", 960000, "Chapter mobilisation, IDs, kits"),
    ]),
    ("NGO partner track (~20 sites)", [
       ("NGO field-team engagement", 2000000, "Field audits via NGO partners"),
       ("CFI desk + 10–15% physical re-checks", 500000, "QA on NGO submissions"),
    ]),
    ("Certified RSA track (~40 sites)", [
       ("CRRI-certified auditor fees (IRC SP:88)", 4500000, "Full audits usable by NHAI without rework"),
    ]),
    ("Expert Consortium & Audit Manual v1.0", [
       ("Expert Consortium (4 meetings)", 1000000, "Sep / Dec / Mar / Jun sittings"),
       ("Audit Manual v1.0 production", 800000, "Authoring, design, print"),
    ]),
    ("Public accountability dashboard", [
       ("Dashboard build, hosting & maintenance", 1200000, "Identification → audit → review → implementation tracker"),
    ]),
    ("Programme management & field ops", [
       ("Programme management (staff allocation)", 1800000, "PM + coordination"),
       ("Field travel & logistics", 700000, "Site visits, DRSC meetings"),
       ("National forum & DRSC presentations", 500000, "12+ DRSC presentations, year-end forum"),
    ]),
 ]),
 ("Hit & Run Compensation", "Core programme · fix two districts, map eight", [
    ("Scheme audit & process documentation", [
       ("Stage-wise pipeline audit fieldwork (Gurugram)", 1200000, "FIR → GI Council disbursal; 20–30 claims tracked end-to-end"),
       ("'Where the Pipeline Breaks' report", 600000, "Scheme audit submitted to MoRTH"),
    ]),
    ("Institutional capacity building", [
       ("Fellow-in-Residence stipend (Tehsildar office)", 1200000, "Embedded fellow · research-embedded, not social work"),
       ("SOP reference cards & training-gap assessment", 800000, "District capacity building"),
    ]),
    ("Victim identification & AASHA", [
       ("DLSA-partnered helpdesks", 800000, "Trauma-ward legal helpdesks"),
       ("AASHA QR deployment", 600000, "Police stations & trauma centres; usage measured monthly"),
    ]),
    ("Policy audit & advocacy", [
       ("Quarterly findings & public brief", 500000, "Publishing cadence"),
       ("Advocacy & MoRTH engagement travel", 300000, "Stakeholder engagement"),
    ]),
 ]),
 ("Civic & GovTech (SATARK)", "Core programme · enforcement co-pilot + data stack", [
    ("SATARK scale to 30 cities", [
       ("SATARK engineering & product", 2500000, "ANPR, ULIP, rule engine, billboard + officer app"),
       ("City deployment & rollout", 2000000, "Toward 30 cities · Gurugram, Pune, Ahmedabad first"),
       ("Cloud, ANPR & ULIP integration", 1500000, "Compute, APIs, camera integration"),
    ]),
    ("Open data stack (4 products)", [
       ("Open data product maintenance & hosting", 1500000, "Defect Repository, Parliament tracker, SC Litigation Tracker, Crash Dashboard"),
       ("Data engineering & updates", 1000000, "Daily / weekly refresh, 146 districts"),
    ]),
    ("GovTech exploration & partnerships", [
       ("GovTech scoping & partnerships", 1000000, "2+ new DPG opportunities scoped"),
       ("Pilot / DPG prototyping", 500000, "ULIP / MoRTH ecosystem"),
    ]),
 ]),
 ("Marketing Campaign", "Enabler · 'Safety is the New Swag' (all four pillars now funded)", [
    ("Media advocacy", [
       ("Editorial partnerships, radio & podcast, PSA", 1500000, "Earned + owned media"),
    ]),
    ("Digital advocacy", [
       ("50 micro-influencers & branded short films", 1400000, "#SafeSwag creator programme"),
       ("Digital IP & content production", 600000, "Short films, reels"),
    ]),
    ("Offline awareness", [
       ("OOH in 3–4 cities", 1000000, "Billboards, transit"),
       ("Guerilla installations & event co-branding", 500000, ""),
    ]),
    ("Community mobilization", [
       ("College competitions & app-based engagement", 1000000, "Real-life stories"),
    ]),
 ]),
 ("Bets (SSI + Gig Rider)", "Seed-funded, stage-gated · scale only if the pilot earns it", [
    ("School Safety Index — seed (₹0.10 Cr)", [
       ("IRC:SP:32 rubric validation & Zone Captain pilot", 700000, "~100 school zones in one anchor city"),
       ("Per-school costing & board-ready design", 300000, "Gate to a funded multi-city index in Year 3"),
    ]),
    ("Gig Rider Safety — seed (₹0.10 Cr)", [
       ("Gig crash-data scoping study", 600000, "National study parked until field model proven"),
       ("Platform + OMI Foundation engagements", 400000, "3+ platforms at the table"),
    ]),
 ]),
 ("Trust Admin (Org Backbone)", "Backbone · 25% of AOP · figures exact from AOP org-backbone slide", [
    ("Core team salaries & benefits", [
       ("Core team salaries & benefits", 4200000, "Aggregate core staff cost (individual figures withheld)"),
       ("Safer Mobility Impact Fellowship (SMIF)", 3600000, "★ Fellow cohort, aggregate (~6 fellows across tracks)"),
    ]),
    ("IRSC institutional support & expert consultation", [
       ("IRSC institutional support & expert consultation", 4200000, "Shared expert / leadership consultation"),
    ]),
    ("IT, SaaS & AI tooling", [
       ("AI tooling — Claude (Team) + Codex / AI subscriptions", 600000, "★ AI assistants for build, research & ops"),
       ("Google Workspace & email", 300000, ""),
       ("Slack & collaboration", 200000, ""),
       ("HRMS 'CFI People' & other SaaS", 200000, ""),
       ("Cloud, hosting & domains", 100000, ""),
    ]),
    ("M&E, statutory & impact audit", [
       ("M&E, statutory & impact audit", 600000, ""),
    ]),
    ("Travel, branding & compliance", [
       ("Travel, branding & compliance", 1100000, ""),
    ]),
 ]),
 ("Contingency (4%)", "Reserve", [
    ("Contingency", [
       ("Contingency reserve (4%)", 2400000, "Held against programme + admin spend"),
    ]),
 ]),
]

wb=Workbook()

# ============ Sheet 1: 3-Level Bifurcation ============
ws=wb.active; ws.title="3-Level Bifurcation"
ws.sheet_view.showGridLines=False
ws.sheet_properties.outlinePr.summaryBelow=False
ws.merge_cells("A1:F1")
t=ws["A1"]; t.value="CFI · AOP 2026–27 · 3-level budget bifurcation   (Vertical › Activity › Cost line · ₹6.00 Cr)"
t.font=Font(name="Calibri",size=12,bold=True,color=WHITE); t.fill=PatternFill("solid",fgColor=INDIGO)
t.alignment=Alignment(horizontal="left",vertical="center"); ws.row_dimensions[1].height=30
heads=["Particulars","Level","Budget (₹)","In lacs","% of AOP","Description / basis"]
for j,h in enumerate(heads,1):
    c=ws.cell(row=2,column=j,value=h); c.font=Font(name="Calibri",size=10,bold=True,color=WHITE)
    c.fill=PatternFill("solid",fgColor=INK); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=border
ws.row_dimensions[2].height=24

GRAND_ROW=3  # grand total placed at row 3
def style(row, particular, level, fill, indent, budget, pct, desc, bold, size, color, olevel, is_input):
    a=ws.cell(row=row,column=1,value=particular); a.font=Font(name="Calibri",size=size,bold=bold,color=color)
    a.alignment=Alignment(horizontal="left",vertical="center",indent=indent,wrap_text=True); a.fill=PatternFill("solid",fgColor=fill); a.border=border
    lc=ws.cell(row=row,column=2,value=level); lc.font=Font(name="Calibri",size=8,color=MUT); lc.alignment=Alignment(horizontal="center",vertical="center"); lc.fill=PatternFill("solid",fgColor=fill); lc.border=border
    bc=ws.cell(row=row,column=3,value=budget); bc.font=Font(name="Calibri",size=size,bold=bold,color=color); bc.number_format=RUPEE
    bc.alignment=Alignment(horizontal="right",vertical="center"); bc.fill=PatternFill("solid",fgColor=fill); bc.border=border
    dc=ws.cell(row=row,column=4,value=f"=C{row}/100000"); dc.font=Font(name="Calibri",size=size,bold=bold,color=color); dc.number_format=LACS
    dc.alignment=Alignment(horizontal="right",vertical="center"); dc.fill=PatternFill("solid",fgColor=fill); dc.border=border
    pc=ws.cell(row=row,column=5,value=(f"=C{row}/$C${GRAND_ROW}" if pct else None)); pc.font=Font(name="Calibri",size=size,bold=bold,color=(INDIGO if bold else MUT)); pc.number_format=PCT
    pc.alignment=Alignment(horizontal="right",vertical="center"); pc.fill=PatternFill("solid",fgColor=fill); pc.border=border
    ec=ws.cell(row=row,column=6,value=desc); ec.font=Font(name="Calibri",size=8.5,color=MUT); ec.alignment=Alignment(horizontal="left",vertical="center",wrap_text=True); ec.fill=PatternFill("solid",fgColor=fill); ec.border=border
    if olevel: ws.row_dimensions[row].outline_level=olevel

# grand total row (row 3)
r=3
l1_rows=[]
style(3,"TOTAL — Annual Operating Plan 2026–27 (₹6.00 Cr)","",TINT2,0,None,True,"71% into programmes · 25% backbone · 4% contingency",True,11,INK,0,False)
r=4
for vert, vnote, acts in DATA:
    l1_row=r; l1_rows.append(l1_row)
    style(r,vert,"Vertical",TINT,0,None,True,vnote,True,10.5,INDIGO,0,False); r+=1
    l2_rows=[]
    for act, lines in acts:
        l2_row=r; l2_rows.append(l2_row)
        style(r,act,"Activity",MIST,1,None,True,"",True,10,INK,1,False); r+=1
        l3_first=r
        for line, amt, desc in lines:
            style(r,line,"Cost line",WHITE,2,amt,False,desc,False,9,INK,2,True); r+=1
        l3_last=r-1
        ws.cell(row=l2_row,column=3).value=f"=SUM(C{l3_first}:C{l3_last})"
    # L1 subtotal = sum of its L2 subtotal cells
    ws.cell(row=l1_row,column=3).value="="+"+".join(f"C{x}" for x in l2_rows)
# grand = sum of L1 rows
ws.cell(row=GRAND_ROW,column=3).value="="+"+".join(f"C{x}" for x in l1_rows)

for col,w in zip("ABCDEF",[52,10,15,9,9,50]): ws.column_dimensions[col].width=w
ws.freeze_panes="A3"

# ============ Sheet 2: By Vertical ============
wv=wb.create_sheet("By Vertical")
wv.sheet_view.showGridLines=False
wv.merge_cells("A1:D1")
t=wv["A1"]; t.value="AOP 2026–27 · Summary by vertical (Level 1)"; t.font=Font(size=12,bold=True,color=WHITE)
t.fill=PatternFill("solid",fgColor=INDIGO); t.alignment=Alignment(horizontal="left",vertical="center"); wv.row_dimensions[1].height=28
for j,h in enumerate(["Vertical","Budget (₹)","In lacs","% of AOP"],1):
    c=wv.cell(row=2,column=j,value=h); c.font=Font(bold=True,color=WHITE); c.fill=PatternFill("solid",fgColor=INK)
    c.alignment=Alignment(horizontal="center",vertical="center"); c.border=border
vr=3
for i,(l1_row,(vert,_,_)) in enumerate(zip(l1_rows,DATA)):
    row=vr+i; fill=WHITE if i%2==0 else BAND
    a=wv.cell(row=row,column=1,value=vert); a.font=Font(bold=True,size=10); a.alignment=Alignment(horizontal="left",vertical="center"); a.fill=PatternFill("solid",fgColor=fill); a.border=border
    b=wv.cell(row=row,column=2,value=f"='3-Level Bifurcation'!C{l1_row}"); b.number_format=RUPEE; b.alignment=Alignment(horizontal="right"); b.fill=PatternFill("solid",fgColor=fill); b.border=border
    c=wv.cell(row=row,column=3,value=f"=B{row}/100000"); c.number_format=LACS; c.alignment=Alignment(horizontal="right"); c.fill=PatternFill("solid",fgColor=fill); c.border=border
    d=wv.cell(row=row,column=4,value=f"=B{row}/$B${vr+len(DATA)}"); d.number_format=PCT; d.alignment=Alignment(horizontal="right"); d.fill=PatternFill("solid",fgColor=fill); d.border=border
gr=vr+len(DATA)
a=wv.cell(row=gr,column=1,value="Grand Total"); a.font=Font(bold=True); a.fill=PatternFill("solid",fgColor=MIST); a.border=border
b=wv.cell(row=gr,column=2,value=f"=SUM(B{vr}:B{gr-1})"); b.font=Font(bold=True); b.number_format=RUPEE; b.alignment=Alignment(horizontal="right"); b.fill=PatternFill("solid",fgColor=MIST); b.border=border
c=wv.cell(row=gr,column=3,value=f"=B{gr}/100000"); c.font=Font(bold=True); c.number_format=LACS; c.alignment=Alignment(horizontal="right"); c.fill=PatternFill("solid",fgColor=MIST); c.border=border
d=wv.cell(row=gr,column=4,value=f"=B{gr}/$B${gr}"); d.font=Font(bold=True); d.number_format=PCT; d.alignment=Alignment(horizontal="right"); d.fill=PatternFill("solid",fgColor=MIST); d.border=border
for col,w in zip("ABCD",[34,16,10,10]): wv.column_dimensions[col].width=w

# ============ Sheet 3: Notes & basis ============
wn=wb.create_sheet("Notes & basis")
wn.sheet_view.showGridLines=False; wn.column_dimensions["A"].width=112
notes=[
 ("CFI · AOP 2026–27 — 3-level bifurcation · basis & assumptions", INDIGO, 13, True),
 ("", WHITE,10,False),
 ("EXACT (from the board-approved AOP, 'From programmes to institutions'):", GREEN,11,True),
 ("• The 7 Level-1 verticals and their totals: Rakshak ₹1.85 Cr · Hit & Run ₹0.60 Cr · Civic & GovTech (SATARK) ₹1.00 Cr ·", INK,10,False),
 ("   Marketing Campaign ₹0.60 Cr · Bets ₹0.20 Cr · Trust Admin ₹1.51 Cr · Contingency ₹0.24 Cr = ₹6.00 Cr.", INK,10,False),
 ("• Trust Admin's five headers (salaries ₹0.78 Cr, IRSC support ₹0.42 Cr, IT/SaaS/AI ₹0.14 Cr, M&E ₹0.06 Cr,", INK,10,False),
 ("   travel/branding/compliance ₹0.11 Cr) are taken verbatim from the org-backbone slide.", INK,10,False),
 ("", WHITE,10,False),
 ("FIRST-PASS (designed by me to sum exactly to each vertical — edit freely, everything rolls up):", AMBER,11,True),
 ("• All Level-2 activities and Level-3 cost lines inside the programmes (Rakshak, Hit & Run, SATARK, Campaign, Bets).", INK,10,False),
 ("• The salaries split into core staff (₹42 L) vs SMIF fellowship (₹36 L) within the ₹0.78 Cr salaries line.", INK,10,False),
 ("• The itemisation of the ₹0.14 Cr IT/SaaS/AI tooling line into individual tools.", INK,10,False),
 ("", WHITE,10,False),
 ("FELLOW COSTS kept visible, as requested (marked ★):", INDIGO,11,True),
 ("• Rakshak — Student team stipends: ₹36,000 per team on completion (~90 institutional sites) [confirmed in Slack].", INK,10,False),
 ("• Hit & Run — Fellow-in-Residence stipend at the Tehsildar office.", INK,10,False),
 ("• Trust Admin — Safer Mobility Impact Fellowship (SMIF) cohort, shown aggregate (~6 fellows across Strategy & Ops,", INK,10,False),
 ("   Policy, Research, Cause Marketing tracks). Individual salaries are deliberately NOT itemised (privacy).", INK,10,False),
 ("", WHITE,10,False),
 ("AI TOOLING category (as requested):", INDIGO,11,True),
 ("• Sits under Trust Admin › IT, SaaS & AI tooling. 'AI tooling' is its own line: Claude Team plan + Codex / AI subscriptions.", INK,10,False),
 ("   (Anthropic receipts are recurring monthly and forwarded to finances; a Codex 2–3 seat subscription was requested via Cars24.)", INK,10,False),
 ("• The ₹0.14 Cr split across AI, Workspace, Slack, HRMS and cloud is indicative — adjust to actual run-rate.", INK,10,False),
 ("", WHITE,10,False),
 ("Sources: finalized AOP deck (v-final, ₹6.00 Cr) · Team Charter / Ownership Map (Aug 2026) · Slack #rakshak-updates,", MUT,9,False),
 ("#cfi-core, #builders-cfi · Gmail (Anthropic receipts, SMIF fellowship, Codex request).", MUT,9,False),
]
for i,(txt,color,size,bold) in enumerate(notes,1):
    c=wn.cell(row=i,column=1,value=txt); c.font=Font(name="Calibri",size=size,bold=bold,color=color)
    c.alignment=Alignment(horizontal="left",vertical="center")

out="CFI_AOP_Bifurcation-Finalized_v1_2026-08-19.xlsx"
wb.save(out)
print("saved",out)
