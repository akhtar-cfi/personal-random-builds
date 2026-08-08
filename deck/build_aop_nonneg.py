#!/usr/bin/env python3
"""Add the 'Non-negotiables in each project' slide to the AOP team-share deck."""
import io, re, copy
from pptx import Presentation
from pptx.util import Inches as I, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

BRAND = RGBColor(0x4A, 0x35, 0xFF)
INK   = RGBColor(0x1A, 0x1C, 0x1C)
MUT   = RGBColor(0x77, 0x75, 0x89)
LINE  = RGBColor(0xED, 0xEE, 0xF2)
DISC  = RGBColor(0xF3, 0xF1, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SRC = 'aop_nocost.pptx'
OUT = 'aop_nng.pptx'

prs = Presentation(SRC)

# ---- harvest icon blobs from the KPI slide (index 10) ----
icons = {}
for sh in prs.slides[10].shapes:
    if sh.shape_type == 13:
        d = sh._element.nvPicPr.cNvPr.get('descr') or ''
        for key in ('hard-hat', 'scale', 'radar', 'sprout'):
            if d.startswith(key):
                icons[key] = sh.image.blob
assert set(icons) == {'hard-hat', 'scale', 'radar', 'sprout'}, icons.keys()

# ---- content ----
def B(lead, rest):
    return (lead, rest)

RAKSHAK = [
    B('Location selection is gated:', 'only crash-prone locations to be assigned and selected this year, with data validation done beforehand.'),
    B('Audit quality top notch:', 'the lever we go all in on. If needed, we will put in resources and money for evaluation, validation and top-notch audits.'),
    B('Solutions actually implemented:', 'actually life-saving road design. Our North Star metric is solutions implemented, not just approved. Not minor traffic lights or small fixes: actual budgets, actually fixing the road.'),
    B('Approval means a signed letter', 'explicitly accepting at least one recommendation. Acknowledgement, presentation or verbal interest does not count.'),
    B('Implemented means verified:', 'dated before-and-after photos with location proof, or authority confirmation. A team’s own update alone is never proof.'),
    B('Activity is not outcome:', 'audits done, approvals, implementation reported and implementation verified stay separate numbers, never one headline.'),
    B('The location outlives the team:', 'every site has one case with full history. Students graduate; the case and the authority relationship stay with us.'),
    B('Scale never dilutes quality:', 'more teams only after audit quality holds at the current scale.'),
]
COMP = [
    B('Actually fix the claim rate', 'in at least 2 districts; the change should be directly visible in the numbers.'),
    B('Metric to measure:', 'not meetings held, trainings conducted, or reports and briefs published, but actual change in claim rate.'),
    B('A clear path to scale up', 'in other states and nationally going forward; every fix designed so another district can adopt it.'),
    B('Follow real claims end to end:', '20–30 families tracked from FIR to money in the bank. That is where the truth of the process lives.'),
    B('Fix the stage where claims die:', 'stage-wise audit of the scheme pipeline, not general awareness drives.'),
    B('Depth before breadth:', 'go deep in the model districts first; replicate only what has actually worked there.'),
    B('Changes land in government systems:', 'SOPs, formats and portals that officials themselves use. Authorities keep every substantive decision.'),
    B('Honest denominators:', 'claim rate measured the same way every quarter, with the population and as-of date defined.'),
]
SATARK = [
    B('Direct, measurable change in violations:', 'that is the metric, not screenings, downloads or dashboards.'),
    B('Police remain the decision-makers:', 'SATARK flags, officers decide; every interception logged in the field.'),
    B('Screening counts when it converts', 'to enforcement action on the ground.'),
    B('Results before new cities:', 'expansion only after the current city shows measurable change.'),
]
PILOTS = [
    B('Scaled up only when there is a clear path to impact:', 'not research, not press, not marketing.'),
    B('20% of effort can go there', 'for the rest of the stuff, but it will not be core until we are very sure of the path to impact.'),
    B('Every pilot has a decision date:', 'a board-approved design and a scale-or-stop call. Nothing drifts on.'),
    B('Pilots borrow no core capacity:', 'Rakshak and compensation teams are not pulled away to run experiments.'),
]

# ---- helpers ----
blank = prs.slide_masters[0].slide_layouts[6]
slide = prs.slides.add_slide(blank)
SP = slide.shapes

def no_autofit(tf):
    el = tf._txBody.bodyPr
    for tag in ('a:normAutofit', 'a:spAutoFit'):
        e = el.find(qn(tag))
        if e is not None:
            el.remove(e)

def tbox(x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = SP.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    no_autofit(tf)
    return tf

def run(p, txt, sz, bold, color, font, spc=None):
    r = p.add_run(); r.text = txt
    r.font.size = Pt(sz); r.font.bold = bold
    r.font.color.rgb = color; r.font.name = font
    if spc:
        r.font._rPr.set('spc', str(spc))
    return r

def rrect(x, y, w, h, radius=0.16):
    sh = SP.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    sh.adjustments[0] = radius / min(w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = WHITE
    sh.line.color.rgb = LINE; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh

# ---- background + stars ----
bg = SP.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = WHITE; bg.line.fill.background()
bg.shadow.inherit = False
for (sx, sy, ss) in ((12.40, 0.32, 0.50), (12.10, 0.72, 0.30)):
    st = SP.add_shape(MSO_SHAPE.STAR_4_POINT, I(sx), I(sy), I(ss), I(ss))
    st.fill.solid(); st.fill.fore_color.rgb = BRAND; st.line.fill.background()
    st.shadow.inherit = False

# ---- header ----
tf = tbox(0.70, 0.55, 12.0, 0.30)
run(tf.paragraphs[0], 'NON-NEGOTIABLES', 10, True, BRAND, 'Montserrat', spc=240)

tf = tbox(0.70, 0.92, 12.0, 0.50)
p = tf.paragraphs[0]
run(p, 'Non-negotiables ', 27, True, INK, 'Montserrat')
run(p, 'in each project.', 27, True, BRAND, 'Montserrat')

tf = tbox(0.70, 1.50, 12.0, 0.28)
run(tf.paragraphs[0], 'Where effort must go this year, and what does not count as progress.', 10.5, False, MUT, 'Geist')

# ---- cards ----
CT = 1.98          # cards top
CB = 6.95          # cards bottom
GAP = 0.24
CW = (11.93 - 2 * GAP) / 3          # 3.8167
X1, X2, X3 = 0.70, 0.70 + CW + GAP, 0.70 + 2 * (CW + GAP)
SAT_H = 2.15
PIL_T = CT + SAT_H + 0.16

def card(x, y, w, h, icon_key, name, bullets, body_sz=8):
    rrect(x, y, w, h)
    disc = SP.add_shape(MSO_SHAPE.OVAL, I(x + 0.18), I(y + 0.16), I(0.40), I(0.40))
    disc.fill.solid(); disc.fill.fore_color.rgb = DISC; disc.line.fill.background()
    disc.shadow.inherit = False
    SP.add_picture(io.BytesIO(icons[icon_key]), I(x + 0.26), I(y + 0.24), I(0.24), I(0.24))
    tf = tbox(x + 0.70, y + 0.16, w - 0.88, 0.40, anchor=MSO_ANCHOR.MIDDLE)
    run(tf.paragraphs[0], name, 12, True, INK, 'Montserrat')
    tf = tbox(x + 0.18, y + 0.70, w - 0.36, h - 0.86)
    for i, (lead, rest) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.14
        p.space_after = Pt(4.5)
        run(p, '•  ', body_sz, True, BRAND, 'Geist')
        run(p, lead + ' ', body_sz, True, INK, 'Geist')
        run(p, rest, body_sz, False, MUT, 'Geist')

card(X1, CT, CW, CB - CT, 'hard-hat', 'Project Rakshak', RAKSHAK)
card(X2, CT, CW, CB - CT, 'scale', 'Hit & Run Compensation', COMP)
card(X3, CT, CW, SAT_H, 'radar', 'SATARK', SATARK)
card(X3, PIL_T, CW, CB - PIL_T, 'sprout', 'Other Pilot Programmes', PILOTS)

# ---- footer ----
tf = tbox(0.70, 7.12, 8.0, 0.30)
run(tf.paragraphs[0], 'CRASHFREE INDIA  ·  AOP 2026-27', 8.5, True, MUT, 'Montserrat', spc=160)
tf = tbox(11.40, 7.12, 1.50, 0.30)
tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
run(tf.paragraphs[0], '12 / 14', 8.5, True, MUT, 'Montserrat', spc=160)

# ---- move new slide to position 12 (after KPIs, index 11) ----
lst = prs.slides._sldIdLst
ids = list(lst)
new_id = ids[-1]
lst.remove(new_id)
lst.insert(11, new_id)

# ---- renumber existing footers NN / 13 -> NN / 14 with shift ----
pat = re.compile(r'^(\d{2}) / 13$')
for pos, s in enumerate(prs.slides, start=1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if pat.match(r.text.strip()):
                    r.text = f'{pos:02d} / 14'

prs.save(OUT)
print('saved', OUT, 'slides:', len(prs.slides.__iter__.__self__._sldIdLst))
