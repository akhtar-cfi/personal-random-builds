#!/usr/bin/env python3
"""CFI AOP 2026-27 — flat pivot-ready bifurcation, 4-level classification.
Vertical | Category (focus area) | Cost group (high-level) | Further bifurcation (line item)
| Description | Budget allocated | Remarks for Keshav.  Board-approved Rs 6.00 Cr."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

INDIGO="4A35FF"; INK="1A1C1C"; MUT="6B6980"; WHITE="FFFFFF"; GREEN="0A9955"; AMBER="B26A00"; DK="2A2740"
RUPEE='"₹"#,##0'
VTINT={"Trust Admin":"EAE7FF","Technology":"E0F1EF","Education":"FBEFD8","Enabler":"E4F4EA",
       "Infrastructure":"E4EDFB","Policy":"F0E7FA","Post-emergency care":"FBE7EC","Contingency":"EBECF1"}
thin=Side(style="thin", color="E4E5EE"); border=Border(left=thin,right=thin,top=thin,bottom=thin)

# (Vertical, Category[focus area], Cost group[high-level], Line item[further bifurcation], Description, Budget, Remarks)
ROWS=[
 # ---- PROJECT RAKSHAK ----
 ("Infrastructure","Project Rakshak","Institutional track","Student team stipends","₹36,000/team on completion · ~90 sites",3240000,
  "GL: Salary & Incentives (stipends). Milestone-linked — accrue quarterly, release on sign-off. Volume driver = # teams; true-up to actual cohort. Confirm TDS on student stipends with CA. Programme (restricted)."),
 ("Infrastructure","Project Rakshak","Institutional track","Expert mentor honoraria","Faculty & road-safety expert mentors",1300000,
  "GL: Professional fees. TDS u/s 194J (10%). Per-session/mentor — keep engagement letters."),
 ("Infrastructure","Project Rakshak","Institutional track","Campus outreach, onboarding & merchandise","Chapter mobilisation, IDs, kits, ads",960000,
  "GL: Marketing & Branding + Office. Mixed (merch/print + digital ads + travel). Claim GST input where invoiced."),
 ("Infrastructure","Project Rakshak","NGO partner track","NGO field-team engagement","~20 sites via NGO partners",2000000,
  "GL: Professional fees / grant-to-partner. Confirm TDS 194C vs grant treatment. Milestone per site. Restricted."),
 ("Infrastructure","Project Rakshak","NGO partner track","CFI desk + 10–15% physical re-checks","QA on NGO submissions",500000,
  "GL: Travel + Professional fees. Mostly staff time + travel; book travel to Travel head."),
 ("Infrastructure","Project Rakshak","Certified RSA track","CRRI-certified auditor fees (IRC SP:88)","~40 sites, NHAI-usable audits",4500000,
  "GL: Professional fees. Largest single vendor line. TDS 194J + GST input. Milestone/per-audit — keep PO + deliverable."),
 ("Infrastructure","Project Rakshak","Expert Consortium & Audit Manual","Expert Consortium (4 meetings)","Sep / Dec / Mar / Jun sittings",1000000,
  "GL: Professional fees (194J) + Travel + venue. Accrue per meeting."),
 ("Infrastructure","Project Rakshak","Expert Consortium & Audit Manual","Audit Manual v1.0 production","Authoring, design, print",800000,
  "GL: Professional fees + print. Programme opex; GST input on print. Mostly one-time."),
 ("Technology","Project Rakshak","Public dashboard","Public accountability dashboard","Build + hosting; identification→implementation tracker",1200000,
  "GL: Tech Infra. Split build (one-time) vs hosting (recurring). If in-house, it is staff time — confirm vendor vs internal."),
 ("Infrastructure","Project Rakshak","Programme management & field ops","Programme management (staff allocation)","PM + coordination time on Rakshak",1800000,
  "GL: Salary & Incentives — ALLOCATION, not a new payout. Ensure NOT double-counted in Trust Admin salaries."),
 ("Infrastructure","Project Rakshak","Programme management & field ops","Field travel & logistics","Site visits, DRSC meetings",700000,
  "GL: Travel. Reimbursements, no TDS."),
 ("Infrastructure","Project Rakshak","Programme management & field ops","National forum & DRSC presentations","12+ DRSC presentations, year-end forum",500000,
  "GL: Events + Travel. Venue & logistics — GST input on vendor invoices."),
 # ---- HIT & RUN COMPENSATION ----
 ("Post-emergency care","Hit & Run Compensation","Scheme audit & documentation","Stage-wise pipeline audit fieldwork (Gurugram)","FIR→GI Council; 20–30 claims tracked",1200000,
  "GL: Professional fees + Travel. Field researchers/surveyors — TDS 194C/194J. Restricted programme."),
 ("Policy","Hit & Run Compensation","Scheme audit & documentation","'Where the Pipeline Breaks' report","Scheme audit submitted to MoRTH",600000,
  "GL: Professional fees + print. Writing/design + printing. GST input."),
 ("Post-emergency care","Hit & Run Compensation","Institutional capacity building","Fellow-in-Residence stipend (Tehsildar office)","Embedded fellow",1200000,
  "GL: Salary & Incentives (fellowship). On HRMS retainer — recurring monthly, TDS on salary. Tag to Hit & Run."),
 ("Post-emergency care","Hit & Run Compensation","Institutional capacity building","SOP reference cards & training-gap assessment","District capacity",800000,
  "GL: Professional fees + print. Consultant + printing. Mostly one-time."),
 ("Post-emergency care","Hit & Run Compensation","Victim identification & AASHA","DLSA-partnered helpdesks","Trauma-ward legal helpdesks",800000,
  "GL: Professional fees / grant + Office. Confirm DLSA MoU terms."),
 ("Post-emergency care","Hit & Run Compensation","Victim identification & AASHA","AASHA QR deployment","Police stations & trauma centres",600000,
  "GL: Tech Infra + Marketing. QR standees/print + app hosting."),
 ("Policy","Hit & Run Compensation","Policy audit & advocacy","Quarterly findings & public brief","Publishing cadence",500000,
  "GL: Professional fees + Marketing. Recurring quarterly — accrue per quarter."),
 ("Policy","Hit & Run Compensation","Policy audit & advocacy","Advocacy & MoRTH engagement travel","Stakeholder engagement",300000,
  "GL: Travel. Reimbursements, no TDS."),
 # ---- CIVIC & GOVTECH (SATARK) ----
 ("Technology","Civic & GovTech (SATARK)","SATARK platform & scale-up","SATARK engineering & product","ANPR, ULIP, rule engine, officer app",2500000,
  "GL: Tech Infra / Professional fees. In-house devs = salary cross-charge (avoid double count); vendor = 194J + GST. Confirm capex vs opex."),
 ("Technology","Civic & GovTech (SATARK)","SATARK platform & scale-up","City deployment & rollout","Toward 30 cities · Gurugram, Pune, Ahmedabad first",2000000,
  "GL: Professional fees + Travel + Tech. Volume driver = # cities; milestone per go-live."),
 ("Technology","Civic & GovTech (SATARK)","SATARK platform & scale-up","Cloud, ANPR & ULIP integration","Compute, APIs, cameras",1500000,
  "GL: Tech Infra. Recurring cloud — reconcile monthly to usage. Foreign SaaS → GST under RCM; watch forex."),
 ("Technology","Civic & GovTech (SATARK)","Open data stack","Open data product maintenance & hosting","Defect Repo, Parliament, SC Tracker, Crash Dashboard",1500000,
  "GL: Tech Infra. Recurring hosting for 4 products — monthly."),
 ("Technology","Civic & GovTech (SATARK)","Open data stack","Data engineering & updates","Daily/weekly refresh, 146 districts",1000000,
  "GL: Salary / Professional fees. Cross-charge if internal; vendor = 194J."),
 ("Technology","Civic & GovTech (SATARK)","GovTech exploration","GovTech scoping & partnerships","2+ new DPG opportunities",1000000,
  "GL: Professional fees + Travel. Exploratory — keep unspent visible; do not front-load."),
 ("Technology","Civic & GovTech (SATARK)","GovTech exploration","Pilot / DPG prototyping","ULIP / MoRTH ecosystem",500000,
  "GL: Tech Infra. Small pilots — release against scoped need."),
 # ---- MARKETING CAMPAIGN ----
 ("Education","Marketing Campaign","Media advocacy","Editorial partnerships, radio & podcast, PSA","Earned + owned media",1500000,
  "GL: Marketing & Branding. Agency/media invoices — TDS 194C/194J + GST input."),
 ("Education","Marketing Campaign","Digital advocacy","50 micro-influencers & branded short films","#SafeSwag creator programme",1400000,
  "GL: Marketing & Branding. Influencer payouts — 194C/194J; keep contracts. Track CAC (see Aug ads learning)."),
 ("Education","Marketing Campaign","Digital advocacy","Digital IP & content production","Short films, reels",600000,
  "GL: Marketing & Branding. Production vendor + GST input."),
 ("Education","Marketing Campaign","Offline awareness","OOH in 3–4 cities","Billboards, transit",1000000,
  "GL: Marketing & Branding. OOH media buy — 194C + GST input."),
 ("Education","Marketing Campaign","Offline awareness","Guerilla installations & event co-branding","",500000,
  "GL: Marketing & Branding / Events."),
 ("Education","Marketing Campaign","Community mobilization","College competitions & app-based engagement","Real-life stories",1000000,
  "GL: Marketing + Office. Keep winner records for prize payouts."),
 # ---- BETS ----
 ("Infrastructure","Bets — School Safety Index","School Safety Index (seed)","IRC:SP:32 rubric validation & Zone Captain pilot","~100 school zones in one city",700000,
  "GL: Professional fees + Travel. SEED — release ONLY against board-approved design (stage gate). Ring-fence at ₹0.10 Cr."),
 ("Infrastructure","Bets — School Safety Index","School Safety Index (seed)","Per-school costing & board-ready design","Gate to Year-3 funded index",300000,
  "GL: Professional fees. SEED. Deliverable = board paper."),
 ("Policy","Bets — Gig Rider","Gig Rider (seed)","Gig crash-data scoping study","National study parked till model proven",600000,
  "GL: Professional fees. SEED — researcher fees 194J. Ring-fence at ₹0.10 Cr."),
 ("Policy","Bets — Gig Rider","Gig Rider (seed)","Platform + OMI Foundation engagements","3+ platforms at the table",400000,
  "GL: Travel + Professional fees. SEED. Relationship-building."),
 # ---- ORG BACKBONE (TRUST ADMIN) ----
 ("Trust Admin","Org Backbone","Core team salaries","Core team salaries & benefits","Aggregate core staff (individual figures withheld)",4200000,
  "GL: Salary & Incentives. Core/shared payroll ONLY (not cross-charged to programmes). Ensure programme staff-allocations are NOT double-counted. Monthly; TDS 192, PF/ESI."),
 ("Enabler","Org Backbone","Fellowship (SMIF)","Safer Mobility Impact Fellowship","~6 fellows across tracks (aggregate)",3600000,
  "GL: Salary & Incentives (fellowship). HRMS fixed-term agreements. Recurring monthly. Individual amounts confidential. If a fellow is dedicated to a programme, consider cross-charging it."),
 ("Trust Admin","Org Backbone","IRSC institutional support","IRSC institutional support & expert consultation","Shared expert/leadership",4200000,
  "GL: Professional fees. TDS 194J + GST. IRSC is a related party — confirm inter-entity treatment; keep MoU."),
 ("Technology","Org Backbone","IT, SaaS & AI tooling","AI tooling — Claude (Team) + Codex / AI subscriptions","AI assistants for build/research/ops",600000,
  "GL: Tech Infra (software subscriptions). ← dedicated AI-tooling line. Recurring monthly, mostly USD (Anthropic/OpenAI) — GST under RCM on import of service (pay & claim) + forex/card fees. Receipts already reaching finance."),
 ("Technology","Org Backbone","IT, SaaS & AI tooling","Google Workspace & email","Per-seat",300000,
  "GL: Tech Infra. Recurring per-seat — scales with headcount. RCM watch."),
 ("Technology","Org Backbone","IT, SaaS & AI tooling","Slack & collaboration","Per-seat",200000,
  "GL: Tech Infra. Recurring per-seat. RCM watch."),
 ("Technology","Org Backbone","IT, SaaS & AI tooling","HRMS 'CFI People' & other SaaS","",200000,
  "GL: Tech Infra. HRMS + misc SaaS. Recurring."),
 ("Technology","Org Backbone","IT, SaaS & AI tooling","Cloud, hosting & domains","Not product-specific",100000,
  "GL: Tech Infra. Domains/hosting. Recurring."),
 ("Trust Admin","Org Backbone","M&E & audit","M&E, statutory & impact audit","Statutory audit + M&E",600000,
  "GL: Audit & Compliance. TDS 194J. Year-end weighted."),
 ("Trust Admin","Org Backbone","Overheads (travel/branding/compliance)","Travel, branding & compliance","Org travel, branding, compliance",1100000,
  "GL: mixed (Travel + Marketing + Compliance). Split into 3 lines if you want cleaner cost-head pivots."),
 # ---- CONTINGENCY ----
 ("Contingency","Contingency","Contingency reserve","Contingency reserve (4%)","Held against programme + admin spend",2400000,
  "GL: Contingency (do NOT pre-allocate). Release only on approved overruns and RE-TAG to the actual head when spent."),
]
assert sum(r[5] for r in ROWS)==60000000, sum(r[5] for r in ROWS)

wb=Workbook()
ws=wb.active; ws.title="AOP Bifurcation"
heads=["Vertical","Category (focus area)","Cost group (high-level)","Further bifurcation (line item)","Description","Budget allocated","Remarks for Keshav — tagging & finance notes"]
for j,h in enumerate(heads,1):
    c=ws.cell(row=1,column=j,value=h); c.font=Font(name="Calibri",size=10.5,bold=True,color=WHITE)
    c.fill=PatternFill("solid",fgColor=INDIGO); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=border
ws.row_dimensions[1].height=38
for i,(v,cat,grp,item,desc,amt,rem) in enumerate(ROWS):
    r=i+2; band="FFFFFF" if i%2==0 else "F7F8FB"; vt=VTINT.get(v,"F2F2F5")
    def put(col,val,align="left",fmt=None,size=9.5,color=INK,bold=False,wrap=True,fill=band,valign="center"):
        cc=ws.cell(row=r,column=col,value=val); cc.font=Font(name="Calibri",size=size,bold=bold,color=color)
        cc.alignment=Alignment(horizontal=align,vertical=valign,wrap_text=wrap); cc.fill=PatternFill("solid",fgColor=fill); cc.border=border
        if fmt: cc.number_format=fmt
    put(1,v,bold=True,fill=vt,color=DK)
    put(2,cat,bold=True,color=INDIGO)
    put(3,grp,bold=True,color=DK,size=9.5)
    put(4,item)
    put(5,desc,size=9,color=MUT)
    put(6,amt,align="right",fmt=RUPEE,wrap=False,bold=True)
    put(7,rem,size=8.5,valign="top")
    ws.row_dimensions[r].height=32
last=len(ROWS)+1
ws.auto_filter.ref=f"A1:G{last}"
for col,w in zip("ABCDEFG",[16,22,25,30,32,14,58]): ws.column_dimensions[col].width=w
ws.freeze_panes="A2"; ws.sheet_view.showGridLines=False

# ---- Sheet 2: Summary ----
wc=wb.create_sheet("Summary"); wc.sheet_view.showGridLines=False
DB=f"'AOP Bifurcation'!$F$2:$F${last}"; DV=f"'AOP Bifurcation'!$A$2:$A${last}"; DC=f"'AOP Bifurcation'!$B$2:$B${last}"; DG=f"'AOP Bifurcation'!$C$2:$C${last}"
def band(cell,txt,fill=INDIGO,size=11,span=None):
    if span: wc.merge_cells(span)
    c=wc[cell]; c.value=txt; c.font=Font(bold=True,color=WHITE,size=size); c.fill=PatternFill("solid",fgColor=fill); c.alignment=Alignment(horizontal="left",vertical="center")
def hdr(row,cols):
    for j,t in enumerate(cols,1):
        c=wc.cell(row=row,column=j,value=t); c.font=Font(bold=True,color=WHITE,size=9.5); c.fill=PatternFill("solid",fgColor=INK)
        c.alignment=Alignment(horizontal=("left" if j==1 else "right"),vertical="center"); c.border=border
band("A1","Reconciliation — ties to the AOP (₹6.00 Cr)",span="A1:C1"); wc.row_dimensions[1].height=24
wc["A2"]="Grand total"; wc["A2"].font=Font(bold=True); wc["A2"].border=border
gt=wc["B2"]; gt.value=f"=SUM({DB})"; gt.number_format=RUPEE; gt.font=Font(bold=True); gt.alignment=Alignment(horizontal="right"); gt.border=border
wc["C2"]="expected ₹6,00,00,000"; wc["C2"].font=Font(color=MUT,size=9)
verts=[]; cats=[]
for v,cat,*_ in ROWS:
    if v not in verts: verts.append(v)
    if cat not in cats: cats.append(cat)
band("A4","By Vertical (functional)",fill=INK,size=10,span="A4:C4"); hdr(5,["Vertical","Budget (₹)","% of AOP"]); r=6
for v in verts:
    a=wc.cell(row=r,column=1,value=v); a.font=Font(size=10); a.fill=PatternFill("solid",fgColor=VTINT.get(v,"FFFFFF")); a.border=border
    b=wc.cell(row=r,column=2,value=f"=SUMIF({DV},A{r},{DB})"); b.number_format=RUPEE; b.alignment=Alignment(horizontal="right"); b.border=border
    p=wc.cell(row=r,column=3,value=f"=B{r}/$B$2"); p.number_format='0.0%'; p.font=Font(color=MUT,size=9); p.alignment=Alignment(horizontal="right"); p.border=border
    r+=1
wc.cell(row=r,column=1,value="Total").font=Font(bold=True); wc.cell(row=r,column=1).border=border
tb=wc.cell(row=r,column=2,value=f"=SUM(B6:B{r-1})"); tb.number_format=RUPEE; tb.font=Font(bold=True); tb.alignment=Alignment(horizontal="right"); tb.border=border; wc.cell(row=r,column=3).border=border
start=r+2
band(f"A{start}","By Category (core focus area)",fill=INK,size=10,span=f"A{start}:C{start}"); hdr(start+1,["Category (focus area)","Budget (₹)","% of AOP"]); r=start+2
for cat in cats:
    a=wc.cell(row=r,column=1,value=cat); a.font=Font(size=10,bold=True,color=INDIGO); a.border=border
    b=wc.cell(row=r,column=2,value=f"=SUMIF({DC},A{r},{DB})"); b.number_format=RUPEE; b.alignment=Alignment(horizontal="right"); b.border=border
    p=wc.cell(row=r,column=3,value=f"=B{r}/$B$2"); p.number_format='0.0%'; p.font=Font(color=MUT,size=9); p.alignment=Alignment(horizontal="right"); p.border=border
    r+=1
wc.cell(row=r,column=1,value="Total").font=Font(bold=True); wc.cell(row=r,column=1).border=border
cb=wc.cell(row=r,column=2,value=f"=SUM(B{start+2}:B{r-1})"); cb.number_format=RUPEE; cb.font=Font(bold=True); cb.alignment=Alignment(horizontal="right"); cb.border=border; wc.cell(row=r,column=3).border=border
for col,w in zip("ABC",[34,18,12]): wc.column_dimensions[col].width=w

# ---- Sheet 3: Read me ----
wr=wb.create_sheet("Read me"); wr.sheet_view.showGridLines=False; wr.column_dimensions["A"].width=120
lines=[
 ("CFI · AOP 2026–27 — budget bifurcation (flat, pivot-ready · 4-level)",INDIGO,13,True),("",WHITE,10,False),
 ("COLUMNS",GREEN,11,True),
 ("A Vertical (functional) · B Category (focus area / programme) · C Cost group (high-level bifurcation) ·",INK,10,False),
 ("D Further bifurcation (line item) · E Description · F Budget allocated · G Remarks for Keshav.",INK,10,False),
 ("Every value repeats down the rows (no merged cells) so you can pivot at any of the four levels.",INK,10,False),
 ("",WHITE,10,False),
 ("HOW TO PIVOT",GREEN,11,True),
 ("Select 'AOP Bifurcation' ▸ Insert ▸ PivotTable. Rows: any of Vertical / Category / Cost group / line item; Values: Sum of Budget.",INK,10,False),
 ("'Summary' reconciles to ₹6.00 Cr with a by-Vertical and a by-Category view.",INK,10,False),
 ("",WHITE,10,False),
 ("BASIS",AMBER,11,True),
 ("EXACT (board-approved AOP): Category totals — Rakshak ₹1.85 Cr, Hit & Run ₹0.60, SATARK ₹1.00, Campaign ₹0.60,",INK,10,False),
 ("   Bets ₹0.20, Org Backbone ₹1.51, Contingency ₹0.24 = ₹6.00 Cr; plus Trust Admin's five headers.",INK,10,False),
 ("FIRST-PASS (edit freely): cost-group & line-item splits inside programmes, the ₹42L core vs ₹36L SMIF split,",INK,10,False),
 ("   the ₹0.14 Cr AI/SaaS itemisation, and the functional Vertical tag on each line.",INK,10,False),
 ("",WHITE,10,False),
 ("FOR KESHAV",INDIGO,11,True),
 ("Remarks give GL head + TDS/GST/RCM + recurring-vs-one-time per line. Watch double-count on 'staff allocation' lines.",INK,10,False),
 ("Fellows: Rakshak stipends, Hit & Run Fellow-in-Residence, SMIF (Org Backbone). AI tooling is its own line (IT/SaaS/AI group).",INK,10,False),
 ("",WHITE,10,False),
 ("Sources: finalized AOP deck (₹6.00 Cr) · Team Charter / Ownership Map · Slack · Gmail (Anthropic receipts, SMIF, Codex).",MUT,9,False),
]
for i,(t,c,s,b) in enumerate(lines,1):
    cc=wr.cell(row=i,column=1,value=t); cc.font=Font(name="Calibri",size=s,bold=b,color=c); cc.alignment=Alignment(horizontal="left",vertical="center")

out="CFI_AOP_Bifurcation-for-Finance_v4_2026-08-19.xlsx"
wb.save(out); print("saved",out,"| rows:",len(ROWS),"| total:",sum(r[5] for r in ROWS))
