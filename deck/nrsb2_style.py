# NRSB v2 deck style — white paper, dark literal titles, sparse blue, tinted icon rows, photo-forward.
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageOps
from cfi_helpers import icon_png, _strip_style  # lucide icon rasterizer + shadow fix

INK   = RGBColor(0x16,0x18,0x1D)
MUT   = RGBColor(0x6E,0x71,0x80)
LINE  = RGBColor(0xE8,0xE9,0xEE)
MIST  = RGBColor(0xF7,0xF7,0xFA)
PAPER = RGBColor(0xFF,0xFF,0xFF)
BRAND = RGBColor(0x4A,0x35,0xFF)
BR_T  = RGBColor(0xEF,0xED,0xFF)   # brand tint
GREEN = RGBColor(0x0A,0x99,0x55)
GR_T  = RGBColor(0xE9,0xF7,0xF0)
AMBER = RGBColor(0xE5,0x89,0x00)
AM_T  = RGBColor(0xFD,0xF3,0xE3)
RED   = RGBColor(0xE1,0x19,0x00)
RD_T  = RGBColor(0xFD,0xEF,0xEC)
TEAL  = RGBColor(0x0E,0x8A,0x8A)
TE_T  = RGBColor(0xE7,0xF4,0xF4)
LAVW  = RGBColor(0xC8,0xC0,0xFF)

HEAD="Montserrat"; BODY="Geist"
SW=Inches(13.333); SH=Inches(7.5)
ML=Inches(0.62); MR=Inches(0.62)
CW=SW-ML-MR

BUILD=os.path.dirname(os.path.abspath(__file__))
SP=os.path.dirname(BUILD)
V5=os.path.join(SP,"assets/v5/deckready")
V3=os.path.join(SP,"assets/v3")
V3N=os.path.join(SP,"assets/v3/new")
CACHE=os.path.join(BUILD,"crop_cache"); os.makedirs(CACHE,exist_ok=True)
LOGO_BLUE=os.path.join(BUILD,"logo_blue.png"); LOGO_WHITE=os.path.join(BUILD,"logo_white.png")

def a5(name): return os.path.join(V5,name)
def a3(name): return os.path.join(V3N,name)

prs=None; blank=None
def new_deck():
    global prs, blank
    prs=Presentation(); prs.slide_width=SW; prs.slide_height=SH
    blank=prs.slide_layouts[6]
    return prs

def slide(bg=PAPER):
    s=prs.slides.add_slide(blank)
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,SW,SH)
    r.fill.solid(); r.fill.fore_color.rgb=bg; r.line.fill.background(); r.shadow.inherit=False
    _strip_style(r); return s

def rect(s,x,y,w,h,fill=None,line=None,lw=0.75,rounded=False,radius=None):
    shp=MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    r=s.shapes.add_shape(shp,x,y,w,h)
    if rounded and radius is not None:
        try: r.adjustments[0]=radius
        except Exception: pass
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb=fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb=line; r.line.width=Pt(lw)
    r.shadow.inherit=False; _strip_style(r); return r

def circle(s,x,y,d,fill,line=None):
    c=s.shapes.add_shape(MSO_SHAPE.OVAL,x,y,d,d); c.fill.solid(); c.fill.fore_color.rgb=fill
    if line is None: c.line.fill.background()
    else: c.line.color.rgb=line; c.line.width=Pt(0.75)
    c.shadow.inherit=False; _strip_style(c); return c

def hline(s,x,y,w,color=LINE,wt=0.75):
    ln=s.shapes.add_connector(1,x,y,x+w,y); ln.line.color.rgb=color; ln.line.width=Pt(wt)
    ln.shadow.inherit=False; _strip_style(ln); return ln

def _spc(run,p): run._r.get_or_add_rPr().set('spc',str(int(p*100)))

def text(s,x,y,w,h,t,*,font=BODY,size=12,bold=False,color=INK,align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP,ls=1.28,tracking=None,italic=False):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame
    tf.margin_left=tf.margin_right=0; tf.margin_top=tf.margin_bottom=0
    tf.word_wrap=True; tf.vertical_anchor=anchor
    if isinstance(t,str):
        for i,line in enumerate(t.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.alignment=align; p.line_spacing=ls
            r=p.add_run(); r.text=line
            r.font.name=font; r.font.size=Pt(size); r.font.bold=bold
            r.font.italic=italic; r.font.color.rgb=color
            if tracking is not None: _spc(r,tracking)
    else:
        p=tf.paragraphs[0]; p.alignment=align; p.line_spacing=ls
        for seg,st in t:
            r=p.add_run(); r.text=seg
            r.font.name=st.get('font',font); r.font.size=Pt(st.get('size',size))
            r.font.bold=st.get('bold',bold); r.font.italic=st.get('italic',italic)
            r.font.color.rgb=st.get('color',color)
            if 'tracking' in st: _spc(r,st['tracking'])
    return tb

def icon(s,name,x,y,size,hexcolor,sw=2.0):
    return s.shapes.add_picture(icon_png(name,hexcolor,sw=sw),x,y,width=size,height=size)

def icon_disc(s,name,x,y,d,tint,hexcolor,frac=0.52,sw=2.0):
    circle(s,x,y,d,tint)
    isz=Emu(int(d*frac)); off=Emu(int(d*(1-frac)/2))
    icon(s,name,x+off,y+off,isz,hexcolor,sw=sw)

_HEXMAP={}
def hexs(rgb): return f"#{rgb}"

def crop_fill(path,w_in,h_in,tag=""):
    """Pre-crop image to exactly fill w×h (inches) box; returns cached path."""
    key=f"{os.path.basename(path)}_{w_in:.2f}x{h_in:.2f}{tag}".replace('/','_').replace(' ','_')
    out=os.path.join(CACHE,key+".jpg")
    if not os.path.exists(out):
        im=ImageOps.exif_transpose(Image.open(path)).convert("RGB")
        tw,th=w_in,h_in; ar_t=tw/th; ar=im.width/im.height
        if ar>ar_t:
            nw=int(im.height*ar_t); x0=(im.width-nw)//2; im=im.crop((x0,0,x0+nw,im.height))
        else:
            nh=int(im.width/ar_t); y0=max(0,int((im.height-nh)*0.32)); im=im.crop((0,y0,im.width,y0+nh))
        im.thumbnail((1600,1600)); im.save(out,quality=87)
    return out

def photo(s,path,x,y,w,h=None,border=True,caption=None,fill=False,cap_color=MUT):
    """Place photo. fill=True crops to exact w×h; else scales by width."""
    if fill and h is not None:
        p=s.shapes.add_picture(crop_fill(path,w/914400,h/914400),x,y,width=w,height=h)
    else:
        p=s.shapes.add_picture(path,x,y,width=w)
        if h is None: h=p.height
    if border:
        rect(s,x,y,w,h,fill=None,line=LINE,lw=1.0)
    if caption:
        text(s,x,y+h+Inches(0.05),w,Inches(0.25),caption,size=8.5,color=cap_color,ls=1.1)
    return p

def chip(s,x,y,w,h,t,fill=MIST,color=INK,size=9.5,bold=True,line=None):
    rect(s,x,y,w,h,fill=fill,line=line,rounded=True,radius=0.5)
    text(s,x+Inches(0.02),y,w-Inches(0.04),h,t,font=HEAD,size=size,bold=bold,color=color,
         align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,ls=1.0)

def sect_tab(s,cur,name):
    """running-agenda chip, top-right."""
    w=Inches(2.9)
    text(s,SW-MR-w,Inches(0.30),w,Inches(0.24),f"SECTION {cur} OF 5  ·  {name}",
         font=HEAD,size=8,bold=True,color=MUT,tracking=1.6,align=PP_ALIGN.RIGHT)

def header(s,rail,title_runs,lead=None,sect=None):
    """rail: '02 · WHAT WE HAVE BUILT'; title_runs: list of (txt, brand?) tuples; returns content y."""
    if sect: sect_tab(s,*sect)
    text(s,ML,Inches(0.30),Inches(8),Inches(0.24),rail.upper(),font=HEAD,size=9,bold=True,color=BRAND,tracking=2.2)
    hline(s,ML,Inches(0.62),Inches(0.55),color=BRAND,wt=2.0)
    runs=[(t,{'color':BRAND} if b else {'color':INK}) for t,b in title_runs]
    text(s,ML,Inches(0.76),Inches(12.1),Inches(0.95),runs,font=HEAD,size=22.5,bold=True,ls=1.12)
    ny=Inches(1.62)
    if lead:
        text(s,ML,ny,Inches(11.9),Inches(0.4),lead,size=11.5,color=MUT,ls=1.3)
        ny=ny+Inches(0.44)
    return ny

def footer(s,wg,idx,total):
    hline(s,ML,Inches(7.06),CW)
    text(s,ML,Inches(7.14),Inches(8),Inches(0.25),
         f"Crashfree India  ·  Expression of Interest — Working Group {wg}",size=8.5,color=MUT,ls=1.0)
    text(s,SW-MR-Inches(1.2),Inches(7.14),Inches(1.2),Inches(0.25),f"{idx:02d} / {total:02d}",
         size=8.5,color=MUT,align=PP_ALIGN.RIGHT,ls=1.0)

def stat(s,x,y,w,value,label,color=BRAND,vsize=26,lsize=9.5,card=False,h=Inches(1.05)):
    if card: rect(s,x,y,w,h,fill=MIST,rounded=True,radius=0.14)
    pad=Inches(0.14) if card else 0
    text(s,x+pad,y+(Inches(0.10) if card else 0),w-2*pad,Inches(0.55),value,font=HEAD,size=vsize,bold=True,color=color,ls=1.0)
    text(s,x+pad,y+(Inches(0.60) if card else Inches(0.5)),w-2*pad,h-Inches(0.6),label,size=lsize,color=MUT,ls=1.12)

def irow(s,x,y,w,name,title,desc,tint,col,d=Inches(0.4),tsize=12,dsize=10.5,dy=Inches(0.0)):
    icon_disc(s,name,x,y,d,tint,col)
    text(s,x+d+Inches(0.18),y-Inches(0.02)+dy,Inches(2.6),Inches(0.5),title,font=HEAD,size=tsize,bold=True,color=INK,ls=1.05)
    text(s,x+d+Inches(0.18)+Inches(2.7),y-Inches(0.02)+dy,w-d-Inches(2.95),Inches(0.75),desc,size=dsize,color=MUT,ls=1.22)

def num_disc(s,x,y,d,n,color=BRAND,tint=BR_T):
    circle(s,x,y,d,tint)
    text(s,x,y-Inches(0.012),d,d,str(n),font=HEAD,size=13,bold=True,color=color,
         align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,ls=1.0)

def save(path):
    prs.save(path); print("saved",path)
