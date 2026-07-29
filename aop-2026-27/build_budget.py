import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BRAND = "4A35FF"; LAV = "F3F1FF"; DARK="1A1C1C"; MUTED="777589"; BORDERC="D9D9E3"
thin = Side(style="thin", color=BORDERC)
box = Border(left=thin,right=thin,top=thin,bottom=thin)

wb = openpyxl.Workbook()

# (Vertical, Programme, Line item, Basis of estimate, [Q1,Q2,Q3,Q4] in Rs Lakh)
D = [
("Project Rakshak","Cohort 2 — Audit Pipeline","Student team stipends","100 teams x Rs 36,000 (4-month cohort; performance-linked tranches)",[9,18,9,0]),
("Project Rakshak","Cohort 2 — Audit Pipeline","Mentor honorarium","25 faculty/expert mentors x Rs 24,000",[1.5,2.5,2,0]),
("Project Rakshak","Cohort 2 — Audit Pipeline","Certified RSA audits (Track C)","~40 NH/complex-junction sites x Rs 25,000 (CRRI-certified RSAs, IRC SP:88)",[3,4,3,0]),
("Project Rakshak","Cohort 2 — Audit Pipeline","NGO partner grants (Track B)","~20 district/peri-urban sites; incl. 10-15% physical re-checks by CFI",[3,5,4,0]),
("Project Rakshak","Programme Management","Programme team (fellows, PM, advocacy)","2-3 fellows + programme staff stipends, 12 months",[8,9,8,8]),
("Project Rakshak","Institutionalisation","Expert Consortium + Rakshak Audit Manual v1.0","4 consortium meetings/yr + manual drafting, design & print",[2,1,2,1]),
("Project Rakshak","Institutionalisation","Public dashboard & tech maintenance","Hosting, dev upkeep of accountability dashboard",[1.25,1.25,1.25,1.25]),
("Project Rakshak","Institutionalisation","Post-intervention impact audits","Before/after assessment at implemented sites",[0,2,3,3]),
("Project Rakshak","Engagement & Visibility","Rakshak Finale + national awards","Annual finale event + prize pool for top teams (trimmed from 50L in prior draft)",[0,0,10,20]),
("Project Rakshak","Engagement & Visibility","Travel & field logistics","Team + mentor travel to anchor districts and audit sites",[3,3,2,2]),
("Project Rakshak","Engagement & Visibility","Media, documentation & communication","Site stories, audit report design, public communication",[3,4,4,4]),
("Project Rakshak","Engagement & Visibility","DRSC & stakeholder engagement","12+ DRSC presentations; district administration consultations",[2,3,2,2]),
("Project Rakshak","Engagement & Visibility","Campus ambassadors & outreach","Student mobilisation across partner universities",[2,2,1,0]),
("Civic & GovTech (SATARK)","SATARK Platform","Product engineering team","2-3 engineers + product lead: SATARK rule engine, officer enforcement app, control-room views (team size & rates assumed - no cost data in SATARK docs)",[10,10,10,10]),
("Civic & GovTech (SATARK)","SATARK Platform","New city deployments & police integration","Software-side go-lives for the 30-city scale target (Gurugram, Pune, Ahmedabad first). ASSUMES camera/billboard/infra costs are borne by city police & partners - CFI funds integration, rules setup & go-live support (~Rs 0.7L/city; validate)",[2,6,6,6]),
("Civic & GovTech (SATARK)","SATARK Platform","Cloud, ULIP API & data infrastructure","Hosting, ANPR data pipelines, ULIP query volumes, security & uptime",[3,3,3,3]),
("Civic & GovTech (SATARK)","SATARK Platform","Officer training & field support","Role-based training, enforcement-drive support, false-positive review loop with police",[1,2,2.5,2.5]),
("Civic & GovTech (SATARK)","Data Stack Public Goods","Four open data products","Defect Repository, Road Safety in Parliament, SC Litigation Tracker, Crash Data Dashboard: cloud automation, AI pipeline costs, human-review console, annual RAI refresh",[3,3,3,3]),
("Civic & GovTech (SATARK)","Data Stack Public Goods","Product design, dashboards & impact evaluation","Design upkeep + instrumentation: interceptions logged, data citations in DRSCs/press",[2,2,2,2]),
("Hit & Run Compensation","Programme Management","Vertical team","Vertical lead + 2 research fellows (~Rs 50K/mo) + rolling intern roster (~Rs 15K/mo)",[7.5,7.5,7.5,7.5]),
("Hit & Run Compensation","Field Operations","Field travel","Gurugram + SW Delhi + Jaipur visits; watchlist district trips (air + lodging + per diem)",[2,2,2,2]),
("Hit & Run Compensation","Field Operations","AASHA go-to-market","QR standees + print collateral at PS/hospitals/DRSC venues; WhatsApp Business hosting",[2,1,0.5,0.5]),
("Hit & Run Compensation","Advocacy","Stakeholder roundtables","3-4 roundtables anchored to quarterly audit-findings releases",[1.5,1.5,1.5,1.5]),
("Hit & Run Compensation","Audit & Research","District process mapping & comparative notes","Stage-wise pipeline audit (Gurugram deep-dive) + comparative mapping toward 8+ districts by Jun 2027 (SW Delhi, Jaipur, watchlist)",[2,3,1,1]),
("Hit & Run Compensation","Advocacy","Victim narratives, scheme audit report & public brief","Case studies, year-end audit report production, 4-6 pg public brief",[0.5,1,1.5,2]),
("Bets (Seed Fund)","School Safety Index (bet)","Rubric validation + one-city pilot","IRC:SP:32-2023 rubric (44 checks / 13 sections) validated on ~100 schools in one city; released only against a board-approved pilot design",[0,2,4,4]),
("Bets (Seed Fund)","Gig Rider Safety (bet)","Scoping study + platform engagement","Crash-registry scoping + first conversations with delivery platforms & OMI; released only against a board-approved study design",[0,2,4,4]),
("Marketing Campaign","Safety is the New Swag","Digital advocacy","50 micro-influencers, branded short films, #SafeSwag digital IPs",[2,5,6,5]),
("Marketing Campaign","Safety is the New Swag","Media advocacy","Editorial partnerships, radio/podcast integration, PSA distribution",[1,4,5,4]),
("Marketing Campaign","Safety is the New Swag","Offline awareness","OOH in 3-4 cities, guerilla installations, event co-branding",[0,3,5,4]),
("Marketing Campaign","Safety is the New Swag","Community mobilization","College competitions, app-based engagement, real-life stories",[1,2,3,2]),
("Marketing Campaign","Safety is the New Swag","Campaign management","Dedicated campaign resource (was untagged 'people cost' in prior draft)",[2,2,2,2]),
("Trust Admin","People","Salaries & benefits","Core team salaries, EPF, insurance, allowances",[19.5,19.5,19.5,19.5]),
("Trust Admin","Institutional","Institutional support & expert consultation","IRSC institutional charge (~Rs 3.5L/mo) + leadership/expert consultation fees",[10.5,10.5,10.5,10.5]),
("Trust Admin","Technology","IT, SaaS & AI tools","Google Workspace, Slack, Zoom, Claude, Lovable, domain & cloud",[3.5,3.5,3.5,3.5]),
("Trust Admin","Governance","M&E + statutory & impact audit","Monitoring & evaluation, statutory audit, impact measurement",[0.5,1.5,1.5,2.5]),
("Trust Admin","Operations","Travel & networking","Leadership travel to conferences, stakeholder meetings",[1,2,2,2]),
("Trust Admin","Operations","Branding & outreach collateral","Brand kit, brochures, merchandise, content",[1,1,0.5,0.5]),
("Trust Admin","Operations","Virtual office & compliance","Registered virtual office rent for compliance",[0.25,0.25,0.25,0.25]),
("Contingency","Contingency","Contingency reserve (4%)","4% of Rs 6.00 Cr total budget; buffer for partner-driven and unbudgeted costs",[6,6,6,6]),
]

tot = sum(sum(q) for *_ ,q in D)
assert abs(tot-600)<1e-9, tot

ws = wb.active; ws.title = "Budget Detail"
ws.sheet_view.showGridLines = False
hdr = ["Vertical","Programme / Activity","Line Item","Basis of Estimate",
       "Q1 (Jul-Sep)","Q2 (Oct-Dec)","Q3 (Jan-Mar)","Q4 (Apr-Jun)","Total (Rs Lakh)","Total (Rs)"]
ws["A1"]="Crashfree India — AOP 2026-27 Budget Detail"
ws["A1"].font = Font(name="Arial", size=14, bold=True, color=DARK)
ws["A2"]="All quarterly figures in Rs Lakh. Blue cells are editable inputs; totals recalculate automatically. Period: July 2026 - June 2027."
ws["A2"].font = Font(name="Arial", size=9, italic=True, color=MUTED)
r0=4
for j,h in enumerate(hdr,1):
    c=ws.cell(row=r0,column=j,value=h)
    c.font=Font(name="Arial",size=10,bold=True,color="FFFFFF")
    c.fill=PatternFill("solid",fgColor=BRAND)
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    c.border=box
r=r0+1
band=False; prev_v=None
for (v,p,li,basis,qs) in D:
    if v!=prev_v: band = not band; prev_v=v
    fill = PatternFill("solid",fgColor=LAV) if band else None
    vals=[v,p,li,basis]+qs
    for j,val in enumerate(vals,1):
        c=ws.cell(row=r,column=j,value=val)
        c.border=box
        if fill and j<=4: c.fill=fill
        if j<=4:
            c.font=Font(name="Arial",size=9.5,color=DARK)
            c.alignment=Alignment(vertical="top",wrap_text=True)
        else:
            c.font=Font(name="Arial",size=9.5,color="0000FF")
            c.number_format='#,##0.00;-#,##0.00;"-"'
            c.alignment=Alignment(horizontal="right",vertical="top")
    ws.cell(row=r,column=9,value=f"=SUM(E{r}:H{r})")
    ws.cell(row=r,column=10,value=f"=I{r}*100000")
    for j in (9,10):
        c=ws.cell(row=r,column=j); c.border=box
        c.font=Font(name="Arial",size=9.5,bold=(j==9),color=DARK)
        c.number_format = '#,##0.00' if j==9 else '#,##0'
        c.alignment=Alignment(horizontal="right",vertical="top")
    r+=1
last=r-1
ws.cell(row=r,column=1,value="GRAND TOTAL")
for j in range(1,11):
    c=ws.cell(row=r,column=j); c.border=box
    c.fill=PatternFill("solid",fgColor=DARK)
    c.font=Font(name="Arial",size=10,bold=True,color="FFFFFF")
for j,col in zip(range(5,11),"EFGHIJ"):
    c=ws.cell(row=r,column=j,value=f"=SUM({col}{r0+1}:{col}{last})")
    c.number_format = '#,##0' if j==10 else '#,##0.00'
    c.alignment=Alignment(horizontal="right")
widths=[20,26,34,52,11,11,11,11,13,14]
for j,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes="A5"

ws2 = wb.create_sheet("Budget Summary",0)
ws2.sheet_view.showGridLines=False
ws2["A1"]="Crashfree India — Annual Operating Plan 2026-27"
ws2["A1"].font=Font(name="Arial",size=15,bold=True,color=DARK)
ws2["A2"]="Budget Summary  |  July 2026 - June 2027  |  Total: Rs 6.00 Cr"
ws2["A2"].font=Font(name="Arial",size=10,color=MUTED)
verts=["Project Rakshak","Civic & GovTech (SATARK)","Hit & Run Compensation","Marketing Campaign","Trust Admin","Bets (Seed Fund)","Contingency"]
hdr2=["Vertical","Q1 (Jul-Sep)","Q2 (Oct-Dec)","Q3 (Jan-Mar)","Q4 (Apr-Jun)","Total (Rs Lakh)","Total (Rs Cr)","% of Total"]
r0=4
for j,h in enumerate(hdr2,1):
    c=ws2.cell(row=r0,column=j,value=h)
    c.font=Font(name="Arial",size=10,bold=True,color="FFFFFF")
    c.fill=PatternFill("solid",fgColor=BRAND); c.border=box
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
r=r0+1
for v in verts:
    ws2.cell(row=r,column=1,value=v)
    for j,col in zip(range(2,6),"EFGH"):
        ws2.cell(row=r,column=j,value=f"=SUMIF('Budget Detail'!$A$5:$A${last},$A{r},'Budget Detail'!{col}$5:{col}${last})")
    ws2.cell(row=r,column=6,value=f"=SUM(B{r}:E{r})")
    ws2.cell(row=r,column=7,value=f"=F{r}/100")
    ws2.cell(row=r,column=8,value=f"=F{r}/$F${r0+len(verts)+1}")
    for j in range(1,9):
        c=ws2.cell(row=r,column=j); c.border=box
        c.font=Font(name="Arial",size=10,bold=(j in(1,6)),color=DARK)
        if j in (2,3,4,5,6): c.number_format='#,##0.00;-#,##0.00;"-"'
        if j==7: c.number_format='0.00'
        if j==8: c.number_format='0.0%'
        if j>1: c.alignment=Alignment(horizontal="right")
        if r%2==0 and j==1: c.fill=PatternFill("solid",fgColor=LAV)
    r+=1
ws2.cell(row=r,column=1,value="TOTAL")
for j,col in zip(range(2,9),"BCDEFGH"):
    ws2.cell(row=r,column=j,value=f"=SUM({col}{r0+1}:{col}{r-1})")
for j in range(1,9):
    c=ws2.cell(row=r,column=j); c.border=box
    c.fill=PatternFill("solid",fgColor=DARK)
    c.font=Font(name="Arial",size=10,bold=True,color="FFFFFF")
    if j in (2,3,4,5,6): c.number_format='#,##0.00'
    if j==7: c.number_format='0.00'
    if j==8: c.number_format='0.0%'
    if j>1: c.alignment=Alignment(horizontal="right")
for j,w in enumerate([26,13,13,13,13,15,13,11],1): ws2.column_dimensions[get_column_letter(j)].width=w
note_r=r+2
ws2.cell(row=note_r,column=1,value="Programme spend (Rakshak + Civic & GovTech + H&R + Marketing + Bets seed): Rs 4.25 Cr (~71%)  |  Trust Admin: Rs 1.51 Cr (~25%)  |  Contingency: Rs 0.24 Cr (4%)")
ws2.cell(row=note_r,column=1).font=Font(name="Arial",size=9,italic=True,color=MUTED)

ws3=wb.create_sheet("Notes & Assumptions")
ws3.sheet_view.showGridLines=False
notes=[
("HOW TO USE THIS WORKBOOK",""),
("Editable cells","Blue-font cells on 'Budget Detail' (quarterly amounts, Rs Lakh) are the only inputs. All totals, the summary sheet, and %s recalculate automatically."),
("Structure","Each line = Vertical > Programme > Line Item with an explicit basis of estimate (unit x rate wherever available)."),
("Period","Q1 = Jul-Sep 2026, Q2 = Oct-Dec 2026, Q3 = Jan-Mar 2027, Q4 = Apr-Jun 2027."),
("",""),
("KEY ASSUMPTIONS (from AOP source documents; verify before board sign-off)",""),
("Rakshak","100 student teams @ Rs 36K (Cohort 2 AOP); 25 mentors @ Rs 24K; Track C = ~40 sites @ ~Rs 25K/site RSA fee (rate assumed - confirm with empanelment EOIs); NGO Track B grants Rs 12L."),
("Civic & GovTech","NEW vertical, set at exactly Rs 1.00 Cr per Akhtar's instruction. Target: SATARK scaled to 30 cities by Jun 2027 (live today: Jaipur + Bengaluru) + four open data products + exploration of new govtech/DPG opportunities. SATARK docs contain NO cost data - every line is a planning assumption. CRITICAL: Rs 1 Cr funds the software/integration side only; a 30-city rollout assumes city-side infra (ANPR cameras, billboards) is funded by police/partners. Validate before board sign-off."),
("Hit & Run","Sized to the Rs 47-66L envelope in the Crash Claim AOP review deck (set at Rs 60L). Vertical lead compensation assumed inside 'Vertical team' line, per org policy."),
("Bets (Seed Fund)","School Safety Index and Gig Rider Safety are no longer funded programmes this year - each carries a Rs 10L seed, released only against a board-approved pilot/study design. If a bet proves out, it returns as a funded Year 3 programme."),
("Marketing","'Safety is the New Swag' campaign consolidated at Rs 60L: all four campaign pillars now carry budgets (prior draft had 3 of 4 pillars at Rs 0) + Rs 8L campaign management (was an untagged Rs 30L 'people cost')."),
("Trust Admin","Salaries Rs 78L; IRSC institutional charge Rs 42L (=~Rs 3.5L/mo, matches mid-year actuals); IT+AI tools consolidated at Rs 14L."),
("Contingency","Exactly 4% of Rs 6.00 Cr = Rs 24L (prior draft said '4%' but held Rs 20L = 3.2%)."),
("",""),
("CHANGES IN v4 (vs v3, Jul 2026 - per Akhtar's restructure instruction)",""),
("1","School Safety Index removed as a funded vertical (was Rs 100L) - moved to the Bets seed fund at Rs 10L."),
("2","Gig Rider Safety removed as a funded vertical (was Rs 20L) - moved to the Bets seed fund at Rs 10L."),
("3","NEW Civic & GovTech (SATARK) vertical added at exactly Rs 100L: SATARK platform Rs 80L + data-stack public goods Rs 20L."),
("4","Grand total unchanged at Rs 600.0L; programme spend unchanged at Rs 425L (~71%); contingency unchanged at Rs 24L (4%)."),
("",""),
("CHANGES vs 'Budget Bifurcation_v2' (prior draft, Rs 6.20 Cr - carried from v3)",""),
("1","Total rebalanced from Rs 619.7L to exactly Rs 600.0L."),
("2","Rakshak trimmed Rs 210L -> Rs 185L: Finale+prizes 50->30L, media 25->15L; consortium/manual line added (institutionalisation focus)."),
("3","Marketing: 3 campaign pillars had descriptions but Rs 0 budget; all 4 pillars now funded; Rs 30L untagged people cost reduced to Rs 8L campaign management."),
("4","Missing 'Rs Lakh' values fixed (salaries, AI tools, marketing people cost had blank lakh column)."),
("5","Every line now has quarterly phasing + basis of estimate; summary + quarterly rollups are formula-driven."),
("6","H&R kept at Rs 60L (within the Rs 47-66L envelope from the vertical AOP deck)."),
]
r=1
for a,b in notes:
    ca=ws3.cell(row=r,column=1,value=a); cb=ws3.cell(row=r,column=2,value=b)
    if b=="" and a: ca.font=Font(name="Arial",size=11,bold=True,color="4A35FF")
    else:
        ca.font=Font(name="Arial",size=9.5,bold=True,color=DARK)
        cb.font=Font(name="Arial",size=9.5,color=DARK)
        cb.alignment=Alignment(wrap_text=True,vertical="top")
    r+=1
ws3.column_dimensions["A"].width=22; ws3.column_dimensions["B"].width=120

wb.save("CFI_AOP_Budget_2026-27_v4.xlsx")
print("saved, target total (lakh):", tot)
# NOTE: after saving, run the xlsx skill's recalc.py so formula caches are populated.
