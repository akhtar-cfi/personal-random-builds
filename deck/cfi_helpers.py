# Crashfree India deck helper library — per CFI design skill §14.2/14.3 (canonical, reuse verbatim)
import os
import cairosvg
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BRAND      = RGBColor(0x4A,0x35,0xFF)  # primary indigo — line-2 titles, bars, key numbers
BRAND_LITE = RGBColor(0xF3,0xF1,0xFF)  # lavender wash — card fills, chips
WHITE      = RGBColor(0xFF,0xFF,0xFF)
DARK       = RGBColor(0x1A,0x1C,0x1C)  # headings + body
MUTED      = RGBColor(0x77,0x75,0x89)  # secondary text, captions, footers
BORDER     = RGBColor(0xED,0xEE,0xF2)  # hairlines, card outlines
SUCCESS    = RGBColor(0x00,0xAA,0x44)  # status: on-track (NEVER decorative)
WARNING    = RGBColor(0xF5,0x7C,0x00)  # status: caution
DANGER     = RGBColor(0xF1,0x00,0x15)  # critical/death stats ONLY
LAV        = RGBColor(0xC8,0xC0,0xFF)  # light text on brand backgrounds
LAV_DK     = RGBColor(0x6A,0x55,0xFF)  # hairlines/glyphs inside brand blocks

HEAD = "Montserrat"
BODY = "Geist"

SLIDE_W = Inches(13.333); SLIDE_H = Inches(7.5)
M_L = Inches(0.7); M_R = Inches(0.7)
CONTENT_W = SLIDE_W - M_L - M_R

BUILD_DIR = os.path.dirname(os.path.abspath(__file__))
ICON_DIR = os.path.join(BUILD_DIR, "node_modules/lucide-static/icons")
ICON_CACHE = os.path.join(BUILD_DIR, "icon_cache")
os.makedirs(ICON_CACHE, exist_ok=True)

def icon_png(name, hexcolor, px=256, sw=2.0):
    key = f"{name}_{hexcolor}_{px}_{sw}".replace('#','')
    out = os.path.join(ICON_CACHE, key+".png")
    if os.path.exists(out): return out
    svg = open(os.path.join(ICON_DIR, name+".svg")).read()
    svg = svg.replace('stroke="currentColor"', f'stroke="{hexcolor}"')
    svg = svg.replace('stroke-width="2"', f'stroke-width="{sw}"')
    cairosvg.svg2png(bytestring=svg.encode(), write_to=out, output_width=px, output_height=px)
    return out

prs = Presentation(); prs.slide_width=SLIDE_W; prs.slide_height=SLIDE_H
blank = prs.slide_layouts[6]

def _strip_style(shape):
    """Remove <p:style> so LibreOffice doesn't apply theme effect (drop shadows)."""
    sp = shape._element
    for child in list(sp):
        if child.tag.endswith('}style'):
            sp.remove(child)

def new_slide(bg=WHITE):
    s = prs.slides.add_slide(blank)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,SLIDE_W,SLIDE_H)
    r.fill.solid(); r.fill.fore_color.rgb=bg; r.line.fill.background(); r.shadow.inherit=False
    _strip_style(r)
    return s

def rect(slide,x,y,w,h,fill,line=None,rounded=False):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp,x,y,w,h); s.fill.solid(); s.fill.fore_color.rgb=fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb=line; s.line.width=Pt(0.75)
    s.shadow.inherit=False; _strip_style(s); return s

def circle(slide,x,y,d,fill,line=None):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL,x,y,d,d); c.fill.solid(); c.fill.fore_color.rgb=fill
    if line is None: c.line.fill.background()
    else: c.line.color.rgb=line; c.line.width=Pt(0.75)
    c.shadow.inherit=False; _strip_style(c); return c

def hline(slide,x,y,w,color=BORDER,wt=0.75):
    ln=slide.shapes.add_connector(1,x,y,x+w,y); ln.line.color.rgb=color; ln.line.width=Pt(wt)
    ln.shadow.inherit=False; _strip_style(ln); return ln

def _spc(run,p): run._r.get_or_add_rPr().set('spc',str(int(p*100)))

def text(slide,x,y,w,h,t,*,font=BODY,size=14,bold=False,color=DARK,align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP,ls=1.3,tracking=None,italic=False):
    """t = str (\n => paragraphs) OR list[(seg, style_overrides_dict)] for mixed inline runs."""
    tb=slide.shapes.add_textbox(x,y,w,h); tf=tb.text_frame
    tf.margin_left=tf.margin_right=0; tf.margin_top=tf.margin_bottom=0
    tf.word_wrap=True; tf.vertical_anchor=anchor
    if isinstance(t,str):
        for i,line in enumerate(t.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.alignment=align; p.line_spacing=ls
            r=p.add_run(); r.text=line
            r.font.name=font; r.font.size=Pt(size); r.font.bold=bold
            r.font.italic=italic; r.font.color.rgb=color
            if tracking is not None: _spc(r,tracking)
    else:
        p=tf.paragraphs[0]; p.alignment=align; p.line_spacing=ls
        for seg,st in t:
            r=p.add_run(); r.text=seg
            r.font.name=st.get('font',font); r.font.size=Pt(st.get('size',size))
            r.font.bold=st.get('bold',bold); r.font.italic=st.get('italic',italic)
            r.font.color.rgb=st.get('color',color)
            if 'tracking' in st: _spc(r,st['tracking'])
    return tb

def eyebrow(slide,x,y,t,color=BRAND,w=Inches(12)):
    return text(slide,x,y,w,Inches(0.3),t.upper(),font=HEAD,size=10,bold=True,color=color,tracking=2.4,ls=1.0)

def icon(slide,name,x,y,size,hexcolor,sw=2.0):
    return slide.shapes.add_picture(icon_png(name,hexcolor,sw=sw),x,y,width=size,height=size)

def logo(slide,x,y,width,dark=False):
    p = os.path.join(BUILD_DIR, "logo_white.png" if dark else "logo_blue.png")
    return slide.shapes.add_picture(p,x,y,width=width)

def sparkle(slide,x,y,d,color=BRAND):
    try: st=slide.shapes.add_shape(MSO_SHAPE.STAR_4_POINT,x,y,d,d)
    except Exception: st=slide.shapes.add_shape(MSO_SHAPE.STAR_5_POINT,x,y,d,d)
    st.fill.solid(); st.fill.fore_color.rgb=color; st.line.fill.background(); st.shadow.inherit=False
    _strip_style(st); return st

def sparkles_tr(slide,color=BRAND):
    sparkle(slide,Inches(12.40),Inches(0.32),Inches(0.50),color)
    sparkle(slide,Inches(12.10),Inches(0.72),Inches(0.30),color)

def footer(slide,idx,total,label):
    text(slide,M_L,Inches(7.12),Inches(8),Inches(0.3),label,font=HEAD,size=8.5,color=MUTED,tracking=1.6,bold=True)
    text(slide,Inches(11.4),Inches(7.12),Inches(1.4),Inches(0.3),f"{idx:02d} / {total:02d}",
         font=HEAD,size=8.5,color=MUTED,tracking=1.6,bold=True,align=PP_ALIGN.RIGHT)

def std_header(slide, eb, title_l1, title_l2=None, lead=None):
    """CFI content-slide header: eyebrow + two-line title (line2 = brand) + optional lead.
       Returns the y-coordinate where body content can begin."""
    eyebrow(slide,M_L,Inches(0.55),eb)
    text(slide,M_L,Inches(0.92),Inches(12),Inches(0.6),title_l1,font=HEAD,size=27,bold=True,color=DARK,ls=1.04)
    ny=Inches(1.46)
    if title_l2:
        text(slide,M_L,Inches(1.42),Inches(12),Inches(0.6),title_l2,font=HEAD,size=27,bold=True,color=BRAND,ls=1.04)
        ny=Inches(2.02)
    if lead:
        text(slide,M_L,ny,Inches(11.9),Inches(0.5),lead,font=BODY,size=12.5,color=MUTED,ls=1.4)
        ny=ny+Inches(0.6)
    return ny

def save(path):
    prs.save(path)
    print("saved", path)
