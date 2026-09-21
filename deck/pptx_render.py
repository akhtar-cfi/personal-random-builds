#!/usr/bin/env python3
"""LibreOffice-free pptx renderer: python-pptx shapes -> PIL raster (PNG per slide + combined PDF).
Faithful enough for layout/overflow/color/bold/alignment QA; uses installed Montserrat/Geist VF.
Usage: python3 pptx_render.py deck.pptx [dpi]   ->  deck_pXX.png + deck.pdf
"""
import sys, os, io
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageDraw, ImageFont

EMU = 914400
_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

def raw_geo(shape):
    """Read (x,y,w,h) EMU from raw xfrm, tolerant of float coords python-pptx rejects."""
    try:
        xfrm = shape._element.find('.//{%s}xfrm' % _A)
        if xfrm is None:
            return None
        off = xfrm.find('{%s}off' % _A); ext = xfrm.find('{%s}ext' % _A)
        return (float(off.get('x')), float(off.get('y')), float(ext.get('cx')), float(ext.get('cy')))
    except Exception:
        return None
FONT_FILES = {
    "Montserrat": os.path.expanduser("~/.fonts/Montserrat-VF.ttf"),
    "Geist": os.path.expanduser("~/.fonts/Geist-VF.ttf"),
}
INK = (0x16, 0x18, 0x1D)
_fc = {}

def font(family, size_px, bold):
    size_px = max(6, int(round(size_px)))
    key = (family, size_px, bold)
    if key in _fc:
        return _fc[key]
    path = FONT_FILES.get(family, FONT_FILES["Geist"])
    try:
        f = ImageFont.truetype(path, size_px)
    except Exception:
        f = ImageFont.truetype(FONT_FILES["Geist"], size_px)
    try:
        f.set_variation_by_axes([700 if bold else 400])
    except Exception:
        pass
    _fc[key] = f
    return f

def hx(rgbcolor):
    try:
        s = str(rgbcolor)
        return (int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))
    except Exception:
        return None

def solid_fill(shape):
    try:
        if shape.fill.type == 1:  # MSO_FILL.SOLID
            return hx(shape.fill.fore_color.rgb)
    except Exception:
        pass
    return None

def line_props(shape):
    try:
        ln = shape.line
        if ln.fill.type == 1:
            c = hx(ln.color.rgb)
            w = ln.width.pt if ln.width is not None else 0.75
            return c, w
    except Exception:
        pass
    return None, None

def run_spc(run):
    try:
        v = run._r.get_or_add_rPr().get('spc')
        return int(v) / 100.0 if v is not None else 0.0  # points
    except Exception:
        return 0.0

def draw_text(d, shape, geo, DPI):
    tf = shape.text_frame
    gx, gy, gw, gh = geo
    L = gx + (tf.margin_left or 0)
    T = gy + (tf.margin_top or 0)
    W = gw - (tf.margin_left or 0) - (tf.margin_right or 0)
    Hbox = gh
    box_w_px = W / EMU * DPI
    x0 = L / EMU * DPI
    y0 = T / EMU * DPI

    def p2px(pt):
        return pt / 72.0 * DPI

    # build lines across paragraphs
    para_lines = []  # list of (lines, space_after_px) ; lines = list of list of tokens
    for para in tf.paragraphs:
        align = para.alignment or PP_ALIGN.LEFT
        ls = para.line_spacing if isinstance(para.line_spacing, float) else 1.2
        runs = para.runs
        # tokenize into words carrying style, splitting explicit \n as hard breaks
        tokens = []  # (text, fname, size_px, bold, color, spc_px, hardbreak_after)
        for r in runs:
            fname = r.font.name or "Geist"
            size = r.font.size.pt if r.font.size is not None else 12
            bold = bool(r.font.bold)
            col = hx(r.font.color.rgb) or INK
            spc = p2px(run_spc(r))
            segs = r.text.split("\n")
            for si, seg in enumerate(segs):
                words = seg.split(" ")
                for wi, w in enumerate(words):
                    tok = {"t": w, "f": fname, "s": p2px(size), "b": bold, "c": col, "spc": spc,
                           "space": (wi < len(words) - 1)}
                    tokens.append(tok)
                if si < len(segs) - 1 and tokens:
                    tokens[-1]["br"] = True
        # greedy wrap
        lines = []
        cur = []
        cur_w = 0.0
        def tok_w(tok):
            f = font(tok["f"], tok["s"], tok["b"])
            w = f.getlength(tok["t"])
            if tok["spc"]:
                w += tok["spc"] * max(0, len(tok["t"]))
            return w
        space_w = lambda tok: font(tok["f"], tok["s"], tok["b"]).getlength(" ")
        for tok in tokens:
            tw = tok_w(tok)
            add = tw + (space_w(tok) if cur else 0)
            if cur and cur_w + add > box_w_px + 1:
                lines.append(cur)
                cur = [tok]
                cur_w = tw
            else:
                cur.append(tok)
                cur_w += add
            if tok.get("br"):
                lines.append(cur)
                cur = []
                cur_w = 0.0
        if cur:
            lines.append(cur)
        if not lines:
            lines = [[]]
        para_lines.append((lines, align, ls, p2px(para.space_after.pt if para.space_after is not None else 0)))

    # measure total height
    def line_h(line, ls):
        mx = max([t["s"] for t in line], default=p2px(12))
        return mx * 1.2 * ls
    total_h = 0.0
    for lines, align, ls, sa in para_lines:
        for ln in lines:
            total_h += line_h(ln, ls)
        total_h += sa
    # vertical anchor
    va = tf.vertical_anchor
    box_h_px = Hbox / EMU * DPI
    if va == MSO_ANCHOR.MIDDLE:
        y = y0 + max(0, (box_h_px - total_h) / 2)
    elif va == MSO_ANCHOR.BOTTOM:
        y = y0 + max(0, box_h_px - total_h)
    else:
        y = y0
    # draw
    for lines, align, ls, sa in para_lines:
        for ln in lines:
            lh = line_h(ln, ls)
            # line width
            lw = 0.0
            for i, t in enumerate(ln):
                f = font(t["f"], t["s"], t["b"])
                w = f.getlength(t["t"]) + (t["spc"] * max(0, len(t["t"])) if t["spc"] else 0)
                lw += w
                if i < len(ln) - 1 and t["space"]:
                    lw += f.getlength(" ")
            if align == PP_ALIGN.CENTER:
                x = x0 + (box_w_px - lw) / 2
            elif align == PP_ALIGN.RIGHT:
                x = x0 + (box_w_px - lw)
            else:
                x = x0
            base = y + (lh - max([t["s"] for t in ln], default=p2px(12)) * 1.2)  # baseline-ish top
            for i, t in enumerate(ln):
                f = font(t["f"], t["s"], t["b"])
                if t["spc"]:
                    cx = x
                    for ch in t["t"]:
                        d.text((cx, base), ch, font=f, fill=t["c"])
                        cx += f.getlength(ch) + t["spc"]
                    x = cx
                else:
                    d.text((x, base), t["t"], font=f, fill=t["c"])
                    x += f.getlength(t["t"])
                if i < len(ln) - 1 and t["space"]:
                    x += f.getlength(" ")
            y += lh
        y += sa

def draw_shape(img, d, shape, DPI):
    st = str(shape.shape_type)
    geo = raw_geo(shape)
    if geo is None:
        return
    gx, gy, gw, gh = geo
    if "PICTURE" in st:
        try:
            blob = shape.image.blob
            im = Image.open(io.BytesIO(blob)).convert("RGBA")
            w = max(1, int(gw / EMU * DPI)); h = max(1, int(gh / EMU * DPI))
            im = im.resize((w, h))
            img.paste(im, (int(gx / EMU * DPI), int(gy / EMU * DPI)), im)
        except Exception:
            pass
        return
    x0 = gx / EMU * DPI; y0 = gy / EMU * DPI
    x1 = (gx + gw) / EMU * DPI; y1 = (gy + gh) / EMU * DPI
    fill = solid_fill(shape)
    lc, lw = line_props(shape)
    lw_px = max(1, int(round((lw or 0) / 72.0 * DPI))) if lc else 0
    ast = None
    try:
        ast = str(shape.auto_shape_type)
    except Exception:
        ast = None
    is_conn = "CONNECTOR" in st or "LINE" in st
    if is_conn and lc:
        d.line([(x0, y0), (x1, y1)], fill=lc, width=max(1, lw_px))
    elif ast and "OVAL" in ast:
        d.ellipse([x0, y0, x1, y1], fill=fill, outline=lc, width=lw_px if lc else 1)
    elif ast and "STAR" in ast:
        d.ellipse([x0, y0, x1, y1], fill=fill)
    elif ast and "ROUNDED_RECTANGLE" in ast:
        try:
            adj = shape.adjustments[0]
        except Exception:
            adj = 0.1
        r = max(1, int(adj * min(x1 - x0, y1 - y0)))
        d.rounded_rectangle([x0, y0, x1, y1], radius=r, fill=fill, outline=lc, width=lw_px if lc else 1)
    elif ast is not None or fill is not None or lc is not None:
        d.rectangle([x0, y0, x1, y1], fill=fill, outline=lc, width=lw_px if lc else 1)
    # text
    try:
        if shape.has_text_frame and shape.text_frame.text.strip():
            draw_text(d, shape, geo, DPI)
    except Exception:
        pass

def render(path, dpi=150):
    prs = Presentation(path)
    W = int(prs.slide_width / EMU * dpi); H = int(prs.slide_height / EMU * dpi)
    base = os.path.splitext(path)[0]
    pngs = []
    for i, slide in enumerate(prs.slides):
        img = Image.new("RGB", (W, H), "white")
        d = ImageDraw.Draw(img)
        for shape in slide.shapes:
            try:
                draw_shape(img, d, shape, dpi)
            except Exception as e:
                print("  shape err:", e)
        out = f"{base}_p{i+1:02d}.png"
        img.save(out)
        pngs.append(img)
    if pngs:
        pngs[0].save(f"{base}.pdf", save_all=True, append_images=pngs[1:], resolution=dpi)
    print(f"rendered {len(pngs)} pages -> {base}_pNN.png + {base}.pdf")
    return len(pngs)

if __name__ == "__main__":
    p = sys.argv[1]
    dpi = int(sys.argv[2]) if len(sys.argv) > 2 else 150
    render(p, dpi)
